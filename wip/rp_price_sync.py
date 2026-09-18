#!/usr/bin/env python3
"""
rp_price_sync.py - bring the RP WIP file's CONTRACT / ETC to the NEWEST DOCUMENT (shared/rp_price):
JobTread's approved proposal, the newest proposal PDF in the job folder, the takeoff cost sheet -
whichever is newest wins, a pick below what is already billed is held. Every change is written to
the change log (shared/wip_audit) with the paper that won and its date, and the estimator gets the
list of jobs where JobTread is missing, has no approved proposal, or is older than the folder.

Dry-run by default (repo rule 7): prints the table and changes nothing.
  python3 wip/rp_price_sync.py             # the table + the estimator list
  python3 wip/rp_price_sync.py --write     # update the RP WIP file (backup first) + the change log
Then the guarded writer carries it to the tab: rp_wip_reader --emit-review -> decisions -> --apply-review.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import shutil
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))

from openpyxl import load_workbook                              # noqa: E402
from openpyxl.styles import Font                                # noqa: E402
from shared import paths, jobtread, rp_price, wip_audit, pnl_paths  # noqa: E402
from shared.xlsx_verify import assert_clean                     # noqa: E402

GREEN = "FF00B050"


def rp_file() -> Path:
    return paths.get_path("RP_WIP_FILE", paths.onedrive_base() / "RP WIP TO FIX_Final.xlsx")


def master_path() -> Path:
    p = paths.get_path("WIP_EXCEL_PATH", paths.onedrive_base() / "Company Files - WIP Report" / "WIP - MASTER.xlsx")
    return p if p.exists() else p.with_name("WIP - MASTER new.xlsx")


def read_billed(master: Path) -> dict:
    """{line: billed to date} from the master's RP tab (the last QBO sync)."""
    try:
        wb = load_workbook(master, read_only=True)
    except OSError:
        return {}
    tab = next((t for t in ("Test - RP", "WIP-RP") if t in wb.sheetnames), None)
    if not tab:
        return {}
    ws = wb[tab]
    rows = list(ws.iter_rows(values_only=True))
    hr = next((i for i, r in enumerate(rows[:15]) if r and "PROJECT #" in [str(x).strip() for x in r if x]), None)
    if hr is None:
        return {}
    hdr = [str(x).strip() if x else "" for x in rows[hr]]
    pc, bc = hdr.index("PROJECT #"), hdr.index("BILLED TO DATE") if "BILLED TO DATE" in hdr else None
    out = {}
    for r in rows[hr + 1:]:
        p = str(r[pc] or "").strip().upper() if pc < len(r) else ""
        if p.startswith("RP") and bc is not None and bc < len(r):
            try:
                out[p] = float(r[bc]) if r[bc] not in (None, "") else None
            except (TypeError, ValueError):
                out[p] = None
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--write", action="store_true", help="update the RP WIP file (backup first) and the change log")
    ap.add_argument("--no-jobtread", action="store_true", help="skip JobTread (folder paper only)")
    ap.add_argument("--only", help="one line, e.g. RP7120 or RP6858-FTW")
    ap.add_argument("--json", help="also dump every resolution to this JSON path")
    a = ap.parse_args()

    RP, M = rp_file(), master_path()
    wb = load_workbook(RP)
    ws = wb["RP WIP"]
    lines = {}
    for r in range(3, ws.max_row + 1):
        j = str(ws.cell(r, 1).value or "").strip().upper()
        if j.startswith("RP") and j not in lines:
            lines[j] = r
    if a.only:
        lines = {k: v for k, v in lines.items() if k == a.only.upper()}
    billed = read_billed(M)
    jt = {} if a.no_jobtread else jobtread.jobs(set(lines), log=print)
    print(f"{len(lines)} line(s) · JobTread: {sum(1 for v in jt.values() if v['docs'])} with an approved proposal, {len(jt)} exist")

    def num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    results, changes, estimator = [], [], []
    for line, r in lines.items():
        K, E = num(ws.cell(r, 4).value), num(ws.cell(r, 5).value)
        folder = None
        for c in (15, 16):
            v = str(ws.cell(r, c).value or "")
            if v.startswith("/"):
                p = Path(v)
                folder = p if p.is_dir() else p.parent
                break
        if folder is None:
            folder = pnl_paths.rp_job_folder(line.replace("-FTW", ""))
        res = rp_price.resolve(line, folder, jt.get(line), K, E, billed.get(line),
                               str(ws.cell(r, 15).value or ""), str(ws.cell(r, 16).value or ""))
        res["row"] = r
        results.append(res)
        if res["changes"]["contract"] or res["changes"]["etc"]:
            changes.append(res)
        jt_flags = [f for f in res["flags"] if "JobTread" in f]
        if jt_flags:
            estimator.append((line, "; ".join(jt_flags)))

    print(f"\n{'LINE':12} {'CONTRACT now':>13} {'-> pick':>13}  {'ETC now':>12} {'-> pick':>12}  WHO WON (date)")
    for res in results:
        k, e = res["contract"], res["etc"]
        ck = f"{k['amount']:>13,.2f}" if (k and res["changes"]["contract"]) else f"{'same':>13}"
        ce = f"{e['amount']:>12,.2f}" if (e and res["changes"]["etc"]) else f"{'same':>12}"
        who = " · ".join(x for x in (
            f"K: {k['src'][:38]} ({rp_price.fmt_date(k['date'])})" if k and res["changes"]["contract"] else "",
            f"E: {e['src'][:38]} ({rp_price.fmt_date(e['date'])})" if e and res["changes"]["etc"] else "") if x)
        cur = res["current"]
        print(f"{res['line']:12} {(cur['contract'] or 0):>13,.2f} {ck}  {(cur['etc'] or 0):>12,.2f} {ce}  {who}")
        for f in res["flags"]:
            print(f"{'':12}   ⚠ {f}")
    print(f"\n{len(changes)} line(s) would change · {len(estimator)} for the estimator")
    if a.json:
        Path(a.json).write_text(json.dumps(results, default=str, indent=1), encoding="utf-8")

    if estimator:
        print("\nFOR THE ESTIMATOR (JobTread)")
        for line, why in estimator:
            print(f"  {line:12} {why}")

    if not a.write:
        print("\n(dry run - nothing written; add --write)")
        return 0
    if not changes:
        print("\nnothing to write")
        return 0
    backups = M.parent / "WIP History" / "Backups"
    backups.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M")
    shutil.copy(RP, backups / f"pre-pricesync_{stamp}_{RP.name}")
    now = dt.datetime.now().isoformat(timespec="seconds")
    today = dt.date.today().strftime("%m/%d/%Y")
    entries = []
    for res in changes:
        r, line = res["row"], res["line"]
        k, e = res["contract"], res["etc"]
        note_bits = []
        if res["changes"]["contract"]:
            old = num(ws.cell(r, 4).value)
            ws.cell(r, 4).value = round(k["amount"], 2)
            ws.cell(r, 15).value = f"{k['src']} ({rp_price.fmt_date(k['date'])})"
            entries.append({"at": now, "tab": "RP WIP file", "project_no": line, "field": "contract", "old": old, "new": round(k["amount"], 2),
                            "source": f"{k['src']} dated {rp_price.fmt_date(k['date'])} - newest document wins", "actor": "rp_price_sync", "run": f"rp_price_sync --write {today}", "note": None})
            note_bits.append(f"contract {k['src']} {rp_price.fmt_date(k['date'])}")
        if res["changes"]["etc"]:
            old = num(ws.cell(r, 5).value)
            ws.cell(r, 5).value = round(e["amount"], 2)
            ws.cell(r, 16).value = f"{e['src']} ({rp_price.fmt_date(e['date'])})"
            entries.append({"at": now, "tab": "RP WIP file", "project_no": line, "field": "etc", "old": old, "new": round(e["amount"], 2),
                            "source": f"{e['src']} dated {rp_price.fmt_date(e['date'])} - newest document wins", "actor": "rp_price_sync", "run": f"rp_price_sync --write {today}", "note": None})
            note_bits.append(f"ETC {e['src']} {rp_price.fmt_date(e['date'])}")
        for c in (4, 5):
            ws.cell(r, c).number_format = '"$"#,##0.00'
            ws.cell(r, c).font = Font(name="Calibri", size=11, color=GREEN, bold=True)
        if num(ws.cell(r, 4).value) and num(ws.cell(r, 5).value):
            ws.cell(r, 12).value = "Ready"
        ws.cell(r, 13).value = f"{today}: newest document wins - {'; '.join(note_bits)} · " + str(ws.cell(r, 13).value or "")
    wb.save(RP)
    assert_clean(RP)
    n = wip_audit.log_changes(entries)
    print(f"\nwrote {len(changes)} line(s) to {RP.name} (backup pre-pricesync_{stamp}) · change log: {n} entries")
    return 0


if __name__ == "__main__":
    sys.exit(main())
