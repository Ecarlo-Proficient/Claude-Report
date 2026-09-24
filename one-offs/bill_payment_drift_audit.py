#!/usr/bin/env python3
"""
bill_payment_drift_audit.py - checks QBO quietly rewrote after they were paid.

Read-only on the raw QBO mirror (shared/qbo_mirror). No QBO call, no QBO write.

When a PAID bill is edited or deleted, QBO rewrites the check that paid it on
its own - no prompt, no warning, the audit log says "Indirect edit by System":
  - the check goes (partly or fully) UNAPPLIED - money that left the bank now
    floats on the vendor as a credit, and every bill it paid reopens;
  - the re-entered copy of the bill shows OPEN - it looks owed a second time;
  - the loan-repayment / vendor credit that rode on the check is freed;
  - later, QBO AUTO-APPLIES the floating money to the next bill entered for
    that vendor - a check dated 08/28 ends up "paying" a bill from 09/15, so
    the next real check goes out short by that amount.
Subs get hit hardest (weekly bills + loan credits on every check). The worked
example is check 48314 (the owner 2026-09-24): paid UC41 on 08/28, UC41 was
deleted 09/17 and re-entered 09/18, the check floated $4,635.70 and QBO then
put $135.40 of it on UC044 (09/15).

For every bill payment since --since this flags:
  FLOATING   check total > what it applies to bills (net of credits)
  LATE LINK  check applied to a bill dated AND entered after the check was
             written - QBO put floating money on the next bill
  COPY OPEN  an open bill for the same vendor entered after the check but
             dated on/before it (the re-entered copy of a bill the check paid)
  CREDIT     an open vendor credit dated on/before the check (freed from it)
and names the cause when the mirror still holds it (a deleted bill that linked
to the check; the change log's pre-strip copy of the check).

    python3 one-offs/bill_payment_drift_audit.py                 since 2025-01-01
    python3 one-offs/bill_payment_drift_audit.py --check 48314   one check, printed in full
    python3 one-offs/bill_payment_drift_audit.py --subs          subs only

Refresh the mirror first (python3 ledger/refresh_mirror.py) so it reads today.
Output: <CompanyHealth>/Bill Payment Drift.xlsx (one Table) + stdout summary.
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from shared import paths, xlsx_verify      # noqa: E402
from shared import qbo_mirror as mirror    # noqa: E402
from shared.qbo_costs import sub_evidence  # noqa: E402

EPS = 0.005
MIN_FLOAT = 1.00           # cents of rounding on old checks are not drift


def _t(s) -> dt.datetime | None:
    try:
        return mirror.parse_iso(str(s))
    except (TypeError, ValueError):
        return None


def _d(s) -> dt.date | None:
    try:
        return dt.date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def _fmt(d) -> str:
    return d.strftime("%m/%d/%Y") if d else ""


def _money(x) -> str:
    return f"${x:,.2f}"


def _created(rec: dict) -> dt.datetime | None:
    return _t((rec.get("MetaData") or {}).get("CreateTime"))


def _vendor(rec: dict) -> tuple:
    r = rec.get("VendorRef") or {}
    return str(r.get("value") or ""), str(r.get("name") or "")


def _links(pay: dict) -> list:
    """[(TxnType, TxnId, amount)] - what the payment applies, line by line."""
    out = []
    for ln in pay.get("Line") or []:
        for lt in ln.get("LinkedTxn") or []:
            out.append((lt.get("TxnType"), str(lt.get("TxnId")), round(float(ln.get("Amount") or 0), 2)))
    return out


def pay_url(txn_id: str) -> str:
    return f"https://qbo.intuit.com/app/billpayment?txnId={txn_id}"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--since", default="2025-01-01", help="check date floor (default 2025-01-01)")
    ap.add_argument("--check", help="one check number - print its full story")
    ap.add_argument("--subs", action="store_true", help="subs only")
    ap.add_argument("--out", type=Path, default=None)
    a = ap.parse_args()

    con = mirror.connect()
    print(f"mirror as of {mirror._meta_get(con, 'last_refresh') or '?'} (refresh first for today)")

    bills = {r["Id"]: r for r in mirror.load("Bill", include_deleted=True, con=con)}
    deleted_bills = {r[0] for r in con.execute("SELECT id FROM qbo_bill WHERE deleted=1")}
    credits = mirror.load("VendorCredit", con=con)
    pays = [p for p in mirror.load("BillPayment", con=con) if str(p.get("TxnDate") or "") >= a.since]

    sub_vendor = set()
    open_bills = defaultdict(list)
    for b in bills.values():
        vid = _vendor(b)[0]
        if sub_evidence(b.get("PrivateNote")):
            sub_vendor.add(vid)
        if b["Id"] not in deleted_bills and float(b.get("Balance") or 0) > EPS:
            open_bills[vid].append(b)
    open_credits = defaultdict(list)
    for c in credits:
        if float(c.get("Balance") or 0) > EPS:
            open_credits[_vendor(c)[0]].append(c)
    # deleted bills that still name the payment that paid them -> the cause
    killed_by = defaultdict(list)
    for bid in deleted_bills:
        b = bills.get(bid) or {}
        for lt in b.get("LinkedTxn") or []:
            if "BillPayment" in str(lt.get("TxnType")):
                killed_by[str(lt.get("TxnId"))].append(b)
    # the change log's strip events (kept from 09/23 on with the pre-strip copy)
    strips = defaultdict(list)
    for c in mirror.changes(con, entity="BillPayment", limit=100000):
        if "payment unapplied" in c["flags"] or (c.get("lines_before") or 0) > (c.get("lines_after") or 0):
            strips[c["rec_id"]].append(c)

    rows = []
    for p in pays:
        vid, vname = _vendor(p)
        is_sub = vid in sub_vendor
        if a.subs and not is_sub:
            continue
        if a.check and str(p.get("DocNumber") or "") != a.check:
            continue
        total = round(float(p.get("TotalAmt") or 0), 2)
        lk = _links(p)
        billed = sum(x[2] for x in lk if x[0] in ("Bill", "JournalEntry"))   # a JE to AP pays like a bill
        credited = sum(x[2] for x in lk if x[0] == "VendorCredit")
        floating = round(total - (billed - credited), 2)
        p_made, p_date = _created(p), _d(p.get("TxnDate"))

        late = []
        card = p.get("PayType") == "CreditCard"      # a card-fee bill added after an Amex payment is habit, not drift
        for typ, bid, amt in lk:
            if typ != "Bill" or card:
                continue
            b = bills.get(bid) or {}
            bm, bd = _created(b), _d(b.get("TxnDate"))
            # a check cannot pay for work billed after it was written: the bill is dated AFTER the
            # check and entered after it. A re-entered copy (dated on/before the check) re-applied
            # by hand is a repair, not drift; the 05/2025 import re-stamped CreateTime on both sides.
            if p_made and bm and bm > p_made + dt.timedelta(minutes=1) and bd and p_date and bd > p_date:
                later = [(pid, x) for pid, x in _other_payers(pays, bid, p["Id"])]
                late.append({"bill": b.get("DocNumber") or bid, "bill_date": _d(b.get("TxnDate")),
                             "entered": bm, "amount": amt, "others": later})
        copies, freed = [], []
        if floating > MIN_FLOAT or late:
            for b in open_bills.get(vid, []):
                bm, bd = _created(b), _d(b.get("TxnDate"))
                if p_made and bm and bm > p_made and bd and p_date and bd <= p_date:
                    copies.append(b)
            for c in open_credits.get(vid, []):
                cd = _d(c.get("TxnDate"))
                if cd and p_date and cd <= p_date and (p_date - cd).days <= 31:
                    freed.append(c)
        if floating <= MIN_FLOAT and not late and not a.check:
            continue

        cause = [f"bill {b.get('DocNumber') or b['Id']} deleted" for b in killed_by.get(p["Id"], [])]
        for s in strips.get(p["Id"], []):
            cause.append(f"stripped {_fmt(_d(s['changed_at']))} (before copy kept)")
        todo = []
        if copies:
            todo.append("re-apply this check to " + ", ".join(
                f"{b.get('DocNumber')} ({_fmt(_d(b.get('TxnDate')))}, open {_money(float(b.get('Balance') or 0))})"
                for b in copies) + " - do NOT pay the copy again")
        if freed:
            todo.append("put back credit " + ", ".join(
                f"{c.get('DocNumber') or c['Id']} {_money(float(c.get('Balance') or 0))}" for c in freed))
        for L in late:
            todo.append(f"take {_money(L['amount'])} off {L['bill']} (entered {_fmt(L['entered'])}, after this "
                        f"check) - {L['bill']} then shows {_money(L['amount'])} open = what the sub was short-paid")
        if floating > MIN_FLOAT and not copies and not late:
            todo.append("find the bill(s) this check paid and re-apply it")
        rows.append({
            "vendor": vname, "sub": "Sub" if is_sub else "", "check": p.get("DocNumber") or "",
            "date": p_date, "total": total, "applied": round(billed - credited, 2), "floating": floating,
            "late": "; ".join(f"{L['bill']} {_money(L['amount'])}" for L in late),
            "copy": "; ".join(f"{b.get('DocNumber')} open {_money(float(b.get('Balance') or 0))}" for b in copies),
            "credit": "; ".join(f"{c.get('DocNumber') or c['Id']} {_money(float(c.get('Balance') or 0))}" for c in freed),
            "cause": "; ".join(cause),
            "edited": _t((p.get("MetaData") or {}).get("LastUpdatedTime")),
            "todo": "; ".join(todo), "id": p["Id"], "_late": late, "_copies": copies, "_freed": freed, "_p": p,
        })
    con.close()

    rows.sort(key=lambda r: (r["sub"] != "Sub", -(r["floating"] + sum(L["amount"] for L in r["_late"]))))
    if a.check:
        for r in rows:
            _story(r, bills)
        if not rows:
            print(f"check {a.check}: not found since {a.since}")
        return 0

    n_sub = sum(1 for r in rows if r["sub"])
    print(f"\nchecks QBO rewrote after payment since {a.since}: {len(rows)} ({n_sub} subs)")
    print(f"  floating (paid, applied to nothing): {_money(sum(r['floating'] for r in rows if r['floating'] > MIN_FLOAT))}"
          f" on {sum(1 for r in rows if r['floating'] > MIN_FLOAT)} checks")
    print(f"  auto-applied to a bill entered later: {sum(1 for r in rows if r['_late'])} checks, "
          f"{_money(sum(L['amount'] for r in rows for L in r['_late']))}")
    print(f"  re-entered copy sitting OPEN (double-pay risk): {sum(len(r['_copies']) for r in rows)} bills, "
          f"{_money(sum(float(b.get('Balance') or 0) for r in rows for b in r['_copies']))}")
    for r in rows[:25]:
        print(f"  {r['sub']:<3} {r['check']:<8} {_fmt(r['date'])} {r['vendor'][:28]:<28} total {_money(r['total']):>12}"
              f"  floating {_money(r['floating']):>11}  {('late: ' + r['late']) if r['late'] else ''}")
    _write(rows, a.out or paths.companyhealth_dir() / "Bill Payment Drift.xlsx", a.since)
    return 0


def _other_payers(pays, bill_id, skip_id):
    for q in pays:
        if q["Id"] == skip_id:
            continue
        for typ, bid, amt in _links(q):
            if typ == "Bill" and bid == bill_id:
                yield q.get("DocNumber") or q["Id"], amt


def _story(r, bills) -> None:
    p = r["_p"]
    print(f"\ncheck {r['check']} · {r['vendor']}{' (sub)' if r['sub'] else ''} · dated {_fmt(r['date'])} "
          f"· written {_fmt(_created(p))} · last touched {_fmt(r['edited'])}")
    print(f"  total {_money(r['total'])} · applied now {_money(r['applied'])} · FLOATING {_money(r['floating'])}")
    for typ, bid, amt in _links(p):
        b = bills.get(bid) or {}
        print(f"    applied {_money(amt):>11} -> {typ} {b.get('DocNumber') or bid} dated {_fmt(_d(b.get('TxnDate')))} "
              f"entered {_fmt(_created(b))}")
    for L in r["_late"]:
        others = ", ".join(f"check {c} {_money(x)}" for c, x in L["others"]) or "nothing else"
        print(f"  LATE LINK: {L['bill']} was entered {_fmt(L['entered'])}, after this check - QBO put "
              f"{_money(L['amount'])} of the floating money on it; the rest of it was paid by {others}")
    for b in r["_copies"]:
        print(f"  COPY OPEN: {b.get('DocNumber')} dated {_fmt(_d(b.get('TxnDate')))} entered {_fmt(_created(b))} "
              f"open {_money(float(b.get('Balance') or 0))} - reads as owed again")
    for c in r["_freed"]:
        print(f"  CREDIT FREED: {c.get('DocNumber') or c['Id']} {_money(float(c.get('Balance') or 0))} "
              f"({(c.get('PrivateNote') or '')[:60]})")
    if r["cause"]:
        print(f"  cause: {r['cause']}")
    print(f"  to fix: {r['todo'] or '-'}")


def _write(rows, out: Path, since: str) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Checks"
    ws["A1"] = "Checks QBO rewrote after they were paid"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = (f"Since {_fmt(_d(since))} · run {dt.datetime.now().strftime('%m/%d/%Y %I:%M %p')} · "
                f"{len(rows)} checks, {sum(1 for r in rows if r['sub'])} subs")
    ws["A3"] = "Floating"
    ws["B3"] = sum(r["floating"] for r in rows if r["floating"] > MIN_FLOAT)
    ws["C3"] = "Auto-applied to later bills"
    ws["D3"] = sum(L["amount"] for r in rows for L in r["_late"])
    ws["E3"] = "Copies open (double-pay risk)"
    ws["F3"] = sum(float(b.get("Balance") or 0) for r in rows for b in r["_copies"])
    for c in ("B3", "D3", "F3"):
        ws[c].number_format = '"$"#,##0.00'
        ws[c].font = Font(bold=True)
    ws["B3"].fill = PatternFill("solid", fgColor="FCE4D6")
    ws["D3"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws["F3"].fill = PatternFill("solid", fgColor="F8CBAD")

    hdr = ["Vendor", "Sub", "Check #", "Check date", "Check total", "Applied now", "Floating",
           "Auto-applied to later bill", "Re-entered copy open", "Freed credit", "Cause",
           "Last changed", "To fix", "Open check"]
    H = 5
    for i, h in enumerate(hdr, start=1):
        ws.cell(H, i, h).font = Font(bold=True)
    for n, r in enumerate(rows, start=H + 1):
        vals = [r["vendor"], r["sub"], int(r["check"]) if str(r["check"]).isdigit() else r["check"],
                r["date"], r["total"], r["applied"], r["floating"], r["late"], r["copy"], r["credit"],
                r["cause"], r["edited"].date() if r["edited"] else None, r["todo"], "open"]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(n, i, v)
            if i in (4, 12):
                c.number_format = "mm/dd/yyyy"
            elif i in (5, 6, 7):
                c.number_format = '"$"#,##0.00'
        link = ws.cell(n, 14)
        link.hyperlink = pay_url(r["id"])
        link.font = Font(color="0563C1", underline="single")
    last = H + max(len(rows), 1)
    ref = f"A{H}:N{last}"
    t = Table(displayName="Drift", ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    body = f"A{H + 1}:N{last}"
    ws.conditional_formatting.add(f"G{H + 1}:G{last}", CellIsRule(operator="greaterThan", formula=["1"],
                                  fill=PatternFill("solid", fgColor="FCE4D6")))
    ws.conditional_formatting.add(f"H{H + 1}:H{last}", FormulaRule(formula=[f'LEN($H{H + 1})>0'],
                                  fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.conditional_formatting.add(f"I{H + 1}:I{last}", FormulaRule(formula=[f'LEN($I{H + 1})>0'],
                                  fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(body, FormulaRule(formula=[f'$B{H + 1}="Sub"'], font=Font(bold=True)))
    ws.freeze_panes = f"A{H + 1}"
    for i, w in enumerate((30, 6, 10, 11, 13, 13, 13, 30, 34, 22, 30, 12, 70, 8), start=1):
        ws.column_dimensions[ws.cell(H, i).column_letter].width = w
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    xlsx_verify.assert_clean(out)
    print(f"\nwrote {out}")


if __name__ == "__main__":
    sys.exit(main())
