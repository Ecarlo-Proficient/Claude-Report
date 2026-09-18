"""3/31/26 WIP as-of, done the simple way (the user 2026-08-08: "look at WIP
Master, copy it, update costs/billed from QBO at the cutoff").

The snapshot's 'WIP Master' tab already holds the WHOLE report — MFD + CP + RP
(all 90 residential jobs), with contract, ETC (TOTAL COSTS), and the WIP math.
So: copy it, grouped MFD → CP → RP, refresh ONLY billed + costs from QBO dated
on/before the cutoff, and pull CP's ETC from the commercial takeoff. If a CP
takeoff ETC implies gross profit over 30% it's almost certainly a bad read —
that row is left on the placeholder and reported to the owner instead.

No folder scanning of Residential — RP is already in the tab.
"""
import os
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pyxlsb import open_workbook as open_xlsb

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))
sys.path.insert(0, str(_REPO / "wip"))
from shared import qbo_api                                    # noqa: E402
from shared.takeoff_etc import find_takeoff_etc              # noqa: E402
from shared.xlsx_verify import assert_clean, safe_table_name  # noqa: E402

HIST = Path(os.getenv("WIP_HISTORY_DIR",
            str(Path.home() / "Library/CloudStorage/OneDrive-ProficientConcrete,LLC"
                / "Company Files - WIP Report" / "WIP History")))
SNAP, CUTOFF, LABEL = "WIP - 03-31-26.xlsb", "2026-03-31", "March 31, 2026"
CP_ROOTS = [Path("/Volumes/Common/CURRENT PROJECTS/Awarded Projects Commercial projects")]
CP_ROOTS.append(CP_ROOTS[0] / "Completed Projects")

# 0-based columns on the WIP Master tab
PROJ, CONTRACT, ETC, BILLED, COSTS, GPPCT = 0, 4, 5, 6, 7, 18
NCOL = 21
FONT = "Tahoma"
MONEY = '"$"#,##0_);[Red]("$"#,##0)'
PCT = "0.00%"
MONEY_COLS = {4, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15, 16, 17, 19, 20}   # 0-based $
PCT_COLS = {12, 18}
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
SEC_FILL = PatternFill("solid", fgColor="1F3864")
_thin = Side(style="thin", color="000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _key(s):
    """PROJECT cell → QBO key. PROJ_RE keeps the -FTW suffix so RP7607 (slab) and
    RP7607-FTW (flatwork) stay DISTINCT — they're separate QBO projects. A bare
    leading number (12-31 format) is MFD."""
    s = str(s or "").strip()
    m = qbo_api.PROJ_RE.search(s.upper())
    if m:
        return m.group(1)
    m2 = re.match(r"(\d{2,4})\b", s)
    return f"MFD{m2.group(1)}" if m2 else None


def _read_master():
    with open_xlsb(str(HIST / SNAP)) as wb:
        with wb.get_sheet("WIP Master") as sh:
            return [[c.v for c in row] for row in sh.rows()]


def _index_cp():
    idx = {}
    for root in CP_ROOTS:
        try:
            for e in root.iterdir():
                if e.is_dir():
                    m = re.search(r"(CP\d{3,4})", e.name.upper())
                    if m and m.group(1) not in idx:
                        idx[m.group(1)] = e
        except OSError:
            pass
    return idx


def build(out_dir):
    grid = _read_master()
    access, cid = qbo_api.load_credentials()
    proj_map = qbo_api.build_project_customer_map(access, cid)
    cp_folders = _index_cp()

    # count keys so a phased project (MFD192 across 3 rows) isn't QBO-split
    from collections import Counter
    occ = Counter(k for r in grid[5:] if r and len(r) > PROJ
                  for k in [_key(r[PROJ])] if k)

    buckets = {"MFD": [], "CP": [], "RP": []}
    over30, refreshed, phased = [], 0, 0
    cache = {}
    for r in grid[5:]:
        if not r or len(r) <= PROJ:
            continue
        key = _key(r[PROJ])
        if not key:
            continue
        while len(r) < NCOL:
            r.append(None)
        div = "MFD" if key.startswith("MFD") else "CP" if key.startswith("CP") else "RP"
        # billed + costs from QBO as of the cutoff (skip phased — can't split)
        if occ[key] == 1:
            if key not in cache:
                cust = proj_map.get(key)
                cache[key] = None
                if cust:
                    try:
                        t = qbo_api.extract_pl_totals(qbo_api.fetch_project_pl(
                            access, cid, cust["id"], "2019-01-01", CUTOFF))
                        cache[key] = (round(float(t.get("income", 0) or 0), 2),
                                      round((t.get("cogs", 0) or 0) + (t.get("expenses", 0) or 0), 2))
                    except Exception:
                        cache[key] = None
            if cache[key]:
                r[BILLED], r[COSTS] = cache[key]
                refreshed += 1
        else:
            phased += 1
        # CP ETC from the commercial takeoff — with the GP>30% guard
        if div == "CP":
            fol = cp_folders.get(key)
            k = r[CONTRACT] if isinstance(r[CONTRACT], (int, float)) else None
            if fol is not None and k:
                try:
                    _p, etc, note, _f = find_takeoff_etc(fol, key, "slab", "")
                except Exception:
                    etc, note = None, ""
                if isinstance(etc, (int, float)) and etc > 0:
                    gp = (k - etc) / k
                    if gp > 0.30:
                        over30.append((key, k, etc, gp, note))     # report, don't apply
                    else:
                        r[ETC] = round(etc, 2)
                        r[GPPCT] = gp                              # real GP%, not the 9.09% placeholder
        buckets[div].append(r)

    _write(grid, buckets, out_dir)
    print(f"  refreshed {refreshed} rows from QBO · {phased} phased kept as snapshot")
    print(f"  CP ETC applied where GP ≤ 30%; {len(over30)} CP jobs held back (GP>30%)")
    return over30


def _write(grid, buckets, out_dir):
    wb = Workbook()
    ws = wb.active
    ws.title = "WIP as of 3-31-2026"
    # title + 3-row header, copied from the original
    for i in range(5):
        for j in range(NCOL):
            v = grid[i][j] if i < len(grid) and j < len(grid[i]) else None
            c = ws.cell(i + 1, j + 1, v)
            c.font = Font(name=FONT, size=8, bold=(i >= 2))
            if i < 2 and isinstance(v, (int, float)) and 40000 < v < 60000:
                c.number_format = "mmmm d, yyyy"        # REPORT DATE serial → readable date
            if i >= 2:
                c.fill = HDR_FILL
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                c.border = BORDER
    row = 6
    for div, name in (("MFD", "MULTI-FAMILY"), ("CP", "COMMERCIAL"), ("RP", "RESIDENTIAL")):
        sec = ws.cell(row, 1, name)
        sec.font = Font(name=FONT, size=8, bold=True, color="FFFFFF")
        for j in range(1, NCOL + 1):
            ws.cell(row, j).fill = SEC_FILL
        row += 1
        for r in buckets[div]:
            for j in range(NCOL):
                c = ws.cell(row, j + 1, r[j] if j < len(r) else None)
                c.font = Font(name=FONT, size=8)
                c.border = BORDER
                if j in PCT_COLS:
                    c.number_format = PCT
                elif j in MONEY_COLS:
                    c.number_format = MONEY
            row += 1
    ws.freeze_panes = "A6"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 32
    from openpyxl.worksheet.table import Table, TableStyleInfo   # noqa: F401 (kept off — irregular header)
    out = out_dir / "WIP as of 3-31-2026.xlsx"
    wb.save(out)
    assert_clean(out)
    print(f"  ✓ verified clean → {out}")
    return out


if __name__ == "__main__":
    over = build(Path.home() / "Downloads")
    if over:
        print("\n  CP ETCs held back (takeoff implied GP > 30% — verify):")
        for key, k, etc, gp, note in over:
            print(f"    {key}: contract ${k:,.0f} · takeoff ETC ${etc:,.0f} → GP {gp*100:.0f}%  [{note}]")
