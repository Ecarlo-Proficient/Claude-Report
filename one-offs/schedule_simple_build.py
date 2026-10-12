"""schedule_simple_build.py - the plain version of the horizontal RP schedule (no tables).

The owner dropped the tables / hidden columns / formulas build (2026-10-11: "it needs to
flow, it needs to be simple"). This keeps only what needs no upkeep, on a COPY:

  1. The calendar runs Mon-Sat to 12/31 on Main AND Completed (same columns, so a cut
     row lands under the same dates), real dates in every day band (punch-list bands
     too), the month name over each month's first working day. Widths, borders and
     band look copied from the last existing day column of each sheet.
  2. Four cell styles named exactly as the Legend - Task Completed, Projected, Final,
     Hard Stops / Holds / Cancelled - Arial 8 bold, centred, thin border, in the shades
     the owner picked (contrast checked against black text). The Legend swatches use
     them; calendar entries already in the clerk's four colours are switched to them.
     The DATE READY greens and CREW yellows (the old daily-schedule codes) are untouched.
  3. Both sheets freeze below the header with the job columns fixed (I4).

  --dropdowns  also puts a phase dropdown on the Main calendar cells: flatwork rows offer
               the Legend's flatwork phases, foundation rows the foundation phases. Any
               other text is still accepted ("Set forms / Baldo").

  python3 one-offs/schedule_simple_build.py SRC.xlsx OUT.xlsx [--dropdowns]

Read-only on SRC. OUT must not be SRC. Ends with xlsx_verify.assert_clean.
"""
from __future__ import annotations

import copy
import datetime as dt
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter as L  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.views import Selection  # noqa: E402

from shared.xlsx_verify import assert_clean  # noqa: E402

MAIN, DONE, LEGEND = "Main Schedule", "Completed", "Legend"
CAL0 = 9  # column I: Monday 10/05
MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST",
          "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
WEEKDAYS = {"MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "THRUSDAY", "FRIDAY", "SATURDAY"}
DAY_FMT = "dddd d"
# legend name, new fill, the clerk's fill it replaces (theme colours as "theme<n>+<tint>")
STYLES = [
    ("Task Completed", "9DC3E6", "theme4+0.40"),
    ("Projected", "FFF06A", "FFFF00"),
    ("Final", "B4DE9C", "00B050"),
    ("Hard Stops / Holds / Cancelled", "F4A6A6", "FF7C80"),
]


def die(msg: str) -> None:
    sys.exit(f"STOP: {msg}")


def up(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper()


def fill_key(cell) -> str | None:
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    if c.type == "rgb":
        return str(c.rgb)[-6:].upper()
    if c.type == "theme":
        return f"theme{c.theme}" + (f"+{c.tint:.2f}" if c.tint else "")
    return None


def band_rows(ws) -> list:
    """Rows whose first calendar cell is a weekday name (or already a date)."""
    out = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, CAL0).value
        if isinstance(v, dt.datetime) or re.sub(r"[\d\s]", "", up(v)) in WEEKDAYS:
            out.append(r)
    return out


def calendar_days(ws, default=None) -> list:
    """Mon-Sat from the first day in row 2 to 12/31, checked against the typed labels."""
    year = month = None
    for c in range(1, ws.max_column + 1):
        m = re.match(r"([A-Z]+) (\d{4})$", up(ws.cell(1, c).value))
        if m and m.group(1) in MONTHS:
            month, year = MONTHS.index(m.group(1)) + 1, int(m.group(2))
    if not month and default:
        month, year = default
    if not month:
        die(f"{ws.title}: no 'MONTH YYYY' in row 1")
    typed = []
    for c in range(CAL0, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, dt.datetime):
            typed.append(v.day)
        elif v not in (None, ""):
            m = re.search(r"(\d+)$", up(v))
            typed.append(int(m.group(1)) if m else None)
    if not typed or typed[0] is None:
        die(f"{ws.title}: row 2 has no day labels")
    days, d = [], dt.date(year, month, typed[0])
    while d.year == year:
        if d.weekday() < 6:
            days.append(d)
        d += dt.timedelta(days=1)
    for i, t in enumerate(typed):
        if t != days[i].day:
            die(f"{ws.title}: row 2 label #{i + 1} says {t}, the calendar expects {days[i]:%m/%d/%Y}")
    return days


def widths(ws) -> dict:
    """{column index: width}, with range records (min..max) expanded to every column."""
    out = {}
    for k, d in ws.column_dimensions.items():
        lo = d.min or column_index_from_string(k)
        hi = d.max or lo
        for i in range(lo, hi + 1):
            out[i] = d.width
    return out


def extend(ws, days: list, month_style=None) -> int:
    """Grow the calendar to len(days) columns; real dates in every band; month names."""
    old_last = max(c for c in range(CAL0, ws.max_column + 1) if ws.cell(2, c).value not in (None, ""))
    last = CAL0 + len(days) - 1
    w = widths(ws)
    ws.column_dimensions.clear()
    for i, width in w.items():
        if width and i <= old_last:
            ws.column_dimensions[L(i)].width = width
    for c in range(old_last + 1, last + 1):
        ws.column_dimensions[L(c)].width = w.get(old_last) or 17.7
        for r in range(1, ws.max_row + 1):
            src = ws.cell(r, old_last)
            if src.has_style:
                ws.cell(r, c)._style = copy.copy(src._style)
    for c in range(1, ws.max_column + 1):
        if re.match(r"[A-Z]+ \d{4}$", up(ws.cell(1, c).value)):
            month_style = copy.copy(ws.cell(1, c)._style)
            ws.cell(1, c).value = None
    bands = band_rows(ws)
    prev = None
    for i, d in enumerate(days):
        c = CAL0 + i
        for r in bands:
            ws.cell(r, c).value = d
            ws.cell(r, c).number_format = DAY_FMT
        if d.month != prev:
            ws.cell(1, c).value = f"{MONTHS[d.month - 1]} {d.year}"
            if month_style is not None:
                ws.cell(1, c)._style = copy.copy(month_style)
            ws.cell(1, c).alignment = Alignment(horizontal="left")
            prev = d.month
    return last


def freeze(ws) -> None:
    """Freeze below row 3 with A:H fixed, written in Excel's own three-pane form."""
    tl = f"{L(CAL0)}4"
    ws.freeze_panes = tl
    ws.sheet_view.topLeftCell = None
    ws.sheet_view.selection = [
        Selection(pane="topRight", activeCell=f"{L(CAL0)}1", sqref=f"{L(CAL0)}1"),
        Selection(pane="bottomLeft", activeCell="A4", sqref="A4"),
        Selection(pane="bottomRight", activeCell=tl, sqref=tl)]


def add_styles(wb) -> None:
    side = Side(style="thin")
    for name, fill, _old in STYLES:
        st = NamedStyle(name=name)
        st.font = Font(name="Arial", sz=8, b=True, color="FF000000")
        st.fill = PatternFill("solid", fgColor="FF" + fill)
        st.alignment = Alignment(horizontal="center", vertical="center")
        st.border = Border(left=side, right=side, top=side, bottom=side)
        wb.add_named_style(st)


def restyle(ws, last: int) -> dict:
    """Calendar cells in the clerk's four colours -> the matching style. Counts per style."""
    by_old = {old: name for name, _new, old in STYLES}
    bands = set(band_rows(ws))
    n = {name: 0 for name, *_ in STYLES}
    for r in range(3, ws.max_row + 1):
        if r in bands:
            continue
        for c in range(CAL0, last + 1):
            name = by_old.get(fill_key(ws.cell(r, c)) or "")
            if name:
                ws.cell(r, c).style = name
                n[name] += 1
    return n


def legend(wb) -> None:
    lg = wb[LEGEND]
    names = {up(name): name for name, *_ in STYLES}
    hit = 0
    for r in range(1, lg.max_row + 1):
        name = names.get(up(lg.cell(r, 3).value))
        if name:
            lg.cell(r, 2).style = name
            hit += 1
    if hit != len(STYLES):
        die(f"found {hit} of {len(STYLES)} colour rows on the Legend")


def phase_ranges(wb) -> tuple:
    """Legend column B rows of the foundation and flatwork phase lists."""
    lg = wb[LEGEND]
    fnd, flat, sec = [], [], "fnd"
    for r in range(1, lg.max_row + 1):
        if up(lg.cell(r, 2).value).startswith("FLATWORK PHASES"):
            sec = "flat"
        elif isinstance(lg.cell(r, 1).value, (int, float)) and lg.cell(r, 2).value:
            (fnd if sec == "fnd" else flat).append(r)
    if not fnd or not flat:
        die("could not find both phase lists on the Legend")
    return (f"{LEGEND}!$B${min(fnd)}:$B${max(fnd)}", f"{LEGEND}!$B${min(flat)}:$B${max(flat)}")


def dropdowns(wb, last: int) -> None:
    ws = wb[MAIN]
    rows = {}
    for r in range(1, ws.max_row + 1):
        a = up(ws.cell(r, 1).value)
        if a.startswith("FLATWORK") and "flat" not in rows:
            rows["flat"] = r
        elif a.startswith("FOUNDATION") and "fnd" not in rows:
            rows["fnd"] = r
        elif "PUNCH LIST" in a and "punch" not in rows:
            rows["punch"] = r
    if set(rows) != {"flat", "fnd", "punch"}:
        die(f"Main sections not found: {rows}")
    fnd_list, flat_list = phase_ranges(wb)
    for top, bottom, src in ((rows["flat"] + 2, rows["fnd"] - 1, flat_list),
                             (rows["fnd"] + 2, rows["punch"] - 1, fnd_list)):
        dv = DataValidation(type="list", formula1=src, allow_blank=True, showErrorMessage=False)
        dv.add(f"{L(CAL0)}{top}:{L(last)}{bottom}")
        ws.add_data_validation(dv)


def build(src: Path, out: Path, with_dropdowns: bool) -> None:
    if src.resolve() == out.resolve():
        die("OUT must be a new file, never the source")
    wb = load_workbook(src)
    for ws in (wb[MAIN], wb[DONE]):
        if ws.tables:
            die(f"{ws.title} has Excel tables - this build is for the plain schedule")
    days = calendar_days(wb[MAIN])
    if calendar_days(wb[DONE], (days[0].month, days[0].year))[0] != days[0]:
        die("Main and Completed start on different days")
    add_styles(wb)
    month_style = next((copy.copy(c._style) for c in wb[MAIN][1]
                        if re.match(r"[A-Z]+ \d{4}$", up(c.value))), None)
    report = {}
    for ws in (wb[MAIN], wb[DONE]):
        last = extend(ws, days, month_style)
        report[ws.title] = restyle(ws, last)
        freeze(ws)
    legend(wb)
    if with_dropdowns:
        dropdowns(wb, last)
    wb.save(out)
    assert_clean(out)
    print(f"OK  {out}")
    print(f"    {days[0]:%m/%d/%Y} - {days[-1]:%m/%d/%Y}: {len(days)} days, columns {L(CAL0)}:{L(last)}")
    for sheet, n in report.items():
        print(f"    {sheet}: " + ", ".join(f"{k} {v}" for k, v in n.items()))


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(a) != 2:
        die("usage: schedule_simple_build.py SRC.xlsx OUT.xlsx [--dropdowns]")
    build(Path(a[0]), Path(a[1]), "--dropdowns" in sys.argv)
