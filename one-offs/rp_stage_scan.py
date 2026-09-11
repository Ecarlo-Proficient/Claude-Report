#!/usr/bin/env python3
"""
rp_stage_scan.py - which residential jobs bill ONCE and which bill BY SCOPE.

The owner (2026-09-08): "we have two different types, a one invoice one job vs
a scope based invoices for one job like this one" (RP6586). The P&L treats
them differently - a one-invoice job's invoice date is the end of the job
and later bills are suspect; a scope-based job bills each scope as it
finishes (lot prep, piers, grade beams, foundation), so it carries several
invoices months apart and nothing is cut off. This scan reads the evidence
off QBO's own invoices and lists every RP job with its kind.

THE RULE lives in shared/rp_invoicing (project-pnl decides the same way):
every memo reads "RP#### - <address> - <what was billed>", so the third
segment is the invoice's description, not a kind marker. A STAGE invoice
names a piece of the slab build (lot prep, piers, mud slab, grade beams,
walls, foundation) and is not an extra (pump, rock saw, brickledge, bolts,
"... Extra"). Two or more stage invoices = scope-based; one partial stage
(piers, grade beams ...) = scope-based with one stage billed so far; one
"Foundation" invoice, with or without extras, = a one-invoice job. TYPE on
the RP WIP (Tract / Custom) is carried for corroboration, never as the rule.

Output: <CompanyHealth>/RP Invoicing Stages.xlsx  (Jobs + Invoices sheets,
plain) and a summary on stdout. Read-only against QBO and the WIP master.

  python3 one-offs/rp_stage_scan.py                 # invoices since 2025-01-01
  python3 one-offs/rp_stage_scan.py --since 2024-01-01 --out ~/Downloads/x.xlsx
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import defaultdict
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
from shared import paths, wip_contracts, xlsx_verify           # noqa: E402
from shared import rp_invoicing                                  # noqa: E402
from shared.qbo_api import load_credentials, project_of_invoice, query_all  # noqa: E402

_JOB_RE = re.compile(r"^RP\d{4}(?:-FTW)?$", re.I)


def _memo(inv: dict) -> str:
    return ((inv.get("PrivateNote") or "").strip()
            or ((inv.get("CustomerMemo") or {}).get("value") or "").strip())


def _date(s: str):
    try:
        return dt.date.fromisoformat(str(s)[:10])
    except ValueError:
        return None


def _rp_wip_rows() -> dict:
    """{job -> {type, builder, status, contract}} off the WIP master's Test - RP."""
    out: dict = {}
    p = wip_contracts._wip_dir() / "WIP - MASTER new.xlsx"
    if not p.exists():
        return out
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True, data_only=True)
    if "Test - RP" not in wb.sheetnames:
        return out
    ws = wb["Test - RP"]
    rows = ws.iter_rows(values_only=True)
    hdr = None
    for row in rows:
        if any(isinstance(v, str) and v.strip().upper() == "PROJECT #" for v in row):
            hdr = {str(v).strip().upper(): i for i, v in enumerate(row) if v is not None}
            break
    if not hdr:
        return out
    for row in rows:
        j = hdr.get("PROJECT #")
        job = str(row[j] or "").strip().upper() if j is not None and j < len(row) else ""
        if not _JOB_RE.match(job):
            continue

        def g(h):
            i = hdr.get(h)
            return row[i] if i is not None and i < len(row) else None
        out[job] = {"type": g("TYPE"), "builder": g("BUILDER"), "status": g("STATUS"),
                    "contract": g("ORIGINAL CONTRACT")}
    wb.close()
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--since", default="2025-01-01", help="invoices dated on/after (YYYY-MM-DD)")
    ap.add_argument("--out", type=Path,
                    default=paths.companyhealth_dir() / "RP Invoicing Stages.xlsx")
    a = ap.parse_args()

    access, cid = load_credentials()
    invs = query_all(access, cid, "Invoice", where=f"TxnDate >= '{a.since}'")
    by_job: dict = defaultdict(list)
    for inv in invs:
        job = (project_of_invoice(inv) or "").upper()
        if not _JOB_RE.match(job):
            continue
        by_job[job].append(inv)
    wip = _rp_wip_rows()
    print(f"{len(invs)} invoices since {a.since} · {len(by_job)} RP jobs invoiced · "
          f"{len(wip)} RP rows on Test - RP")

    jobs_rows, inv_rows = [], []
    for job in sorted(set(by_job) | set(wip)):
        prof = rp_invoicing.classify([
            {"doc_num": i.get("DocNumber"), "date": i.get("TxnDate", ""),
             "memo": _memo(i), "amount": float(i.get("TotalAmt") or 0)}
            for i in by_job.get(job, [])])
        lst, scopes, flags = prof["invoices"], prof["scopes"], prof["stage_flags"]
        dates = [_date(i["date"]) for i in lst]
        kind = prof["kind"]
        w = wip.get(job, {})
        jobs_rows.append([job, w.get("builder"), w.get("type"), w.get("status"), kind,
                          len(lst), prof["n_stage"], dates[0] if dates else None,
                          dates[-1] if dates else None, prof["span_days"],
                          sum(float(i["amount"] or 0) for i in lst),
                          w.get("contract"),
                          " | ".join(s for s, f in zip(scopes, flags) if f),
                          " | ".join(s for s, f in zip(scopes, flags) if s and not f)])
        for i, sc, f in zip(lst, scopes, flags):
            inv_rows.append([job, i["doc_num"], _date(i["date"]), i["amount"], sc,
                             "stage" if f else ("extra" if sc else ""), i["memo"]])

    kinds = defaultdict(int)
    by_type = defaultdict(lambda: defaultdict(int))
    for row in jobs_rows:
        kinds[row[4]] += 1
        by_type[str(row[2] or "?")][row[4]] += 1
    print("\nBy kind:")
    for k, n in sorted(kinds.items(), key=lambda kv: -kv[1]):
        print(f"  {n:>4}  {k}")
    print("\nBy WIP TYPE (Tract / Custom) x kind:")
    for t, d in sorted(by_type.items()):
        print(f"  {t}: " + ", ".join(f"{k} {n}" for k, n in sorted(d.items(), key=lambda kv: -kv[1])))

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active; ws.title = "Jobs"
    hdr = ["Project #", "Builder", "Type", "WIP status", "Invoicing kind", "Invoices",
           "Stage invoices", "First invoice", "Last invoice", "Span (days)", "Billed", "Contract",
           "Stage scopes", "Extras / other"]
    ws.append(hdr)
    for row in jobs_rows:
        ws.append(row)
    ws2 = wb.create_sheet("Invoices")
    ws2.append(["Project #", "Invoice #", "Date", "Amount", "Scope (memo)", "Stage / extra", "Memo"])
    for row in inv_rows:
        ws2.append(row)
    for w in (ws, ws2):
        for c in w[1]:
            c.font = Font(bold=True)
        w.freeze_panes = "A2"
        for col in w.iter_cols(min_row=2):
            for c in col:
                if isinstance(c.value, dt.date):
                    c.number_format = "mm/dd/yyyy"
                elif isinstance(c.value, float):
                    c.number_format = '"$"#,##0_);[Red]("$"#,##0)'
    for w, widths in ((ws, (10, 28, 8, 10, 32, 9, 9, 13, 13, 11, 13, 13, 50, 50)),
                      (ws2, (10, 11, 12, 13, 34, 12, 70))):
        for i, wd in enumerate(widths, start=1):
            w.column_dimensions[w.cell(1, i).column_letter].width = wd
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    xlsx_verify.assert_clean(a.out)
    print(f"\nwrote {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
