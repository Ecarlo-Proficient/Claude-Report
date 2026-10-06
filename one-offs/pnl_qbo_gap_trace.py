#!/usr/bin/env python3
"""
pnl_qbo_gap_trace.py - trace a project P&L's Reconciliations gap to the
documents behind it (read-only on QBO and on the workbook).

The Reconciliations sheet compares QuickBooks' own P&L for the job (through the
last draw's end) with the Transactions sheet through the same date, per line:
Income / Cost of Goods Sold / Operating Expenses. When a line reads "off by",
this pulls QuickBooks' ProfitAndLossDetail for the job's customer through that
date, sums both sides PER DOCUMENT (invoice / bill / check / journal entry #),
and prints every document whose two sides differ - one row per document with
date, type, amount on each side and the QBO link.

USAGE
  python3 one-offs/pnl_qbo_gap_trace.py "<Project_PnL_XXX.xlsx>" [more workbooks]
"""
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "project-pnl"))
from openpyxl import load_workbook

import completed_pnl as cp
from shared import qbo_api

SECTIONS = {"income": "Income", "cogs": "Cost of Goods Sold", "exp": "Expenses"}
LINK = {"Invoice": "invoice", "Bill": "bill", "Check": "check", "Expense": "expense",
        "Journal Entry": "journal", "Credit Card Credit": "creditcardcredit",
        "Vendor Credit": "vendorcredit", "Deposit": "deposit", "Credit Memo": "creditmemo"}


def _date(v):
    return v.date() if hasattr(v, "date") else v


def workbook_side(path):
    """Per-section {doc: amount} from the Transactions sheet, plus the
    QuickBooks-through date the Reconciliations sheet compares on."""
    wb = load_workbook(str(path))
    rec = wb["Reconciliations"]
    thru = None
    for row in rec.iter_rows(min_row=1, max_row=12):
        for c in row:
            m = re.search(r"through\s+(\d{2})/(\d{2})/(\d{4})", str(c.value or ""))
            if m:
                thru = f"{m[3]}-{m[1]}-{m[2]}"
                break
        if thru:
            break
    ws = wb["Transactions"]
    c0 = next(c for c in (1, 2) for r in range(1, 60)
              if str(ws.cell(r, c).value or "").strip() == "Inv #")
    inv, not_billed = cp._read_invoices(ws, c0)
    side = {k: defaultdict(float) for k in SECTIONS}
    lim = qbo_api_date(thru)
    for i in inv:
        d = _date(i["date"])
        if lim and d and str(d) > lim:
            continue
        side["income"][str(i["doc"] or "").strip()] += i["gross"]
    if not_billed:
        side["income"]["(retainage not billed / JE)"] += not_billed
    for sec in cp._read_costs(ws, c0):
        key = "cogs" if sec["name"].startswith("COST") else "exp"
        for a in sec["accounts"]:
            for v in a["vendors"]:
                for ln in v["lines"]:
                    d = _date(ln["date"])
                    if lim and d and str(d) > lim:
                        continue
                    side[key][str(ln["doc"] or "").strip()] += ln["amt"]
    return thru, side


def qbo_api_date(s):
    return s


def qbo_side(access, cid, customer_id, thru):
    rep = qbo_api.report(access, cid, "ProfitAndLossDetail", params={
        "start_date": "2015-01-01", "end_date": thru, "accounting_method": "Accrual",
        "customer": customer_id})
    cols = [c.get("ColTitle") or c.get("ColType") for c in rep.get("Columns", {}).get("Column", [])]
    side = {k: defaultdict(float) for k in SECTIONS}
    info = {}

    def walk(node, top):
        for r in (node.get("Rows") or {}).get("Row", []) or []:
            if r.get("type") == "Section":
                h = ((r.get("Header") or {}).get("ColData") or [{}])[0].get("value", "")
                # the section that decides Income / COGS / Expenses is the
                # first header that names one ("Ordinary Income/Expenses" wraps them)
                walk(r, top if top else (h if h.strip().lower() in
                                         {v.lower() for v in SECTIONS.values()} else None))
            elif r.get("type") == "Data":
                cd = r.get("ColData", [])
                rec = {cols[i]: (cd[i].get("value"), cd[i].get("id")) for i in range(min(len(cols), len(cd)))}
                key = next((k for k, s in SECTIONS.items()
                            if (top or "").strip().lower() == s.lower()), None)
                if key is None:
                    continue
                amt = float(rec.get("Amount", ("0",))[0] or 0)
                doc = str(rec.get("Num", ("",))[0] or "").strip()
                typ = rec.get("Transaction Type", ("", None))
                side[key][doc] += amt
                info.setdefault((key, doc), (rec.get("Date", ("",))[0], typ[0], typ[1],
                                             rec.get("Name", ("",))[0]))
    walk(rep, None)
    return side, info


def main(paths):
    access, cid = qbo_api.load_credentials()
    pmap = qbo_api.build_project_customer_map(access, cid)
    for p in paths:
        job = re.search(r"((?:CP|MFD|RP)\d+(?:-FTW)?)", Path(p).name).group(1)
        thru, wbs = workbook_side(p)
        qs, info = qbo_side(access, cid, pmap[job]["id"], thru)
        print(f"\n=== {job}  (QuickBooks through {thru})")
        for k, label in SECTIONS.items():
            tq, tw = sum(qs[k].values()), sum(wbs[k].values())
            if abs(tq - tw) < 0.01:
                continue
            print(f"  {label}: QuickBooks {tq:,.2f}  Transactions {tw:,.2f}  gap {tw - tq:,.2f}")
            docs = set(qs[k]) | set(wbs[k])
            rows = [(d, qs[k].get(d, 0.0), wbs[k].get(d, 0.0)) for d in docs]
            rows = [x for x in rows if abs(x[1] - x[2]) >= 0.01]
            for d, q, w in sorted(rows, key=lambda x: -abs(x[2] - x[1])):
                dt_, typ, tid, name = info.get((k, d), ("", "", None, ""))
                url = (f"https://qbo.intuit.com/app/{LINK.get(typ, 'transaction')}?txnId={tid}"
                       if tid else "")
                print(f"    #{d or '(no #)':<26} {str(dt_):<11} {str(typ):<14} QBO {q:>12,.2f}  "
                      f"Tx {w:>12,.2f}  diff {w - q:>11,.2f}  {str(name)[:28]:<28} {url}")


if __name__ == "__main__":
    main(sys.argv[1:])
