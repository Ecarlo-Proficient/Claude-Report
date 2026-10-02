#!/usr/bin/env python3
"""cp_overview_verify.py - independent check of CP Overview.xlsx (read-only).

Every job's BILLED and COSTS re-derived straight from QBO and compared to what
the workbook's own formulas evaluate to:
  billed = invoice income lines; a positive line on the RETAINAGE item is a
           release (collects retainage already billed) and is left out
  costs  = bill / expense LINE amounts on the job's project (shared/qbo_costs),
           payroll left out (it rides the overhead %)
A gap is printed per job; a known one (a retainage double count the Overview
removes on purpose) is explained by its ⚑ line from the build.

USAGE
  python3 one-offs/cp_overview_verify.py                       # the active CP Overview
  python3 one-offs/cp_overview_verify.py "<CP Completed 2026.xlsx>" --legacy
"""
import re
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook

from shared import qbo_api, qbo_costs as qc
from shared.job_lines import JobMatcher

LEGACY = "--legacy" in sys.argv          # the COMPLETED book: pre-2026 jobs, P&Ls built --legacy
_args = [a for a in sys.argv[1:] if a != "--legacy"]
OV = (_args[0] if _args else
      "/Volumes/Common/CURRENT PROJECTS/Awarded Projects Commercial projects/CP Overview.xlsx")
PAY = re.compile(r"payroll|wages|salar|employee benefit|workers.?comp", re.I)

wb = load_workbook(OV)


def ev(ws, ref, depth=0):
    v = ws[ref].value
    if isinstance(v, (int, float)) or v is None:
        return v or 0
    s = str(v)
    if not s.startswith("="):
        return 0
    f = s[1:]
    m = re.fullmatch(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", f)
    if m:
        return sum(ev(ws, f"{m[1]}{i}") for i in range(int(m[2]), int(m[4]) + 1))
    if re.fullmatch(r"[A-Z]+\d+(\+[A-Z]+\d+)*", f):
        return sum(ev(ws, x) for x in f.split("+"))
    m = re.fullmatch(r"'?([^'!]+)'?!([A-Z]+\d+)", f)
    if m:
        return ev(wb[m[1]], m[2])
    raise ValueError(f"{ws.title}!{ref}: {s}")


ov = wb["Overview"]
hdr = {ov.cell(5, c).value: c for c in range(1, 40) if ov.cell(5, c).value}
jobs = {}
for r in range(6, ov.max_row + 1):
    lab = str(ov.cell(r, 2).value or "")
    m = re.match(r"(CP\d+)", lab)
    if m:
        jobs[m[1]] = (ev(ov, f"{ov.cell(r, hdr['Billed']).coordinate}"),
                      ev(ov, f"{ov.cell(r, hdr['Costs']).coordinate}"))

access, cid = qbo_api.load_credentials()
pmap = qbo_api.build_project_customer_map(access, cid)
acct = qc.build_account_map(access, cid)
# EVERY project customer, so a line coded to another job stays with that job;
# a line counts for a job only when it is CODED to the job's project (the
# owner's rule - the project is the source of truth). A line whose memo names
# the job but that is coded elsewhere / not coded is reported, never counted.
c2p = {v["id"]: p for p, v in pmap.items()}
p2name = {v["id"]: p for p, v in pmap.items()}
cost = defaultdict(float)
loose = defaultdict(list)
for ln in qc.iter_cost_lines(access, cid, acct, c2p, since="2020-01-01"):
    if PAY.search(str(ln["account"] or "")):
        continue
    coded = c2p.get(ln["customer_id"]) if ln["customer_id"] else None
    if coded in jobs:
        cost[coded] += ln["amount"]
        continue
    named = re.findall(r"\bCP\s?-?(\d{3,4})\b", str(ln["memo"] or "") + " " + str(ln["description"] or ""))
    for n in {f"CP{x}" for x in named} & set(jobs):
        loose[n].append((ln, coded))

matchers = {}
if LEGACY:
    # the same rule the --legacy P&L uses: project, OR the line text names the
    # job, OR the bill memo names it and only it. The job's CLASS is NOT used -
    # the P&L adds class lines only with +class, and they are mixed (CP697
    # #11012 sits on its class with a memo naming RP6906)
    bills, purchases = qc.pull_expense_txns(access, cid, "2020-01-01")
    for j in jobs:
        if j not in pmap:
            continue
        m = JobMatcher(pmap[j]["id"], j, [], legacy=True, text_rules=True)
        matchers[j] = m
        cost[j] = 0.0
        for t in bills + purchases:
            for ln in t.get("Line") or []:
                det = ln.get("AccountBasedExpenseLineDetail") or ln.get("ItemBasedExpenseLineDetail")
                if not det or not m(det, ln, t):
                    continue
                if PAY.search(qc.cost_leaf(det, acct, fallback="")):
                    continue
                cost[j] += float(ln.get("Amount") or 0)

print(f"{'JOB':<7}{'QBO billed':>14}{'workbook':>14}{'diff':>11}  |{'QBO cost':>14}{'workbook':>14}{'diff':>11}")
bad = 0
for j, (wb_b, wb_c) in sorted(jobs.items()):
    info = pmap.get(j)
    gross = 0.0
    invs = qbo_api.query_all(access, cid, "Invoice", f"CustomerRef = '{info['id']}'")
    if LEGACY and info.get("parent_id"):          # older jobs invoiced on the GC
        have = {i["Id"] for i in invs}
        invs += [i for i in qbo_api.query_all(access, cid, "Invoice",
                                               f"CustomerRef = '{info['parent_id']}'")
                 if i["Id"] not in have and matchers[j].invoice_belongs(i)]
    for inv in invs:
        for ln in inv.get("Line", []):
            if ln.get("DetailType") != "SalesItemLineDetail":
                continue
            item = ((ln.get("SalesItemLineDetail") or {}).get("ItemRef") or {}).get("name", "")
            amt = float(ln.get("Amount") or 0)
            if amt > 0 and not re.search(r"retainage", item, re.I):
                gross += amt
    db, dc = wb_b - gross, wb_c - cost[j]
    flag = "" if abs(db) < 1 and abs(dc) < 1 else "  <-- check"
    bad += bool(flag)
    print(f"{j:<7}{gross:>14,.2f}{wb_b:>14,.2f}{db:>11,.2f}  |{cost[j]:>14,.2f}{wb_c:>14,.2f}{dc:>11,.2f}{flag}")
print(f"\n{len(jobs)} jobs, {bad} to explain")

print("\nCOST LINES THAT NAME A JOB BUT ARE NOT CODED TO IT (not counted - fix in QBO if they are the job's):")
for j in sorted(loose):
    for ln, coded in sorted(loose[j], key=lambda x: -abs(x[0]["amount"])):
        if abs(ln["amount"]) < 100:
            continue
        print(f"  {j:<7}{ln['txn_type']:<8}#{str(ln['doc_number']):<14}{str(ln['txn_date']):<11}"
              f"{str(ln['vendor'])[:28]:<29}{ln['amount']:>11,.2f}  "
              f"{'coded to ' + coded if coded else 'NO project on the line'}")
