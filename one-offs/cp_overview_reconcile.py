#!/usr/bin/env python3
"""
cp_overview_reconcile.py - prove a CP Overview / CP Completed workbook adds up,
with no Excel and no QuickBooks (read-only, the workbook only).

The companion of cp_overview_verify.py: that one ties each job's billed and
costs to QBO; this one proves the workbook is right INSIDE (the owner 10/02/2026:
"verify the overviews reconcile and are correct"):

  1. every formula evaluates (a small built-in evaluator for the functions the
     workbooks use: IF / OR / AND / SUM / SUMPRODUCT / ISNUMBER, refs, ranges)
  2. each job sheet: coverage TOTAL = its rows = its draw blocks = its
     transaction lines; gross / 10% OH / net / coverage recomputed per row;
     every draw link lands on its block
  3. the Overview: billed / costs = the job sheet's TOTAL; every derived figure
     (gross, OH, net, %s, coverage, projection, % billed / complete, vs plan)
     recomputed in Python; money totals = the sum of the job rows; job links work

Exit 1 when anything is off, with one line per problem.

USAGE
  python3 one-offs/cp_overview_reconcile.py "<CP Overview.xlsx>"
  python3 one-offs/cp_overview_reconcile.py "<CP Completed 2026.xlsx>"
"""
import re
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

OH = 0.10


TOK = re.compile(r"""\s*(?:
    (?P<num>\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:'[^']+'|[A-Za-z_][\w\.]*)?!?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<fn>[A-Z]+)\(
  | (?P<op><>|<=|>=|[-+*/&=<>(),:])
)""", re.X)


class Err(Exception):
    pass


class Book:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}

    def value(self, sheet, ref):
        key = (sheet, ref)
        if key in self.cache:
            return self.cache[key]
        v = self.wb[sheet][ref].value
        if isinstance(v, str) and v.startswith("="):
            self.cache[key] = None                      # cycle guard
            v = self.formula(sheet, v[1:])
        elif v is None:
            v = 0.0 if False else None
        self.cache[key] = v
        return v

    # ── parse + evaluate ──
    def formula(self, sheet, f):
        toks = []
        pos = 0
        while pos < len(f):
            m = TOK.match(f, pos)
            if not m or m.end() == pos:
                raise Err(f"cannot tokenise {f[pos:]!r}")
            pos = m.end()
            kind = m.lastgroup
            toks.append((kind, m.group(kind)))
        self.toks, self.i, self.sheet = toks, 0, sheet
        node = self.expr()
        if self.i != len(toks):
            raise Err(f"trailing tokens in {f}")
        return node()

    def peek(self):
        return self.toks[self.i] if self.i < len(self.toks) else (None, None)

    def take(self, val=None):
        t = self.peek()
        if val is not None and t[1] != val:
            raise Err(f"expected {val}, got {t}")
        self.i += 1
        return t

    def expr(self):                                   # comparison
        left = self.concat()
        while self.peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.take()[1]
            right = self.concat()
            left = (lambda a, b, op: lambda: cmp(a(), b(), op))(left, right, op)
        return left

    def concat(self):
        left = self.add()
        while self.peek()[1] == "&":
            self.take()
            right = self.add()
            left = (lambda a, b: lambda: f"{a()}{b()}")(left, right)
        return left

    def add(self):
        left = self.mul()
        while self.peek()[1] in ("+", "-"):
            op = self.take()[1]
            right = self.mul()
            left = (lambda a, b, op: lambda: arith(a(), b(), op))(left, right, op)
        return left

    def mul(self):
        left = self.unary()
        while self.peek()[1] in ("*", "/"):
            op = self.take()[1]
            right = self.unary()
            left = (lambda a, b, op: lambda: arith(a(), b(), op))(left, right, op)
        return left

    def unary(self):
        if self.peek()[1] == "-":
            self.take()
            inner = self.unary()
            return lambda: arith(0.0, inner(), "-")
        if self.peek()[1] == "+":
            self.take()
            return self.unary()
        return self.atom()

    def atom(self):
        kind, val = self.take()
        if kind == "num":
            x = float(val)
            return lambda: x
        if kind == "str":
            s = val[1:-1].replace('""', '"')
            return lambda: s
        if kind == "ref":
            sheet = self.sheet
            if "!" in val:
                sheet, val = val.rsplit("!", 1)
                sheet = sheet.strip("'")
            val = val.replace("$", "")
            if ":" in val:
                a, b = val.split(":")
                cells = rng(a, b)
                return lambda: [self.value(sheet, c) for c in cells]
            return lambda: self.value(sheet, val)
        if kind == "op" and val == "(":
            e = self.expr()
            self.take(")")
            return e
        if kind == "fn":
            args = []
            if self.peek()[1] != ")":
                args.append(self.expr())
                while self.peek()[1] == ",":
                    self.take()
                    args.append(self.expr())
            self.take(")")
            return fn(val, args)
        raise Err(f"unexpected {kind} {val}")


def rng(a, b):
    ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups()
    cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
    out = []
    for r in range(int(ra), int(rb) + 1):
        for c in range(column_index_from_string(ca), column_index_from_string(cb) + 1):
            out.append(f"{get_column_letter(c)}{r}")
    return out


def num(x):
    if x is None or x == "":
        return 0.0
    if isinstance(x, bool):
        return 1.0 if x else 0.0
    if isinstance(x, (int, float)):
        return float(x)
    raise Err(f"#VALUE! {x!r}")


def arith(a, b, op):
    if isinstance(a, list) or isinstance(b, list):
        la = a if isinstance(a, list) else [a] * len(b)
        lb = b if isinstance(b, list) else [b] * len(a)
        return [arith(x, y, op) for x, y in zip(la, lb)]
    if a == "" or b == "":
        raise Err("#VALUE! arithmetic on text")
    x, y = num(a), num(b)
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if y == 0:
        raise Err("#DIV/0!")
    return x / y


def cmp(a, b, op):
    if isinstance(a, list) or isinstance(b, list):
        la = a if isinstance(a, list) else [a] * len(b)
        lb = b if isinstance(b, list) else [b] * len(a)
        return [cmp(x, y, op) for x, y in zip(la, lb)]
    if a is None:
        a = "" if isinstance(b, str) else 0.0
    if b is None:
        b = "" if isinstance(a, str) else 0.0
    if isinstance(a, str) != isinstance(b, str):
        r = {"=": False, "<>": True}.get(op)
        if r is None:
            raise Err("mixed compare")
        return r
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


def fn(name, args):
    if name == "IF":
        def f():
            c = args[0]()
            return args[1]() if (c and c != "") else (args[2]() if len(args) > 2 else False)
        return f
    if name in ("OR", "AND"):
        def f():
            vals = [bool(a()) for a in args]
            return any(vals) if name == "OR" else all(vals)
        return f
    if name == "ISNUMBER":
        return lambda: isinstance(args[0](), (int, float)) and not isinstance(args[0](), bool)
    if name == "SUM":
        def f():
            t = 0.0
            for a in args:
                v = a()
                for x in (v if isinstance(v, list) else [v]):
                    if isinstance(x, (int, float)) and not isinstance(x, bool):
                        t += x
            return t
        return f
    if name == "SUMPRODUCT":
        def f():
            v = args[0]()
            if len(args) > 1:
                for a in args[1:]:
                    v = arith(v, a(), "*")
            return sum(num(x) for x in v)
        return f
    raise Err(f"unknown function {name}")


def main(path: str) -> int:
    wb = load_workbook(path)
    bk = Book(wb)
    problems = []
    def bad(msg):
        problems.append(msg)


    def close(a, b, tol=0.01):
        if a in (None, "") and b in (None, ""):
            return True
        if a in (None, "") or b in (None, ""):
            return False
        return abs(float(a) - float(b)) <= tol


    # 1. every formula evaluates
    n_f = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_f += 1
                    try:
                        bk.value(ws.title, c.coordinate)
                    except (Err, RecursionError) as e:
                        bad(f"{ws.title}!{c.coordinate} does not evaluate: {e}  [{c.value[:60]}]")
    print(f"formulas evaluated: {n_f}")

    # 2. job sheets: draw rows = blocks = lines; TOTAL = rows
    sheet_tot = {}
    links_ok = [0]
    for ws in wb.worksheets[1:]:
        t = ws.title
        v = lambda ref: bk.value(t, ref)  # noqa: E731
        col_b = {r: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
        tot_r = next(r for r, x in col_b.items() if x == "TOTAL")
        hdr_r = next(r for r, x in col_b.items() if x == "Draw")
        tx_r = next(r for r, x in col_b.items() if x == "Date" and r > tot_r)
        table = list(range(hdr_r + 1, tot_r))
        sF = sum(float(v(f"F{r}") or 0) for r in table)
        sG = sum(float(v(f"G{r}") or 0) for r in table)
        if not close(sF, v(f"F{tot_r}")) or not close(sG, v(f"G{tot_r}")):
            bad(f"{t}: coverage TOTAL != sum of its rows")
        # transaction lines: rows with a date or a doc in C but not a band / vendor row
        lineF = lineG = 0.0
        bands = {}
        for r in range(tx_r + 1, ws.max_row + 1):
            b, d = ws.cell(r, 2).value, ws.cell(r, 4).value
            fill = ws.cell(r, 2).fill.fgColor.rgb if ws.cell(r, 2).fill.fill_type else None
            is_band = fill and str(fill).endswith("E6ECF3")
            is_vendor = (b is None and isinstance(d, str) and re.search(r"\(\d+\)$", d or "")
                         and ws.cell(r, 3).value is None)
            if is_band:
                bands[r] = (v(f"F{r}"), v(f"G{r}"))
                continue
            if is_vendor:
                continue
            lineF += float(v(f"F{r}") or 0)
            lineG += float(v(f"G{r}") or 0)
        if not close(lineF, v(f"F{tot_r}")) or not close(lineG, v(f"G{tot_r}")):
            bad(f"{t}: transaction lines {lineF:,.2f} / {lineG:,.2f} != TOTAL "
                f"{v(f'F{tot_r}'):,.2f} / {v(f'G{tot_r}'):,.2f}")
        bF = sum(float(x[0] or 0) for x in bands.values())
        bG = sum(float(x[1] or 0) for x in bands.values())
        if not close(bF, v(f"F{tot_r}")) or not close(bG, v(f"G{tot_r}")):
            bad(f"{t}: draw blocks {bF:,.2f} / {bG:,.2f} != TOTAL")
        # each table row: profit / OH / net / coverage recomputed
        for r in table + [tot_r]:
            F, G = float(v(f"F{r}") or 0), float(v(f"G{r}") or 0)
            if not close(v(f"H{r}"), F - G) or not close(v(f"I{r}"), F * OH) \
                    or not close(v(f"J{r}"), F - G - F * OH):
                bad(f"{t}!row {r}: gross / OH / net wrong")
            cov = v(f"K{r}")
            if ws.cell(r, 2).value in ("Next draw", "Outside draws"):
                continue                     # no coverage on these two rows, by design
            if G and not close(cov, F / G, 1e-9):
                bad(f"{t}!K{r}: coverage {cov} != {F / G}")
        # draw-name links land on their band
        for r in table:
            h = ws.cell(r, 2).hyperlink
            loc = h and (h.location or (h.target or "").lstrip("#"))
            if loc:
                m = re.match(r"'?([^'!]+)'?!B(\d+)", loc)
                if not m or m[1] != t or int(m[2]) not in bands:
                    bad(f"{t}!B{r}: link {loc} does not land on a draw block")
                links_ok[0] += 1
        sheet_tot[t] = (v(f"F{tot_r}"), v(f"G{tot_r}"))

    # 3. Overview rows: base = job sheet; every derived figure recomputed
    ov = wb["Overview"]
    hdr = {ov.cell(5, c).value: c for c in range(1, 40) if ov.cell(5, c).value}
    H = lambda name, r: bk.value("Overview", ov.cell(r, hdr[name]).coordinate) if name in hdr else None  # noqa: E731
    job_rows, tot_rows = [], []
    for r in range(6, ov.max_row + 1):
        lab = str(ov.cell(r, 2).value or "")
        if re.match(r"CP\d+", lab):
            job_rows.append(r)
        elif lab.startswith(("Subtotal", "ALL CP")):
            tot_rows.append(r)
    for r in job_rows:
        job = re.match(r"(CP\d+)", ov.cell(r, 2).value)[1]
        b, c = float(H("Billed", r)), float(H("Costs", r))
        if not close(b, sheet_tot[job][0]) or not close(c, sheet_tot[job][1]):
            bad(f"Overview {job}: billed/costs != its sheet TOTAL")
        exp = {"Gross profit": b - c, "Gross %": (b - c) / b if b else "",
               "10% OH (of billed)": b * OH, "Net profit (10% OH)": b - c - b * OH,
               "Net %": (b - c - b * OH) / b if b else "", "Coverage": b / c if c else ""}
        ctr, etc = H("Contract", r), H("Est. cost (ETC)", r)
        if "Contract" in hdr:
            ctr, etc = float(ctr or 0), float(etc or 0)
            proj = ctr and etc
            pp = ctr - etc if proj else ""
            pn = pp - ctr * OH if proj else ""
            exp.update({"Proj. gross profit": pp, "Proj. gross %": pp / ctr if proj else "",
                        "10% OH (of contract)": ctr * OH if proj else "",
                        "Proj. net (10% OH)": pn, "Proj. net %": pn / ctr if proj else "",
                        "% billed": b / ctr if ctr else "", "% complete": c / etc if etc else ""})
            nm, pm = exp["Net %"], exp["Proj. net %"]
            exp["Net vs plan"] = nm - pm if nm != "" and pm != "" else ""
        for name, want in exp.items():
            got = H(name, r)
            if not close(got, want, 1e-6 if isinstance(want, float) and abs(want) < 5 else 0.01):
                bad(f"Overview {job} {name}: {got} != {want}")
    # totals: money columns = sum of job rows
    money = [k for k in hdr if k in ("Contract", "Est. cost (ETC)", "Billed", "Costs", "Gross profit",
                                     "Proj. gross profit", "10% OH (of contract)", "Proj. net (10% OH)",
                                     "10% OH (of billed)", "Net profit (10% OH)", "Awaiting draw", "Still owed")]
    last_tot = tot_rows[-1]
    for k in money:
        s = sum(float(H(k, r) or 0) for r in job_rows)
        if not close(H(k, last_tot), s):
            bad(f"Overview total {k}: {H(k, last_tot)} != sum of rows {s}")
    b, c = float(H("Billed", last_tot)), float(H("Costs", last_tot))
    if not close(H("Coverage", last_tot), b / c, 1e-9) or not close(H("Gross %", last_tot), (b - c) / b, 1e-9):
        bad("Overview total: coverage / gross % not from the totals")
    # job-name links
    for r in job_rows:
        h = ov.cell(r, 2).hyperlink
        loc = h and (h.location or (h.target or "").lstrip("#"))
        if not loc or loc.split("!")[0].strip("'") not in wb.sheetnames:
            bad(f"Overview row {r}: job link broken")
        else:
            links_ok[0] += 1

    print(f"jobs: {len(job_rows)}   job sheets: {len(sheet_tot)}")
    print(f"links checked: {links_ok[0]}")
    print(f"problems: {len(problems)}")
    for p in problems[:40]:
        print("  ✗", p)
    return 1 if problems else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        raise SystemExit(2)
    raise SystemExit(main(sys.argv[1]))
