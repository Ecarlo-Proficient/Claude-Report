"""schedule_v3_build.py - rebuild the horizontal RP schedule workbook (v2 -> v3).

The owner's 10/2026 schedule (one row per job, the stage typed on a Mon-Sat
calendar) gets, on a COPY, never the original:

  1. Two Excel tables on 'Main Schedule' - Flatwork and Foundation - each with
     a # column (the phase's place in the legend; Data -> Reapply sorts by it)
     and its own PHASE dropdown fed by the legend. (A CHECK column that flagged
     rows in the wrong table was tried and dropped: the owner found it noise.)
  2. The legend phases become two tables (Phases_Foundation, Phases_Flatwork);
     # renumbers itself from position, the dropdowns read them live. The four
     legend colours become workbook cell styles in the calendar's font.
  3. The calendar runs Mon-Sat to 12/31 with real dates, a month name over the
     first working day of each month, punch-list day bands included.
  4. 'Completed' becomes an empty copy of Main (same columns, same dates, both
     tables with one blank row) so a finished row moves with its history; A1
     on each sheet turns red if the two date bands ever disagree.

  python3 one-offs/schedule_v3_build.py SRC.xlsx OUT.xlsx [--fake]

--fake swaps every job's text for made-up values (for testing in Excel).
Read-only on SRC. OUT must not be SRC. Ends with xlsx_verify.assert_clean.
"""
from __future__ import annotations

import copy
import datetime as dt
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side  # noqa: E402
from openpyxl.styles.colors import Color  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.filters import AutoFilter, SortCondition, SortState  # noqa: E402
from openpyxl.worksheet.table import Table, TableFormula  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402

from shared.xlsx_verify import assert_clean  # noqa: E402

GRID = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"),
              bottom=Side(style="thin"))
MAIN, DONE, LEGEND = "Main Schedule", "Completed", "Legend"
MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
          "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
DAY_FMT = "dddd d"
# new column layout: # goes in front of PHASE; the calendar starts at J
COL_NUM, COL_PHASE, COL_DATE, COL_CREW, CAL0 = 6, 7, 8, 9, 10
HEADERS = ["SUPERINTENDENT", "PROJECT", "ADDRESS", "CITY", "BUILDER", "#",
           "PHASE", "DATE READY", "CREW"]
TABLES = {  # table -> (sheet, banner text, phase list name)
    "Flatwork": (MAIN, "FLATWORK", "list_Flat"),
    "Foundation": (MAIN, "FOUNDATION", "list_Fnd"),
    "Done_Flatwork": (DONE, "FLATWORK", "list_Flat"),
    "Done_Foundation": (DONE, "FOUNDATION", "list_Fnd"),
}
STYLES = [  # name, legend meaning, fill, font colour
    ("Task Completed", "Task Completed", Color(theme=8, tint=0.5999938962981048), "FF000000"),
    ("Projected", "Projected", "FFFFEB9C", "FF9C5700"),
    ("Final", "Final", "FFC6EFCE", "FF006100"),
    ("Hold / Cancelled", "Hard Stops / Holds / Cancelled", "FFFFC7CE", "FF9C0006"),
]


def die(msg: str) -> None:
    sys.exit(f"STOP: {msg}")


def clean(v) -> str:
    return re.sub(r"\s+", " ", str(v or "")).strip().upper()


def copy_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy.copy(src._style)


# ---------------------------------------------------------------- formulas
def in_flat(c: str) -> str:
    return f"OR(ISREF({c} col_Flat),ISREF({c} col_DoneFlat))"


def in_fnd(c: str) -> str:
    return f"OR(ISREF({c} col_Fnd),ISREF({c} col_DoneFnd))"


def num_formula(r: int) -> str:
    g = f"G{r}"
    return (f"=IF({g}=\"\",\"\",IFERROR(MATCH({g},IF({in_flat(g)},list_Flat,"
            f"list_Fnd),0),99))")


# ---------------------------------------------------------------- layout
def find_layout(ws) -> dict:
    """Banner / header / data rows of the two sections and the punch lists."""
    rows = {}
    for r in range(1, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if a.startswith("FLATWORK") and "flat" not in rows:
            rows["flat"] = r
        elif a.startswith("FOUNDATION") and "fnd" not in rows:
            rows["fnd"] = r
        elif "PUNCH LIST" in a and "punch" not in rows:
            rows["punch"] = r
    if set(rows) != {"flat", "fnd", "punch"} or not rows["flat"] < rows["fnd"] < rows["punch"]:
        die(f"could not find FLATWORKS / FOUNDATIONS / PUNCH LIST in order: {rows}")
    for k in ("flat", "fnd"):
        if clean(ws.cell(rows[k] + 1, 1).value) != "SUPERINTENDENT":
            die(f"row {rows[k] + 1} should be the SUPERINTENDENT header")
    return rows


def last_used(ws, r0: int, r1: int) -> int:
    last = r0
    for r in range(r0, r1 + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            last = r
    return last


def calendar_days(ws) -> list:
    """Mon-Sat dates from the first day in row 2 through 12/31, checked against
    the day numbers already typed in row 2 so a wrong start date stops the build."""
    month = year = None
    for c in range(1, ws.max_column + 1):
        m = re.match(r"([A-Z]+)\s+(\d{4})", clean(ws.cell(1, c).value))
        if m and m.group(1) in MONTHS:
            month, year = MONTHS.index(m.group(1)) + 1, int(m.group(2))
            break
    if not month:
        die("no 'MONTH YYYY' title in row 1")
    typed = []
    for c in range(CAL0, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, dt.datetime):
            typed.append(v.date())
        elif v not in (None, ""):
            m = re.search(r"(\d+)\s*$", str(v))
            typed.append(int(m.group(1)) if m else None)
    if not typed:
        die("row 2 has no day labels")
    first = typed[0] if isinstance(typed[0], dt.date) else dt.date(year, month, typed[0])
    days, d = [], first
    while d.year == first.year:
        if d.weekday() < 6:
            days.append(d)
        d += dt.timedelta(days=1)
    for i, t in enumerate(typed):
        want = days[i]
        if (t if isinstance(t, int) else t.day) != want.day:
            die(f"row 2 label #{i + 1} says {t}, calendar expects {want:%m/%d/%Y}")
    return days


def insert_columns(ws) -> None:
    """# in front of PHASE (old F)."""
    merges = [str(m) for m in ws.merged_cells.ranges]
    for m in merges:
        ws.unmerge_cells(m)
    # a width record can cover a RANGE of columns (min..max: v2 keeps I:K as one, N:AF as
    # another) - expand it to every column, or all but the first fall back to Excel's
    # narrow default and the dates show ######## (the first v3, 2026-10-06)
    from openpyxl.utils import column_index_from_string as ci
    widths = {}
    for k, v in ws.column_dimensions.items():
        lo, hi = (v.min or ci(k)), (v.max or ci(k))
        for i in range(lo, hi + 1):
            widths[L(i)] = v.width
    ws.insert_cols(6)
    ws.column_dimensions.clear()
    for k, w in widths.items():
        i = ci(k)
        n = i + (i >= 6)
        ws.column_dimensions[L(n)].width = w
    ws.column_dimensions[L(COL_NUM)].width = 5
    for r in range(1, ws.max_row + 1):
        copy_style(ws.cell(r, 5), ws.cell(r, COL_NUM))
    for m in merges:  # only A:H banners exist; widen them over the new columns
        a, b = m.split(":")
        rr = re.sub(r"\D", "", a)
        if a.startswith("A") and b.startswith("H"):
            ws.merge_cells(f"A{rr}:I{rr}")
        else:
            die(f"unexpected merged range {m}")


def extend_calendar(ws, days: list, band_rows: list) -> int:
    """Real dates in the day bands, style copied rightward, month names in row 1."""
    last_old = max(c for c in range(CAL0, ws.max_column + 1) if ws.cell(2, c).value not in (None, ""))
    last = CAL0 + len(days) - 1
    w = ws.column_dimensions[L(last_old)].width or 13
    month_cell = None
    for c in range(1, ws.max_column + 1):
        if re.match(r"[A-Z]+\s+\d{4}", clean(ws.cell(1, c).value)):
            month_cell = ws.cell(1, c)
    month_style = copy.copy(month_cell._style)
    for c in range(1, ws.max_column + 1):
        if re.match(r"[A-Z]+\s+\d{4}", clean(ws.cell(1, c).value)):
            ws.cell(1, c).value = None
    for c in range(last_old + 1, last + 1):
        ws.column_dimensions[L(c)].width = w
        for r in range(1, ws.max_row + 1):
            copy_style(ws.cell(r, last_old), ws.cell(r, c))
    prev = None
    for i, d in enumerate(days):
        c = CAL0 + i
        for r in band_rows:
            ws.cell(r, c).value = d
            ws.cell(r, c).number_format = DAY_FMT
        if d.month != prev:
            ws.cell(1, c).value = f"{MONTHS[d.month - 1]} {d.year}"
            ws.cell(1, c)._style = copy.copy(month_style)
            ws.cell(1, c).alignment = Alignment(horizontal="left")
            prev = d.month
    return last


def table_rows(ws, rows: dict) -> dict:
    """Data rows of each section: header+1 .. one spare blank row."""
    out = {}
    for k, nxt in (("flat", rows["fnd"]), ("fnd", rows["punch"])):
        hdr = rows[k] + 1
        last = last_used(ws, hdr, nxt - 1)
        end = max(last, hdr) + 1  # one blank row to insert into
        if end >= nxt - 1:
            die(f"no blank gap row under the {k} section (rows {hdr}-{nxt}); add one and rerun")
        out[k] = (hdr, end)
    return out


def sort_rows(ws, r0: int, r1: int, key) -> None:
    """Reorder rows r0..r1 (values, styles, comments) by key(row) - stable."""
    snap = []
    for r in range(r0, r1 + 1):
        cells = [(copy.copy(ws.cell(r, c).value), copy.copy(ws.cell(r, c)._style),
                  ws.cell(r, c).comment) for c in range(1, ws.max_column + 1)]
        snap.append((key(r), r, cells, ws.row_dimensions[r].height))
    snap.sort(key=lambda s: (s[0], s[1]))
    for i, (_, _, cells, h) in enumerate(snap):
        r = r0 + i
        for c, (v, st, cm) in enumerate(cells, start=1):
            cell = ws.cell(r, c)
            cell.value, cell._style = v, st
            cell.comment = Comment(cm.text, cm.author) if cm else None
        ws.row_dimensions[r].height = h


def make_table(ws, name: str, hdr: int, end: int, last_col: int) -> None:
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(hdr, i).value = h
    for c in range(CAL0, last_col + 1):  # Excel needs a name per table column; the band
        d = ws.cell(2, c).value          # stays visually blank like v2 (text = fill colour)
        cell = ws.cell(hdr, c)
        cell.value = f"{d:%m/%d}"
        cell.number_format = "@"
        f = copy.copy(cell.font)
        f.color = copy.copy(cell.fill.fgColor) if cell.fill.fill_type else Color(rgb="FFFFFFFF")
        cell.font = f
    ref = f"A{hdr}:{L(last_col)}{end}"
    t = Table(displayName=name, ref=ref)
    t._initialise_columns()
    for col, h in zip(t.tableColumns, [ws.cell(hdr, c).value for c in range(1, last_col + 1)]):
        col.name = str(h)
    # calculated column: Excel fills # into every row added to the table
    t.tableColumns[COL_NUM - 1].calculatedColumnFormula = TableFormula(attr_text=num_formula(hdr + 1)[1:])
    t.autoFilter = AutoFilter(ref=ref)
    t.sortState = SortState(ref=f"A{hdr + 1}:{L(last_col)}{end}", sortCondition=[
        SortCondition(ref=f"{L(COL_NUM)}{hdr + 1}:{L(COL_NUM)}{end}"),
        SortCondition(ref=f"{L(COL_DATE)}{hdr + 1}:{L(COL_DATE)}{end}")])
    ws.add_table(t)
    for r in range(hdr + 1, end + 1):
        ws.cell(r, COL_NUM).value = num_formula(r)
        ws.cell(r, COL_NUM).alignment = Alignment(horizontal="center")
        for c in range(1, last_col + 1):  # one even grid (v2 had rows with no borders)
            ws.cell(r, c).border = GRID


# ---------------------------------------------------------------- legend
def build_legend(wb) -> None:
    lg = wb[LEGEND]
    fnd, flat, swatches, note = [], [], [], None
    sec = "fnd"
    for r in range(1, lg.max_row + 1):
        a, b, c = lg.cell(r, 1).value, lg.cell(r, 2).value, lg.cell(r, 3).value
        if clean(b).startswith("FLATWORK PHASES"):
            sec = "flat"
            continue
        if isinstance(a, (int, float)) and b:
            (fnd if sec == "fnd" else flat).append((clean(b), c))
        elif clean(b) == "FOUNDATION COMPLETED":
            note = (b, c)
        elif lg.cell(r, 2).fill.fill_type and c:
            swatches.append(c)
    if len(fnd) < 5 or len(flat) < 3 or len(swatches) != 4:
        die(f"legend not as expected: {len(fnd)} foundation, {len(flat)} flatwork, {len(swatches)} colours")
    hdr_style = copy.copy(lg.cell(1, 2)._style)
    body_style = copy.copy(lg.cell(3, 2)._style)
    for row in lg.iter_rows():
        for c in row:
            c.value, c._style = None, copy.copy(body_style)
    lg.column_dimensions["A"].width = 6

    def put_table(name, title, items, r0):
        for i, h in enumerate(["#", title, "DESCRIPTION"], start=1):
            lg.cell(r0, i).value = h
            lg.cell(r0, i)._style = copy.copy(hdr_style)
        for k, (ph, desc) in enumerate(items, start=1):
            lg.cell(r0 + k, 1).value = f"=ROW()-ROW({name}[#Headers])"
            lg.cell(r0 + k, 1).alignment = Alignment(horizontal="center")
            lg.cell(r0 + k, 2).value = ph
            lg.cell(r0 + k, 3).value = desc
        ref = f"A{r0}:C{r0 + len(items)}"
        t = Table(displayName=name, ref=ref)
        t._initialise_columns()
        for col, h in zip(t.tableColumns, ["#", title, "DESCRIPTION"]):
            col.name = h
        # no calculated column here: openpyxl-written ones on this table made Excel
        # repair the file (2026-10-06, both a Table[#Headers] and a $A$1 form)
        lg.add_table(t)
        return r0 + len(items)

    end = put_table("Phases_Foundation", "FOUNDATION PHASE", fnd, 1)
    end = put_table("Phases_Flatwork", "FLATWORK PHASE", flat, end + 4)
    r = end + 4
    lg.cell(r, 2).value, lg.cell(r, 3).value = "COLOR (Cell Styles)", "MEANING"
    lg.cell(r, 2)._style = lg.cell(r, 3)._style = copy.copy(hdr_style)
    for i, (name, meaning, *_ignored) in enumerate(STYLES, start=1):
        lg.cell(r + i, 2).style = name
        lg.cell(r + i, 2).value = name
        lg.cell(r + i, 3).value = meaning
    r += len(STYLES) + 2
    if note:
        lg.cell(r, 2).value, lg.cell(r, 3).value = note
        lg.cell(r, 2).font = Font(name="Calibri", sz=11, b=True)
    wb.defined_names["list_Fnd"] = DefinedName("list_Fnd", attr_text="Phases_Foundation[FOUNDATION PHASE]")
    wb.defined_names["list_Flat"] = DefinedName("list_Flat", attr_text="Phases_Flatwork[FLATWORK PHASE]")
    return [p for p, _ in fnd], [p for p, _ in flat]


def add_styles(wb, cal_cell) -> None:
    side = Side(style="thin")
    for name, _m, fill, font in STYLES:
        st = NamedStyle(name=name)
        st.font = Font(name=cal_cell.font.name or "Arial", sz=cal_cell.font.sz or 10, color=font)
        st.fill = PatternFill("solid", fgColor=fill if isinstance(fill, Color) else Color(rgb=fill))
        st.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        st.border = Border(left=side, right=side, top=side, bottom=side)
        wb.add_named_style(st)


# ---------------------------------------------------------------- fake data
def fake_rows(ws, rng, phases_ok, phases_other) -> None:
    import itertools
    cyc = itertools.cycle(phases_ok + [phases_other[0], "OLD PHASE NAME"])
    for i, r in enumerate(range(*rng), start=1):
        if all(ws.cell(r, c).value in (None, "") for c in range(1, 6)):
            continue
        ws.cell(r, 1).value = f"SUPER {i % 4 + 1}"
        ws.cell(r, 2).value = f"RP9{r:03d}"
        ws.cell(r, 3).value = f"{100 + r} TEST STREET"
        ws.cell(r, 4).value = "TESTVILLE"
        ws.cell(r, 5).value = "TEST BUILDER"
        ws.cell(r, COL_PHASE).value = next(cyc)
        ws.cell(r, COL_CREW).value = f"CREW {i % 3 + 1}"
        for c in range(CAL0, ws.max_column + 1):
            if ws.cell(r, c).value not in (None, ""):
                ws.cell(r, c).value = f"TASK {c - CAL0 + 1} / CREW {i % 3 + 1}"


# ---------------------------------------------------------------- main
def build(src: Path, out: Path, fake: bool = False) -> None:
    if src.resolve() == out.resolve():
        die("OUT must be a new file, never the source")
    wb = load_workbook(src)
    ws = wb[MAIN]
    if ws.tables or wb[LEGEND].tables:
        die("source already has tables - it looks like v3 already")
    rows = find_layout(ws)
    insert_columns(ws)
    days = calendar_days(ws)
    punch_bands = [r for r in range(rows["punch"], ws.max_row + 1)
                   if clean(ws.cell(r, CAL0).value).replace(" ", "") == "MONDAY"]
    last_col = extend_calendar(ws, days, [2] + punch_bands)
    trs = table_rows(ws, rows)

    cal_sample = next((ws.cell(r, c) for r in range(trs["flat"][0] + 1, trs["flat"][1])
                       for c in range(CAL0, last_col) if ws.cell(r, c).value), ws.cell(4, CAL0))
    add_styles(wb, cal_sample)
    fnd_list, flat_list = build_legend(wb)

    for k, (hdr, end) in trs.items():  # trim phase text; legend words must match exactly
        for r in range(hdr + 1, end + 1):
            v = ws.cell(r, COL_PHASE).value
            if isinstance(v, str):
                ws.cell(r, COL_PHASE).value = re.sub(r"\s+", " ", v).strip()
    if fake:
        fake_rows(ws, (trs["flat"][0] + 1, trs["flat"][1] + 1), flat_list, fnd_list)
        fake_rows(ws, (trs["fnd"][0] + 1, trs["fnd"][1] + 1), fnd_list, flat_list)
        for r in range(rows["punch"], ws.max_row + 1):
            for c in (1, 3, 4, 5):
                if ws.cell(r, c).value and clean(ws.cell(r, c).value) not in (
                        "NAME", "ADDRESS", "CITY", "BUILDER", "SUPERINTENDENT", "PROJECT"):
                    if "PUNCH LIST" in clean(ws.cell(r, c).value):
                        ws.cell(r, c).value = f"SUPER {r % 4 + 1} PUNCH LIST"
                    else:
                        ws.cell(r, c).value = "TEST"

    for (k, lst) in (("flat", flat_list), ("fnd", fnd_list)):
        hdr, end = trs[k]
        sort_rows(ws, hdr + 1, end - 1, lambda r, lst=lst: (
            lst.index(clean(ws.cell(r, COL_PHASE).value)) if clean(ws.cell(r, COL_PHASE).value) in lst else 99,
            ws.cell(r, COL_DATE).value if isinstance(ws.cell(r, COL_DATE).value, dt.datetime) else dt.datetime.max))

    # Completed = an empty copy of Main, laid out BEFORE Main gets its tables
    old = wb[DONE]
    if old.max_row > 1 or old["A1"].value:
        die("the Completed sheet is not empty - move its rows into the new one by hand")
    idx = wb.sheetnames.index(DONE)
    wb.remove(old)
    done = wb.copy_worksheet(ws)
    done.title = DONE
    wb.move_sheet(done, offset=idx - wb.sheetnames.index(DONE))
    for m in [str(x) for x in done.merged_cells.ranges]:
        done.unmerge_cells(m)
    for k in ("fnd", "flat"):  # drop data rows bottom-up, keep header + one blank row
        hdr, end = trs[k]
        if end - hdr > 1:
            done.delete_rows(hdr + 2, end - hdr - 1)
    for row in done.iter_rows(min_row=4):
        for c in row:
            if c.row > 2 and c.column > 0:
                c.comment = None
    d_rows = find_layout(done)
    for k in ("flat", "fnd"):
        done.merge_cells(f"A{d_rows[k]}:I{d_rows[k]}")
        hdr = d_rows[k] + 1
        for c in range(1, last_col + 1):  # the kept row is a blank slot: no value, no status fill
            done.cell(hdr + 1, c).value = None
            done.cell(hdr + 1, c).fill = PatternFill(fill_type=None)
    for r in list(range(d_rows["flat"] + 3, d_rows["fnd"])) + list(range(d_rows["fnd"] + 3, d_rows["punch"])):
        for c in range(1, last_col + 1):  # gap rows between the tables: nothing on them
            done.cell(r, c).value = None
            done.cell(r, c).fill = PatternFill(fill_type=None)
    for r in range(d_rows["punch"], done.max_row + 1):  # empty punch lists, keep headers + day bands
        band = isinstance(done.cell(r, CAL0).value, dt.datetime)
        for c in range(1, last_col + 1):
            v = clean(done.cell(r, c).value)
            if v and not ("PUNCH LIST" in v or v in ("NAME", "ADDRESS", "CITY", "BUILDER", "DESCRIPTION",
                                                    "DATE", "CREW", "PROJECT", "SUPERINTENDENT")) \
                    and not (band and c >= CAL0):
                done.cell(r, c).value = None
                done.cell(r, c).fill = PatternFill(fill_type=None)
    done.freeze_panes = f"{L(CAL0)}4"

    dmap = {"flat": "Done_Flatwork", "fnd": "Done_Foundation"}
    for k in ("flat", "fnd"):
        hdr = d_rows[k] + 1
        make_table(done, dmap[k], hdr, hdr + 1, last_col)
    make_table(ws, "Flatwork", trs["flat"][0], trs["flat"][1], last_col)
    make_table(ws, "Foundation", trs["fnd"][0], trs["fnd"][1], last_col)
    for nm, tb in (("col_Flat", "Flatwork"), ("col_Fnd", "Foundation"),
                   ("col_DoneFlat", "Done_Flatwork"), ("col_DoneFnd", "Done_Foundation")):
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"{tb}[PHASE]")

    # each table owns its dropdown (Excel bans "which table am I in" tests inside a
    # dropdown rule); the list travels with a row cut into the matching Completed
    # table; a row moved into the wrong table keeps its old list.
    for sh, name in ((ws, "Flatwork"), (ws, "Foundation"), (done, "Done_Flatwork"), (done, "Done_Foundation")):
        t = sh.tables[name]
        r0, r1 = int(re.search(r"(\d+):", t.ref).group(1)) + 1, int(re.search(r"(\d+)$", t.ref).group(1))
        dv = DataValidation(type="list", formula1=TABLES[name][2], allow_blank=True,
                            showErrorMessage=True, errorTitle="Phase",
                            error="Pick a phase from this table's list (Legend tab).")
        dv.add(f"{L(COL_PHASE)}{r0}:{L(COL_PHASE)}{r1}")
        sh.add_data_validation(dv)

    rng = f"${L(CAL0)}$2:${L(last_col)}$2"
    flag = (f"=IFERROR(IF(SUMPRODUCT(--('{MAIN}'!{rng}<>'{DONE}'!{rng}))=0,\"\","
            f"\"DATES DON'T MATCH\"),\"DATES DON'T MATCH\")")
    for sh in (ws, done):
        sh["A1"].value = flag
        sh["A1"].font = Font(name="Arial", sz=10, b=True, color="FFC00000")
    done["C1"].value = "COMPLETED JOBS"
    done["C1"].font = Font(name="Arial", sz=14, b=True)
    ws.freeze_panes = f"{L(CAL0)}4"
    from openpyxl.worksheet.views import Selection
    for sh in (ws, done):  # Excel's own two-way-freeze form: exactly these three panes
        tl = f"{L(CAL0)}4"
        sh.sheet_view.selection = [Selection(pane="topRight", activeCell=f"{L(CAL0)}1", sqref=f"{L(CAL0)}1"),
                                   Selection(pane="bottomLeft", activeCell="A4", sqref="A4"),
                                   Selection(pane="bottomRight", activeCell=tl, sqref=tl)]
    wb.save(out)
    assert_clean(out)
    print(f"OK  {out}")
    print(f"    calendar {days[0]:%m/%d/%Y} - {days[-1]:%m/%d/%Y}: {len(days)} days, columns "
          f"{L(CAL0)}:{L(last_col)}")
    for k, (hdr, end) in trs.items():
        print(f"    {('Flatwork' if k == 'flat' else 'Foundation'):10} rows {hdr}-{end}")
    print(f"    legend: {len(fnd_list)} foundation, {len(flat_list)} flatwork phases; "
          f"styles: {', '.join(s[0] for s in STYLES)}")


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(a) != 2:
        die("usage: schedule_v3_build.py SRC.xlsx OUT.xlsx [--fake]")
    build(Path(a[0]), Path(a[1]), fake="--fake" in sys.argv)
