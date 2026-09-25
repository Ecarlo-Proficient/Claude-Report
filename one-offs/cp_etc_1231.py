"""CP ETC fetch for the 12-31-25 WIP — READ-ONLY.

The 12-31-25 WIP carries a PLACEHOLDER ETC on most commercial rows:
ETC = Total Contract Price / 1.10, i.e. a flat 10% markup, which shows up
as a dead-flat 9.09% gross profit on every one of those rows. That is not a
budget - it is a stand-in. This tool replaces nothing; it goes to the source
(the commercial takeoff's BID sheet) and reports what the real ETC is.

GUARD (same rule as wip_master_asof.py): if the takeoff ETC implies gross
profit over 30%, the read is almost certainly wrong (missing scope, stale
bid revision, wrong takeoff file) and is reported as REVIEW rather than
treated as fact.

Writes NOTHING to the WIP workbook. Optional CSV out for the audit trail.

  python3 ".../one-offs/cp_etc_1231.py"
  python3 ".../one-offs/cp_etc_1231.py" --only CP653,CP689,CP714
"""
import argparse
import csv
import os
import re
import sys
from pathlib import Path

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))
from shared.takeoff_etc import find_takeoff_etc                # noqa: E402

HIST = Path(os.getenv("WIP_HISTORY_DIR",
            str(Path.home() / "Library/CloudStorage/OneDrive-ProficientConcrete,LLC"
                / "Company Files - WIP Report" / "WIP History")))
CP_ROOT = Path(os.getenv("CP_ROOT",
                         "/Volumes/Common/CURRENT PROJECTS/Awarded Projects Commercial projects"))
CP_ROOTS = [CP_ROOT, CP_ROOT / "Completed Projects"]

FALLBACK_PROJ, FALLBACK_CONTRACT, FALLBACK_ETC = 0, 4, 5
PLACEHOLDER_DIVISOR = 1.10
GP_GUARD = 0.30


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").upper().strip())


def find_snapshot():
    """The 12-31-25 snapshot, whatever the estimators named it."""
    if not HIST.is_dir():
        return None, f"WIP History folder not found: {HIST}"
    pats = (r"12[-._]31[-._ ]*2?5", r"2025[-._]12[-._]31")
    hits = [f for f in HIST.iterdir()
            if f.suffix.lower() in (".xlsb", ".xlsx")
            and not f.name.startswith("~$")
            and any(re.search(p, f.name) for p in pats)]
    if not hits:
        names = ", ".join(sorted(f.name for f in HIST.iterdir())[:20])
        return None, f"No 12-31-25 snapshot in {HIST}. Saw: {names}"
    hits.sort(key=lambda f: f.stat().st_mtime, reverse=True)
    return hits[0], None


def read_grid(path):
    if path.suffix.lower() == ".xlsb":
        from pyxlsb import open_workbook as open_xlsb
        with open_xlsb(str(path)) as wb:
            sheet = next((s for s in wb.sheets if "MASTER" in _norm(s)), wb.sheets[0])
            with wb.get_sheet(sheet) as sh:
                return [[c.v for c in row] for row in sh.rows()], sheet
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = next((s for s in wb.sheetnames if "MASTER" in _norm(s)), wb.sheetnames[0])
    grid = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return grid, sheet


def resolve_cols(grid):
    """Columns BY HEADER - the estimators rearrange this file."""
    proj = contract = etc = None
    hdr_row = None
    for i, row in enumerate(grid[:10]):
        for j, cell in enumerate(row or []):
            n = _norm(cell)
            if not n:
                continue
            if proj is None and n in ("PROJECT", "PROJECT #", "JOB", "JOB #", "PROJECT NUMBER"):
                proj, hdr_row = j, i
            if contract is None and "CONTRACT PRICE" in n:
                contract, hdr_row = j, i
            if etc is None and ("ESTIMATED TOTAL COST" in n or n == "TOTAL COSTS"
                                or n == "TOTAL COST" or "ESTIMATED COST" in n):
                etc, hdr_row = j, i
    used_fallback = []
    if proj is None:
        proj, _ = FALLBACK_PROJ, used_fallback.append("PROJECT")
    if contract is None:
        contract, _ = FALLBACK_CONTRACT, used_fallback.append("CONTRACT")
    if etc is None:
        etc, _ = FALLBACK_ETC, used_fallback.append("ETC")
    return proj, contract, etc, hdr_row, used_fallback


def index_cp():
    idx = {}
    for root in CP_ROOTS:
        try:
            for e in root.iterdir():
                if e.is_dir():
                    m = re.search(r"(CP\d{3,4})", e.name.upper())
                    if m and m.group(1) not in idx:
                        idx[m.group(1)] = e
        except OSError:
            pass
    return idx


def money(v):
    return f"${v:,.0f}" if isinstance(v, (int, float)) else "-"


def pct(v):
    return f"{v * 100:.2f}%" if isinstance(v, (int, float)) else "-"


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--only", help="Comma-separated CP numbers, e.g. CP653,CP689")
    ap.add_argument("--snapshot", help="Path to the 12-31-25 workbook (skips auto-discovery)")
    ap.add_argument("--csv", nargs="?", const="", metavar="PATH",
                    help="Also write a CSV (default ~/Downloads/CP ETC 12-31-25.csv)")
    args = ap.parse_args()

    only = None
    if args.only:
        only = {p.strip().upper() for p in args.only.split(",") if p.strip()}

    snap = Path(args.snapshot).expanduser() if args.snapshot else None
    if snap is None:
        snap, err = find_snapshot()
        if snap is None:
            print(f"  ERROR: {err}")
            return 2
    print(f"\n  SNAPSHOT   {snap}")

    grid, sheet = read_grid(snap)
    proj_c, con_c, etc_c, hdr_row, fb = resolve_cols(grid)
    print(f"  SHEET      {sheet!r}  (header row {hdr_row + 1 if hdr_row is not None else '?'})")
    print(f"  COLUMNS    project={proj_c}  contract={con_c}  etc={etc_c}"
          + (f"   [FALLBACK POSITION USED FOR: {', '.join(fb)} - verify]" if fb else ""))

    folders = index_cp()
    if not folders:
        print(f"  ERROR: no CP folders under {CP_ROOT} - is the Common volume mounted?")
        return 2
    print(f"  TAKEOFFS   {len(folders)} CP folders indexed")
    print(f"  GUARD      takeoff ETC implying GP > {GP_GUARD * 100:.0f}% is reported, never trusted\n")

    rows = []
    seen = set()
    for r in grid:
        if not r or len(r) <= max(proj_c, con_c, etc_c):
            continue
        m = re.search(r"(CP\d{3,4})", _norm(r[proj_c]))
        if not m:
            continue
        key = m.group(1)
        if key in seen or (only and key not in only):
            continue
        seen.add(key)
        name = ""
        for j in range(proj_c + 1, min(proj_c + 4, len(r))):
            if isinstance(r[j], str) and r[j].strip():
                name = r[j].strip()
                break
        con = r[con_c] if isinstance(r[con_c], (int, float)) else None
        old = r[etc_c] if isinstance(r[etc_c], (int, float)) else None
        placeholder = bool(con and old and abs(old - con / PLACEHOLDER_DIVISOR) <= 2)
        fol = folders.get(key)
        new = note = src = None
        if fol is None:
            note = "No CP folder found"
        else:
            try:
                p, new, note, _frag = find_takeoff_etc(fol, key, "slab", "")
                src = p.name if p else None
            except Exception as e:
                note = f"{type(e).__name__}: {e}"
        gp_old = (con - old) / con if con and old else None
        gp_new = (con - new) / con if con and isinstance(new, (int, float)) and new else None
        if not isinstance(new, (int, float)) or not new:
            verdict = "NO ETC READ"
        elif gp_new is not None and gp_new > GP_GUARD:
            verdict = "REVIEW - GP>30%"
        elif gp_new is not None and gp_new < 0:
            verdict = "REVIEW - NEGATIVE GP"
        else:
            verdict = "USE"
        rows.append(dict(project=key, name=name, contract=con, wip_etc=old,
                         wip_gp=gp_old, placeholder=placeholder, takeoff_etc=new,
                         takeoff_gp=gp_new, verdict=verdict, note=note or "", file=src or ""))

    if only:
        for k in sorted(only - seen):
            print(f"  NOT ON THE 12-31-25 WIP: {k}")

    w = (10, 34, 13, 13, 8, 13, 8, 20)
    hdr = ("PROJECT", "NAME", "CONTRACT", "WIP ETC", "GP%", "TAKEOFF ETC", "GP%", "VERDICT")
    line = "  " + " ".join(h.ljust(x) if i < 2 else h.rjust(x)
                           for i, (h, x) in enumerate(zip(hdr, w)))
    print(line)
    print("  " + "-" * (len(line) - 2))
    for d in sorted(rows, key=lambda x: x["project"]):
        cells = (d["project"].ljust(w[0]), d["name"][:w[1]].ljust(w[1]),
                 money(d["contract"]).rjust(w[2]),
                 (money(d["wip_etc"]) + ("*" if d["placeholder"] else "")).rjust(w[3]),
                 pct(d["wip_gp"]).rjust(w[4]), money(d["takeoff_etc"]).rjust(w[5]),
                 pct(d["takeoff_gp"]).rjust(w[6]), d["verdict"].rjust(w[7]))
        print("  " + " ".join(cells))
    print("\n  * WIP ETC is the contract / 1.10 placeholder (flat 9.09% GP), not a budget\n")

    for d in sorted(rows, key=lambda x: x["project"]):
        if d["verdict"] != "USE":
            print(f"  {d['project']}: {d['verdict']} - {d['note']}"
                  + (f"  [{d['file']}]" if d["file"] else ""))
        elif d["note"]:
            print(f"  {d['project']}: source {d['note']}  [{d['file']}]")

    n_use = sum(1 for d in rows if d["verdict"] == "USE")
    print(f"\n  {len(rows)} CP rows - {n_use} usable ETC, {len(rows) - n_use} need a human\n")

    if args.csv is not None:
        out = Path(args.csv).expanduser() if args.csv else \
            Path.home() / "Downloads" / "CP ETC 12-31-25.csv"
        out.parent.mkdir(parents=True, exist_ok=True)
        with open(out, "w", newline="") as f:
            wr = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else ["project"])
            wr.writeheader()
            wr.writerows(rows)
        print(f"  CSV -> {out}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
