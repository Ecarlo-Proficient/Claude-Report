#!/usr/bin/env python3
"""
invoice_project_audit.py - invoices billed to the WRONG customer for the job
they name. Read-only against QBO.

The case that started it (the owner 2026-09-08): RP6586's lot-preparation
invoice #33911 was billed to the parent builder, not to the project
sub-customer, with the job only in the memo. The project P&L, the RP WIP
file and every "invoices by customer" pull missed it - the job looked
27.5k under-billed. The owner re-pointed it by hand and asked for a check
that finds every other invoice like it: "any other invoices that are not
under their project that HAS a project in QBO that is clearly missing its
invoice."

WHAT COUNTS. An invoice is OFF ITS PROJECT when its memo (PrivateNote, the
customer-facing memo, or a line description) names exactly ONE job number,
that job EXISTS as a customer in QBO, and the invoice's CustomerRef is not
that customer. The report says who it IS billed to - the project's own
parent (the usual slip), another project, or an unrelated customer - and
how many invoices already sit under the project, so a project with NONE
under it ("clearly missing") stands out. Memos naming no job, or two, are
counted and left alone; a memo naming a job that has NO customer in QBO
goes on its own sheet (that is a missing project, not a mis-pointed
invoice).

WHAT THE FIRST RUN TAUGHT (2026-09-08, 2,464 invoices since 2025-01-01):
  * Until AUGUST 2025 residential invoices were routinely billed to the
    parent builder with the job in the memo - 600 of them, 54-113 a month -
    then it stops dead (1 in September, 2 in October, none after). That is a
    practice change, not 600 slips. So the report splits at
    --practice-since (default 2025-09-01): invoices from then on are the
    headline, earlier ones sit on their own "Before the change" sheet.
  * "Billed to another project" is nearly always the base job and its -FTW
    sibling (313 of 317): a memo "RP4362 - ... - Flatwork" on the RP4362-FTW
    customer is the memo missing its -FTW, and the customer is right. Only a
    sibling whose SCOPE contradicts the customer is flagged (a Foundation
    invoice on the -FTW customer, a Flatwork invoice on the base) - the rest
    are counted as "memo missing -FTW" and left alone.
  * CP / MFD invoices to the parent GC with the job in the memo are common
    practice there; they are listed, but they are the owner's call.

Fixing is a QBO edit - open the invoice, change the customer to the project
- never done by this script. Re-run after the fixes and the rows disappear.

Output: <CompanyHealth>/Invoices Off Project.xlsx
  Off project        one row per mis-pointed invoice, worst first
  No project in QBO  memo names a job with no customer
  By project         every project that has an off-project invoice: under / off / billed
  Summary            counts by division and relationship

  python3 one-offs/invoice_project_audit.py                 # invoices since 2025-01-01
  python3 one-offs/invoice_project_audit.py --since 2024-01-01
  python3 one-offs/invoice_project_audit.py --all           # every invoice on the books
  python3 one-offs/invoice_project_audit.py --practice-since 2025-06-01
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import quote

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
from shared import paths, pnl_paths, xlsx_verify                    # noqa: E402
from shared.qbo_api import PROJ_RE, customer_url, load_credentials, query_all  # noqa: E402
from shared.rp_invoicing import STAGE_RE, scope_of                   # noqa: E402

# a flatwork scope - belongs on the -FTW customer, never the base job
FLATWORK_RE = re.compile(r"flatwork|driveway|sidewalk|patio|lead\s*walk|approach|"
                         r"city\s*walk|porch|steps?\b|pool\s*deck|paving", re.IGNORECASE)


def _n(p: str) -> str:
    return p.upper().replace(" ", "")


def memo_projects(inv: dict) -> set:
    """Every job number the invoice's own text names (customer NOT included)."""
    found = set()
    for text in (inv.get("PrivateNote"), (inv.get("CustomerMemo") or {}).get("value")):
        found |= {_n(x) for x in PROJ_RE.findall(text or "")}
    if not found:
        for ln in inv.get("Line") or []:
            found |= {_n(x) for x in PROJ_RE.findall(ln.get("Description") or "")}
    return found


def invoice_url(txn_id: str, realm: str) -> str:
    return (f"https://qbo.intuit.com/app/login?pagereq="
            f"{quote(f'invoice?txnId={txn_id}')}&deeplinkcompanyid={realm}")


def _date(s) -> dt.date | None:
    try:
        return dt.date.fromisoformat(str(s)[:10])
    except (TypeError, ValueError):
        return None


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--since", default="2025-01-01", help="invoices dated on/after (YYYY-MM-DD)")
    ap.add_argument("--all", action="store_true", help="every invoice, no date filter")
    ap.add_argument("--practice-since", default="2025-09-01",
                    help="invoices from this date are expected ON the project (headline); "
                         "earlier ones report separately (default 2025-09-01, when "
                         "parent-billing of RP invoices stopped)")
    ap.add_argument("--out", type=Path,
                    default=paths.companyhealth_dir() / "Invoices Off Project.xlsx")
    a = ap.parse_args()

    access, realm = load_credentials()
    customers = query_all(access, realm, "Customer")
    cust_by_id = {c["Id"]: c for c in customers if c.get("Id")}
    proj_cust: dict = {}                       # job # -> customer record (first wins)
    for c in customers:
        m = PROJ_RE.search(c.get("DisplayName") or c.get("CompanyName") or "")
        if m:
            proj_cust.setdefault(_n(m.group(0)), c)
    invoices = query_all(access, realm, "Invoice",
                         where="" if a.all else f"TxnDate >= '{a.since}'")
    print(f"{len(customers)} customers ({len(proj_cust)} projects) · "
          f"{len(invoices)} invoices{'' if a.all else ' since ' + a.since}")

    def cust_proj(cid: str) -> str:
        c = cust_by_id.get(cid) or {}
        m = PROJ_RE.search(c.get("DisplayName") or c.get("CompanyName") or "")
        return _n(m.group(0)) if m else ""

    under: Counter = Counter()                 # job -> invoices billed to its own customer
    off_rows, missing_rows = [], []
    n_none = n_two = n_ok = n_memo_ftw = 0
    month: dict = defaultdict(Counter)         # yyyy-mm -> what invoices were billed to
    practice = _date(a.practice_since)
    for inv in invoices:
        cref = inv.get("CustomerRef") or {}
        cid, cname = cref.get("value") or "", cref.get("name") or ""
        cp = cust_proj(cid)
        named = memo_projects(inv)
        ym = str(inv.get("TxnDate") or "")[:7]
        if cp:
            under[cp] += 1
        if not named:
            n_none += 1
            month[ym]["memo names no job"] += 1
            continue
        if len(named) > 1:
            n_two += 1
            continue
        job = named.pop()
        if job == cp:
            n_ok += 1
            month[ym]["on its project"] += 1
            continue
        # the base job and its -FTW sibling are one family: the memo usually
        # just omits the -FTW. Flag only when the SCOPE contradicts the side.
        if cp and cp.split("-")[0] == job.split("-")[0]:
            sc = scope_of(inv.get("PrivateNote") or (inv.get("CustomerMemo") or {}).get("value") or "")
            on_ftw = cp.endswith("-FTW")
            contradicts = ((on_ftw and STAGE_RE.search(sc) and not FLATWORK_RE.search(sc))
                           or (not on_ftw and FLATWORK_RE.search(sc) and not STAGE_RE.search(sc)))
            if not contradicts:
                n_memo_ftw += 1
                month[ym]["memo missing -FTW (customer right)"] += 1
                continue
        target = proj_cust.get(job)
        rec = {"doc": inv.get("DocNumber") or inv.get("Id"), "id": inv.get("Id"),
               "date": _date(inv.get("TxnDate")), "amount": float(inv.get("TotalAmt") or 0),
               "balance": float(inv.get("Balance") or 0), "billed_to": cname, "job": job,
               "division": pnl_paths.division_of(job) or "?",
               "memo": (inv.get("PrivateNote") or (inv.get("CustomerMemo") or {}).get("value") or "")}
        if target is None:
            missing_rows.append(rec)
            continue
        tparent = (target.get("ParentRef") or {}).get("value")
        if cp and cp.split("-")[0] == job.split("-")[0]:
            rel = "its -FTW/base SIBLING (scope contradicts)"
        elif cp:
            rel = "another PROJECT"
        elif tparent and tparent == cid:
            rel = "the project's PARENT"
        else:
            rel = "unrelated customer"
        month[ym][f"billed to {rel}"] += 1
        rec.update({"should_be": target.get("DisplayName") or "", "should_be_id": target["Id"],
                    "rel": rel, "recent": bool(rec["date"] and practice and rec["date"] >= practice)})
        off_rows.append(rec)
    for r in off_rows:
        r["under"] = under.get(r["job"], 0)
    legacy_rows = [r for r in off_rows if not r["recent"]]
    off_rows = [r for r in off_rows if r["recent"]]

    # worst first: projects with NOTHING under them, then by amount
    off_rows.sort(key=lambda r: (r["under"] > 0, -r["amount"]))
    by_job: dict = defaultdict(lambda: {"off": 0, "off_amt": 0.0, "division": "", "should_be": ""})
    for r in off_rows:
        g = by_job[r["job"]]
        g["off"] += 1; g["off_amt"] += r["amount"]
        g["division"], g["should_be"] = r["division"], r["should_be"]

    print(f"\n  ok (billed to their own project): {n_ok}")
    print(f"  memo names no job: {n_none} · names two jobs (left alone): {n_two} · "
          f"memo missing -FTW but customer right: {n_memo_ftw}")
    print(f"  before {a.practice_since} (parent-billing was the practice): "
          f"{len(legacy_rows)} invoice(s), ${sum(r['amount'] for r in legacy_rows):,.0f} "
          f"- on the 'Before the change' sheet")
    print(f"  OFF THEIR PROJECT since {a.practice_since}: {len(off_rows)} invoice(s) on "
          f"{len(by_job)} project(s), ${sum(r['amount'] for r in off_rows):,.0f}")
    print(f"    projects with NOTHING under them: "
          f"{sum(1 for j, g in by_job.items() if under.get(j, 0) == 0)}")
    for k, n in Counter((r["division"], r["rel"]) for r in off_rows).most_common():
        print(f"    {k[0]:<4} billed to {k[1]:<22} {n}")
    print(f"  memo names a job with NO customer in QBO: {len(missing_rows)}")

    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active; ws.title = "Off project"
    ws.append(["Invoice #", "Date", "Amount", "Open balance", "Billed to", "Should be (project)",
               "Project #", "Division", "Billed to is", "Invoices already under the project",
               "Memo", "Open invoice", "Open project"])
    for r in off_rows:
        ws.append([r["doc"], r["date"], r["amount"], r["balance"], r["billed_to"], r["should_be"],
                   r["job"], r["division"], r["rel"], r["under"], r["memo"],
                   invoice_url(r["id"], realm), customer_url(r["should_be_id"], realm)])
    wsl = wb.create_sheet("Before the change")
    wsl.append([f"Invoices dated before {a.practice_since}, when billing the parent builder with the "
                f"job in the memo was the practice - listed, not flagged"])
    wsl.append(["Invoice #", "Date", "Amount", "Open balance", "Billed to", "Should be (project)",
                "Project #", "Division", "Billed to is", "Invoices already under the project",
                "Memo", "Open invoice"])
    for r in sorted(legacy_rows, key=lambda r: (r["under"] > 0, -r["amount"])):
        wsl.append([r["doc"], r["date"], r["amount"], r["balance"], r["billed_to"], r["should_be"],
                    r["job"], r["division"], r["rel"], r["under"], r["memo"],
                    invoice_url(r["id"], realm)])
    ws2 = wb.create_sheet("No project in QBO")
    ws2.append(["Invoice #", "Date", "Amount", "Billed to", "Job named in memo", "Division", "Memo",
                "Open invoice"])
    for r in missing_rows:
        ws2.append([r["doc"], r["date"], r["amount"], r["billed_to"], r["job"], r["division"],
                    r["memo"], invoice_url(r["id"], realm)])
    ws3 = wb.create_sheet("By project")
    ws3.append(["Project #", "Division", "Project customer", "Invoices under it",
                "Invoices off it", "Off amount"])
    for job, g in sorted(by_job.items(), key=lambda kv: (under.get(kv[0], 0) > 0, -kv[1]["off_amt"])):
        ws3.append([job, g["division"], g["should_be"], under.get(job, 0), g["off"], g["off_amt"]])
    ws4 = wb.create_sheet("Summary")
    ws4.append(["Division", "Billed to is", "Invoices", "Amount"])
    tot: dict = defaultdict(lambda: [0, 0.0])
    for r in off_rows:
        t = tot[(r["division"], r["rel"])]; t[0] += 1; t[1] += r["amount"]
    for (d, rel), (n, amt) in sorted(tot.items()):
        ws4.append([d, rel, n, amt])
    ws4.append([]); ws4.append(["Run", dt.datetime.now().strftime("%m/%d/%Y %I:%M %p"),
                                "" if a.all else f"invoices since {a.since}",
                                f"headline = dated on/after {a.practice_since}"])
    ws4.append([]); ws4.append(["Month", "What invoices were billed to (every invoice, by month)"])
    kinds = sorted({k for c in month.values() for k in c})
    ws4.append(["Month"] + kinds)
    for ym in sorted(month):
        ws4.append([ym] + [month[ym].get(k, 0) for k in kinds])
    for w in (ws, wsl, ws2, ws3, ws4):
        hdr_row = 2 if w is wsl else 1
        for c in w[hdr_row]:
            c.font = Font(bold=True)
        w.freeze_panes = f"A{hdr_row + 1}"
        for row in w.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.date):
                    c.number_format = "mm/dd/yyyy"
                elif isinstance(c.value, float):
                    c.number_format = '"$"#,##0_);[Red]("$"#,##0)'
                elif isinstance(c.value, str) and c.value.startswith("https://"):
                    c.hyperlink = c.value; c.value = "open"; c.font = Font(color="0563C1", underline="single")
    for w, widths in ((ws, (11, 12, 13, 13, 34, 40, 11, 9, 22, 12, 60, 12, 12)),
                      (wsl, (11, 12, 13, 13, 34, 40, 11, 9, 22, 12, 60, 12)),
                      (ws2, (11, 12, 13, 34, 14, 9, 60, 12)),
                      (ws3, (11, 9, 40, 12, 12, 13)), (ws4, (10, 24, 10, 13))):
        for i, wd in enumerate(widths, start=1):
            w.column_dimensions[w.cell(1, i).column_letter].width = wd
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    xlsx_verify.assert_clean(a.out)
    print(f"\nwrote {a.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
