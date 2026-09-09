#!/usr/bin/env python3
"""wip_month_end.py - the month-end WIP report as of a cutoff date, in the
bank layout the 12-31-25 and 3-31-26 reports use, with the RESIDENTIAL section
rebuilt from the month's crew schedules (the user 2026-09-09: "CP/MFD you can
do [from the master] but not for RP ... RP is a machine that we need to try to
stop to see where it's at").

HOW THE REPORT IS BUILT (the standard split - see the wip-schedule-standard
reference: estimator data from the master / the folder, accounting data from QBO):
  * MFD + CP: population, contract, ETC, type and bonded = the live master's
    Test-Master tab. Nothing invented. (CP/MFD add a job rarely; the master is
    current for them.)
  * RP: every daily schedule from the prior month-end through the cutoff is
    read ('Main Schedule' tab, PROJECT column, section bands; FLATWORK rows and
    flatwork WRECK rows are the -FTW line). That roster - plus the lines the RP
    WIP file and the master already carry - is the universe. Each line gets a
    STATUS as of the cutoff from what the schedule and QBO show:
      COMPLETED            poured (reached WRECK) and off the cutoff schedule,
                           or billed out (billed reaches the contract); says
                           whether billing is still open
      ACTIVE               on the cutoff schedule before wreck, or seen in the
                           month and neither poured nor billed out (a line that
                           left the schedule mid-stage is flagged to verify)
      COMPLETED EARLIER    not on any schedule in the window, billed out before
                           the month - REMOVE from the RP file
      OFF-SCHEDULE, COSTS  not scheduled, has costs, not billed out (the
                           flatwork-started-off-schedule class)
      BACKLOG              not scheduled, no costs, no billing
    Contract / ETC per line, in this order of trust, each stamped with its
    source: the master (Test-Master) -> the RP WIP file -> JobTread (the
    APPROVED proposal's price + cost; a proposal only fills a line whose
    contract it matches, and never a -FTW line when the base slab's QBO
    billing already matches it - the user 2026-09-09 "peer into JobTread")
    -> the project folder (proposal PDF SUB TOTAL + takeoff cost sheet,
    accepted only when the pair implies a 5-35% margin, else left blank and
    named) -> the General List price for tract builders (flagged, the list is
    unmaintained) -> the invoice total when a line is billed and nothing else
    prices it (flagged).
  * COSTS TO DATE + BILLED TO DATE = QBO project P&L dated 2019-01-01 .. the
    cutoff - the SAME pull the live readers do, only date-bounded.
  * The WIP sheet carries MFD + CP + the RP lines that are ACTIVE or COMPLETED
    with billing still open. Billed-out completed lines, off-schedule and
    backlog lines are on the snapshot sheet only (the same sections the
    bank-facing Test-Master already excludes).

SHEET 2 = 'RP Snapshot <date>': every RP line with its status, stage at the
cutoff, first/last day on the schedule, contract and ETC with sources, QBO
costs and billing, last invoice, whether it is on the WIP sheet, and the RP
file action (ADD / REMOVE / keep) so the RP file can be brought to the day
after the cutoff. Lines whose numbers no source could price say so.

OUTPUT: '<WIP History>/WIP <m-d-yy>.xlsx', fingerprint-scrubbed and verified
with shared.xlsx_verify as the LAST step. Never touches the master, the RP
file or QBO. Operator detail goes to stdout only.

Usage:
  python3 one-offs/wip_month_end.py                       (prior month-end)
  python3 one-offs/wip_month_end.py --cutoff 2026-08-31
  python3 one-offs/wip_month_end.py --cutoff 2026-08-31 --rp-from-master
  python3 one-offs/wip_month_end.py --cutoff 2026-08-31 --out-dir ~/Downloads
"""
from __future__ import annotations

import argparse
import calendar
import datetime as dt
import re
import shutil
import sys
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))
sys.path.insert(0, str(_REPO / "one-offs"))
sys.path.insert(0, str(_REPO / "wip"))
from shared import paths, qbo_api, qbo_vault                   # noqa: E402
from shared.takeoff_etc import _norm, find_takeoff_etc         # noqa: E402
from shared.xlsx_verify import assert_clean                    # noqa: E402
import rp_wip_reader as RP                                     # noqa: E402
from rp_schedule_wip_preview import (                          # noqa: E402
    _is_tract, find_proposal, read_main_schedule)
from rp_jobtread_coverage import ORG_ID as JT_ORG, pave as jt_pave   # noqa: E402

_WIP_DIR = paths.onedrive_base() / "Company Files - WIP Report"
# The master was renamed 'WIP - MASTER new.xlsx' -> 'WIP - MASTER.xlsx' on 2026-09-09;
# take whichever exists so neither name breaks a run (WIP_EXCEL_PATH still wins).
MASTER = paths.get_path(
    "WIP_EXCEL_PATH",
    next((f for f in (_WIP_DIR / "WIP - MASTER.xlsx", _WIP_DIR / "WIP - MASTER new.xlsx") if f.exists()),
         _WIP_DIR / "WIP - MASTER.xlsx"))
HISTORY = paths.get_path(
    "WIP_HISTORY_DIR",
    paths.onedrive_base() / "Company Files - WIP Report" / "WIP History")
RP_FILE = paths.get_path("RP_WIP_FILE", paths.onedrive_base() / "RP WIP TO FIX_Final.xlsx")
SCHED_ROOT = paths.get_path("RP_SCHEDULE_ROOT", "/Volumes/Common/OPERATIONS/SCHEDULE")
_SCHED_RE = re.compile(r"Schedule (\d{1,2})-(\d{1,2})-(\d{2})\.xlsx$", re.IGNORECASE)
QBO_START = "2019-01-01"
FULL_TOL = 0.01          # billed within 1% of the contract = fully billed
GP_LO, GP_HI = 0.05, 0.35   # a folder-read contract/ETC pair must land here

FONT = "Tahoma"
MONEY = '"$"#,##0_);[Red]("$"#,##0)'
PCT = "0.00%"
DATE = "mm/dd/yyyy"
HDR_FILL = PatternFill("solid", fgColor="D9D9D9")
_thin = Side(style="thin", color="000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

HEADERS = [
    "TYPE", "PROJECT #", "PROJECT NAME", "BONDED", "TOTAL CONTRACT PRICE",
    "ESTIMATED TOTAL COSTS", "ORIGINAL PROFIT", "GROSS PROFIT %",
    "COSTS TO DATE", "COST TO COMPLETE", "PERCENT COMPLETE",
    "REVENUES EARNED TO DATE", "PROFIT EARNED TO DATE", "BILLED TO DATE",
    "OVERBILLINGS", "UNDERBILLINGS", "LEFT TO BILL", "FUTURE PROFIT TO EARN",
    "PURE JOB BORROW",
]
FORMULAS = {
    7:  '=IF(OR(E{r}="",F{r}=""),"",E{r}-F{r})',
    8:  '=IF(OR(G{r}="",E{r}="",E{r}=0),"",G{r}/E{r})',
    10: '=IF(OR(F{r}="",I{r}=""),"",F{r}-I{r})',
    11: '=IF(OR(I{r}="",F{r}="",F{r}=0),"",I{r}/F{r})',
    12: '=IF(OR(E{r}="",I{r}="",F{r}="",F{r}=0),"",E{r}*I{r}/F{r})',
    13: '=IF(OR(L{r}="",I{r}=""),"",L{r}-I{r})',
    15: '=IF(OR(N{r}="",L{r}=""),"",MAX(N{r}-L{r},0))',
    16: '=IF(OR(L{r}="",N{r}=""),"",MAX(L{r}-N{r},0))',
    17: '=IF(OR(E{r}="",N{r}=""),"",E{r}-N{r})',
    18: '=IF(OR(G{r}="",M{r}=""),"",G{r}-M{r})',
    19: '=IF(OR(J{r}="",Q{r}=""),"",MAX(J{r}-Q{r},0))',
}
MONEY_COLS = {5, 6, 7, 9, 10, 12, 13, 14, 15, 16, 17, 18, 19}
PCT_COLS = {8, 11}
CENTER_COLS = {1, 2, 4}
WIDTHS = {"A": 27.66, "B": 11.5, "C": 42.83, "D": 7.66, "E": 21.0, "F": 22.0,
          "G": 16.16, "H": 15.33, "I": 14.33, "J": 17.16, "L": 23.83, "M": 22.0,
          "N": 15.33, "O": 13.33, "P": 14.33, "Q": 13.33, "R": 22.0, "S": 16.16}

# ── the RP statuses, in the order the snapshot lists them ──
S_ACTIVE, S_DONE, S_EARLIER, S_OFFSCHED, S_BACKLOG = (
    "ACTIVE", "COMPLETED", "COMPLETED EARLIER", "OFF-SCHEDULE, COSTS", "BACKLOG")
STATUS_ORDER = [S_ACTIVE, S_DONE, S_EARLIER, S_OFFSCHED, S_BACKLOG]
_SLAB_WRECK = re.compile(r"SLAB|FOOTER|FOOTING|PIER|BEAM|GRADE|FOUNDATION")
_FTW_WRECK = re.compile(r"FLATWORK|FTW|PAVING|DRIVEWAY|SIDEWALK|PATIO|POOL DECK|APPROACH")


# ────────────────────────── helpers ──────────────────────────
def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def _label(s):
    """The TYPE label with the master's dash separators dropped, ASCII only
    ('Residential - Custom - Slab' -> 'Residential Custom Slab', the way the
    3-31-26 report writes it)."""
    s = str(s or "")
    s = re.sub(r"\s+[-–—]\s+", " ", s)
    return s.encode("ascii", "ignore").decode()


def _fmt_date(d: dt.date) -> str:
    return f"{d.month}-{d.day}-{d.year % 100:02d}"


def _mdy(d) -> str:
    return d.strftime("%m/%d/%Y") if d else ""


def _prior_month_end(cutoff: dt.date) -> dt.date:
    return cutoff.replace(day=1) - dt.timedelta(days=1)


def _schedule_files():
    """[(date, path)] of every 'Schedule m-d-yy.xlsx' under SCHED_ROOT."""
    out = []
    for f in SCHED_ROOT.glob("*/*/Schedule *.xlsx"):
        m = _SCHED_RE.search(f.name)
        if not m:
            continue
        mo, dy, yy = (int(g) for g in m.groups())
        try:
            out.append((dt.date(2000 + yy, mo, dy), f))
        except ValueError:
            continue
    return sorted(out)


def _fully_billed(K, B) -> bool:
    return bool(K and K > 0 and B is not None and B >= K * (1 - FULL_TOL))


# ────────────────────────── sources ──────────────────────────
def read_master():
    """Test-Master rows (values) + the company prefix from WIP Master!B1."""
    wb = load_workbook(MASTER, data_only=True)
    title = str(wb["WIP Master"]["B1"].value or "").strip()
    company = title.split(" - ")[0].strip() if " - " in title else title
    ws = wb["Test-Master"]
    hdr = next((r for r in range(1, 16) if ws.cell(r, 2).value == "PROJECT #"), 3)
    rows = []
    for r in range(hdr + 1, ws.max_row + 1):
        proj = str(ws.cell(r, 2).value or "").strip().upper()
        if not proj or proj.startswith("TOTAL") or not qbo_api.PROJ_RE.search(proj):
            continue
        rows.append({
            "type": _label(ws.cell(r, 1).value), "proj": proj,
            "name": str(ws.cell(r, 3).value or "").strip(),
            "bonded": "Yes" if str(ws.cell(r, 4).value or "").strip().upper() in ("Y", "YES") else "No",
            "contract": _num(ws.cell(r, 5).value), "etc": _num(ws.cell(r, 6).value),
            "tab_costs": _num(ws.cell(r, 9).value), "tab_billed": _num(ws.cell(r, 14).value),
        })
    wb.close()
    return company, rows


def read_rp_file():
    """Every RP line in the owner's RP WIP file (all bands): contract / ETC /
    address / builder by project #. First copy wins on duplicates."""
    out = {}
    if not RP_FILE.exists():
        return out
    wb = load_workbook(RP_FILE, data_only=True)
    ws = wb["RP WIP"] if "RP WIP" in wb.sheetnames else wb.worksheets[0]
    for r in range(1, ws.max_row + 1):
        job = str(ws.cell(r, 1).value or "").strip().upper()
        if not re.match(r"^RP\d{4}(-FTW)?$", job) or job in out:
            continue
        out[job] = {"address": str(ws.cell(r, 2).value or "").strip(),
                    "builder": str(ws.cell(r, 3).value or "").strip(),
                    "contract": _num(ws.cell(r, 4).value), "etc": _num(ws.cell(r, 5).value)}
    wb.close()
    return out


def _qbo_name(proj_map, proj):
    """The address part of the QBO customer name ('RP####-FTW - <address>' ->
    '<address>'); '' when there is no project."""
    name = (proj_map.get(proj) or {}).get("name") or ""
    return re.sub(r"^\s*" + re.escape(proj) + r"\s*[-:]?\s*", "", name, flags=re.IGNORECASE).strip()


def qbo_totals_at(access, cid, cust_id, cutoff_iso):
    t = qbo_api.extract_pl_totals(qbo_api.fetch_project_pl(access, cid, cust_id, QBO_START, cutoff_iso))
    billed = round(float(t.get("income", 0) or 0), 2)
    costs = round(float(t.get("cogs", 0) or 0) + float(t.get("expenses", 0) or 0), 2)
    return billed, costs


def invoices_by_project(access, cid, cutoff_iso):
    """{project# -> [(date, doc#, total)]} for every invoice dated on/before the
    cutoff since 2025-01-01, attributed with the ONE invoice->job rule."""
    inv = qbo_api.query_all(access, cid, "Invoice",
                            f"TxnDate >= '2025-01-01' AND TxnDate <= '{cutoff_iso}'")
    by = defaultdict(list)
    for i in inv:
        p = qbo_api.project_of_invoice(i)
        if p:
            by[p.upper()].append((dt.date.fromisoformat(i["TxnDate"]), i.get("DocNumber"),
                                  float(i.get("TotalAmt") or 0)))
    for v in by.values():
        v.sort()
    return by


# ────────────────────────── the RP roster ──────────────────────────
def build_roster(prior_end: dt.date, cutoff: dt.date, known_ftw: set):
    """Every RP line on any daily schedule dated prior_end..cutoff ->
    {line: {address, builder, days: [date], stages: [(date, section, desc)]}}.
    FLATWORK-band rows are the -FTW line; a WRECK row is flatwork when its
    description says so, or when it is ambiguous (pads, steps, curb...) and the
    job has a flatwork line anywhere else."""
    files = [(d, f) for d, f in _schedule_files() if prior_end <= d <= cutoff]
    raw = []
    for d, f in files:
        for rec in read_main_schedule(f):
            if rec["job"].upper().startswith("RP"):
                raw.append((d, rec))
    ftw = set(known_ftw) | {rec["job"].upper() + "-FTW" for _d, rec in raw
                            if rec["section"] == "FLATWORK"}
    roster = {}
    for d, rec in raw:
        job = rec["job"].upper()
        desc = _norm(rec["desc"])
        if rec["section"] == "WRECK":
            if _SLAB_WRECK.search(desc):
                scope = "slab"
            elif _FTW_WRECK.search(desc):
                scope = "ftw"
            else:
                scope = "ftw" if job + "-FTW" in ftw else "slab"
        else:
            scope = rec["scope"]
        line = job + ("-FTW" if scope == "ftw" else "")
        r = roster.setdefault(line, {"address": rec["address"], "builder": rec["builder"],
                                     "days": [], "stages": []})
        if d not in r["days"]:
            r["days"].append(d)
        r["stages"].append((d, rec["section"], desc))
    for r in roster.values():
        r["days"].sort()
        r["stages"].sort(key=lambda s: s[0])
    return files, roster


def jt_proposals(numbers):
    """{job# -> [(price, cost, date)]} of APPROVED customerOrder documents in
    JobTread, read-only. {} when the grant key is absent or the API fails."""
    out = {}
    try:
        key = qbo_vault.get("JT_GRANT_KEY")
    except Exception:
        return out
    if not key:
        return out
    for n in sorted(numbers):
        try:
            r = jt_pave(key, {"organization": {"$": {"id": JT_ORG}, "jobs": {
                "$": {"size": 3, "where": {"and": [["number", "=", n]]}},
                "nodes": {"documents": {"$": {"size": 50}, "nodes": {
                    "type": {}, "status": {}, "price": {}, "cost": {}, "createdAt": {}}}}}}})
        except Exception as e:
            print(f"    JobTread {n}: {type(e).__name__}")
            continue
        docs = [d for j in r["organization"]["jobs"]["nodes"] for d in j["documents"]["nodes"]
                if d.get("type") == "customerOrder" and d.get("status") == "approved"]
        if docs:
            out[n] = [(float(d.get("price") or 0), float(d.get("cost") or 0),
                       str(d.get("createdAt") or "")[:10]) for d in docs]
    return out


def resolve_numbers(line, master_rp, rp_file, roster_rec, folders, gl_by_job, invs,
                    jt=None, base_billed=None):
    """(contract, contract_source, etc, etc_source, note) for one RP line."""
    m, f = master_rp.get(line), rp_file.get(line)
    K = E = None
    ks = es = ""
    note = []
    if m and m.get("contract"):
        K, ks = m["contract"], "WIP master"
    elif f and f.get("contract"):
        K, ks = f["contract"], "RP WIP file"
    if m and m.get("etc"):
        E, es = m["etc"], "WIP master"
    elif f and f.get("etc"):
        E, es = f["etc"], "RP WIP file"
    if K and E:
        return K, ks, E, es, ""

    job = line.replace("-FTW", "")
    scope = "ftw" if line.endswith("-FTW") else "slab"
    if jt:
        price = round(sum(p for p, _c, _d in jt), 2)
        cost = round(sum(c for _p, c, _d in jt), 2)
        when = max(d for _p, _c, d in jt)
        src = f"JobTread approved proposal ({when})"
        # A -FTW line: the job's approved proposal is normally the SLAB. Only
        # take it when the slab has not been billed at all, and say so.
        if scope == "ftw" and base_billed:
            what = "matches the base billing" if abs(base_billed - price) <= 0.05 * price \
                else f"base billed {base_billed:,.0f}"
            note.append(f"JobTread proposal {price:,.0f} / cost {cost:,.0f} is the base job's "
                        f"({what}) - not this flatwork line")
        elif scope == "ftw" and not K:
            K, ks = price, src + " - scope unverified (slab or flatwork?)"
            if not E and cost:
                E, es = cost, src + " - scope unverified"
            if K and E:
                return K, ks, E, es, "; ".join(note)
        elif scope == "ftw":
            pass
        elif K and abs(K - price) > 0.05 * K:
            note.append(f"JobTread proposal {price:,.0f} / cost {cost:,.0f} does not match the "
                        f"contract {K:,.0f} - a different scope; not used")
        else:
            if not K:
                K, ks = price, src
            if not E and cost:
                E, es = cost, src
            if K and E:
                return K, ks, E, es, "; ".join(note)
    builder = (roster_rec or {}).get("builder") or (f or {}).get("builder") or ""
    desc = roster_rec["stages"][-1][2] if roster_rec and roster_rec["stages"] else ""
    folder = folders.get(line)
    if folder is not None:
        pk = pe = None
        prop = None
        pnote = tnote = ""
        try:
            prop, pk, pnote = find_proposal(folder, scope, desc)
        except Exception as e:                       # a bad PDF must not stop the run
            prop, pk, pnote = None, None, f"proposal unreadable ({type(e).__name__})"
        try:
            _t, pe, tnote, _frag = find_takeoff_etc(folder, job, scope, desc)
        except Exception as e:
            pe, tnote = None, f"takeoff unreadable ({type(e).__name__})"
        cand_k = K or pk
        cand_e = E or pe
        if cand_k and cand_e:
            gp = (cand_k - cand_e) / cand_k
            if GP_LO <= gp <= GP_HI:
                if not K and pk:
                    K, ks = pk, f"proposal PDF ({prop.name})"
                if not E and pe:
                    E, es = pe, "takeoff cost sheet"
            else:
                note.append(f"folder pair implausible: proposal {pk or 'none'} / takeoff "
                            f"{pe or 'none'} -> {gp:.0%} margin; estimator to price")
        elif cand_k and not K and pk:
            K, ks = pk, f"proposal PDF ({prop.name})"
            note.append(tnote or "no takeoff budget")
        elif not cand_k and pe:
            note.append(f"takeoff reads {pe:,.0f} but no contract to check it against; "
                        + (pnote or "no proposal"))
        elif not cand_k:
            note.append(pnote or "no proposal PDF")
    else:
        note.append("no project folder found")

    if not K and _is_tract(builder):
        rec = gl_by_job.get(job)
        price = (rec or {}).get("flat_bid" if scope == "ftw" else "slab_bid")
        if price:
            K, ks = price, "General List price (unmaintained since 07/30 - verify)"
            note.append("tract builder - contract from P.O.s / price list")
    if not K and invs:
        K, ks = round(sum(a for _d, _n, a in invs), 2), "invoice total (verify - nothing else prices it)"
    return K, ks, E, es, "; ".join(n for n in note if n)


def classify(line, r, K, B, C, last_inv, month_start, cutoff, cutoff_sd):
    """-> (status, stage_at_cutoff, detail)."""
    days = (r or {}).get("days") or []
    stages = (r or {}).get("stages") or []
    on_cutoff = bool(cutoff_sd and cutoff_sd in days)
    seen = bool(days)
    poured = any(sec == "WRECK" for _d, sec, _x in stages)
    billed_out = _fully_billed(K, B)
    inv_in_month = bool(last_inv and month_start <= last_inv <= cutoff)
    stage_now = ""
    if on_cutoff:
        last_day = [s for s in stages if s[0] == cutoff_sd]
        stage_now = f"{last_day[-1][1]}: {last_day[-1][2]}" if last_day else ""
    elif stages:
        stage_now = f"left after {stages[-1][1]}: {stages[-1][2]} ({_mdy(stages[-1][0])})"
    bill = ("billed out" if billed_out else
            f"billing open: {B or 0:,.0f} of {K:,.0f}" if K else
            f"billed {B or 0:,.0f}, no contract")
    if seen:
        if billed_out and (poured or not on_cutoff or inv_in_month):
            d = [bill]
            if on_cutoff:
                d.append("wreck/punch still on the cutoff schedule")
            if last_inv:
                d.append(f"last invoice {_mdy(last_inv)}")
            return S_DONE, stage_now, "; ".join(d)
        if poured and not on_cutoff:
            return S_DONE, stage_now, f"poured (wrecked); {bill}"
        if poured and on_cutoff and stages[-1][1] == "WRECK":
            return S_DONE, stage_now, f"poured, wreck still on the cutoff schedule; {bill}"
        d = [bill]
        if not on_cutoff:
            d.append("LEFT THE SCHEDULE MID-STAGE - verify (poured? dropped?)")
        return S_ACTIVE, stage_now, "; ".join(d)
    # never on a schedule in the window
    if inv_in_month:
        return S_DONE, "never on the schedule", f"invoiced {_mdy(last_inv)} - one-invoice job; {bill}"
    if billed_out:
        return S_EARLIER, "", f"{bill}; last invoice {_mdy(last_inv)} - remove from the RP file"
    if (C or 0) > 0 or (B or 0) > 0:
        return S_OFFSCHED, "", f"{bill}; costs {C or 0:,.0f}; not on any schedule in the window"
    return S_BACKLOG, "", "no costs, no billing, not scheduled"


# ────────────────────────── the report sheet ──────────────────────────
def _style(c, bold=False, fill=None, center=False, fmt=None, border=True, wrap=False):
    c.font = Font(name=FONT, size=8, bold=bold)
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    if center:
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    elif wrap:
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if fmt:
        c.number_format = fmt


def write_report(ws, company, rows, cutoff: dt.date):
    ws["A1"] = f"{company} - WIP REPORT"
    ws["A2"] = f"REPORT DATE: {cutoff.strftime('%b %d, %Y').upper()}"
    for a in ("A1", "A2"):
        ws[a].font = Font(name=FONT, size=8, bold=True)
    for c, h in enumerate(HEADERS, 1):
        _style(ws.cell(3, c, h), bold=True, fill=HDR_FILL, center=True)
    ws.row_dimensions[3].height = 28
    r = 4
    for row in rows:
        vals = {1: row["type"], 2: row["proj"], 3: row["name"], 4: row["bonded"],
                5: row["contract"], 6: row["etc"], 9: row["costs"], 14: row["billed"]}
        for c in range(1, 20):
            v = vals.get(c) if c not in FORMULAS else FORMULAS[c].format(r=r)
            cell = ws.cell(r, c, v)
            _style(cell, center=(c in CENTER_COLS),
                   fmt=MONEY if c in MONEY_COLS else PCT if c in PCT_COLS else None)
        ws.row_dimensions[r].height = 14
        r += 1
    last = r - 1
    tot = r
    ws.cell(tot, 1, "TOTALS")
    for c in range(1, 20):
        cell = ws.cell(tot, c)
        col = cell.column_letter
        if c in MONEY_COLS:
            cell.value = f"=SUBTOTAL(109,{col}4:{col}{last})"
        elif c == 8:
            cell.value = f'=IF(E{tot}=0,"",G{tot}/E{tot})'
        elif c == 11:
            cell.value = f'=IF(F{tot}=0,"",I{tot}/F{tot})'
        _style(cell, bold=True, fill=HDR_FILL,
               fmt=MONEY if c in MONEY_COLS else PCT if c in PCT_COLS else None)
    tab = Table(displayName="Wip" + cutoff.strftime("%b"), ref=f"A3:S{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(tab)
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.sheet_view.selection[0].activeCell = "A4"
    ws.sheet_view.selection[0].sqref = "A4"
    ws.print_area = f"A1:S{tot}"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return last, tot


# ────────────────────────── the RP snapshot sheet ──────────────────────────
SNAP_HDR = ["STATUS", "PROJECT #", "PROJECT NAME", "BUILDER", "STAGE AT CUTOFF",
            "FIRST DAY", "LAST DAY", "DAYS", "TOTAL CONTRACT PRICE", "CONTRACT SOURCE",
            "ESTIMATED TOTAL COSTS", "ETC SOURCE", "COSTS TO DATE", "PERCENT COMPLETE",
            "BILLED TO DATE", "% BILLED", "LAST INVOICE", "ON WIP SHEET", "RP FILE",
            "NOTES"]
SNAP_W = {"A": 18, "B": 12, "C": 40, "D": 28, "E": 34, "F": 11, "G": 11, "H": 6, "I": 16,
          "J": 26, "K": 16, "L": 22, "M": 14, "N": 10, "O": 14, "P": 9, "Q": 11, "R": 9,
          "S": 9, "T": 80}
SNAP_MONEY = {9, 11, 13, 15}
SNAP_PCT = {14, 16}
SNAP_DATE = {6, 7, 17}
SNAP_CENTER = {1, 2, 6, 7, 8, 14, 16, 17, 18, 19}


def write_snapshot(ws, company, snap, cutoff: dt.date):
    ws["A1"] = f"{company} - RESIDENTIAL SNAPSHOT AS OF {cutoff.strftime('%b %d, %Y').upper()}"
    ws["A2"] = ("ACTIVE + COMPLETED with billing open are on the WIP sheet. "
                "COMPLETED EARLIER = remove from the RP file. Sources name where each number came from.")
    for a in ("A1", "A2"):
        ws[a].font = Font(name=FONT, size=8, bold=(a == "A1"))
    for c, h in enumerate(SNAP_HDR, 1):
        _style(ws.cell(3, c, h), bold=True, fill=HDR_FILL, center=True)
    ws.row_dimensions[3].height = 28
    r = 4
    for status in STATUS_ORDER:
        items = [s for s in snap if s["status"] == status]
        if not items:
            continue
        cell = ws.cell(r, 1, f"{status} ({len(items)})")
        cell.font = Font(name=FONT, size=8, bold=True)
        r += 1
        for s in items:
            K, E, C, B = s["contract"], s["etc"], s["costs"], s["billed"]
            vals = [status, s["line"], s["name"], s["builder"], s["stage"],
                    s["first"], s["last"], s["days"] or None, K, s["k_src"], E, s["e_src"],
                    C, (C / E) if (E and C is not None) else None, B,
                    (B / K) if (K and B is not None) else None, s["last_inv"],
                    "Yes" if s["on_wip"] else "No", s["file_action"], s["notes"]]
            for col, v in enumerate(vals, 1):
                cell = ws.cell(r, col, v)
                fmt = (MONEY if col in SNAP_MONEY else PCT if col in SNAP_PCT
                       else DATE if col in SNAP_DATE else None)
                _style(cell, center=(col in SNAP_CENTER), fmt=fmt, wrap=(col == 20))
            ws.row_dimensions[r].height = 14
            r += 1
        r += 1
    last = r - 2
    for col, w in SNAP_W.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.sheet_view.selection[0].activeCell = "A4"
    ws.sheet_view.selection[0].sqref = "A4"
    ws.print_area = f"A1:T{max(last, 4)}"
    ws.print_title_rows = "1:3"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ────────────────────────── fingerprint scrub ──────────────────────────
def scrub(path: Path, company: str) -> None:
    """Rewrite the zip: Excel's own app properties, the company as creator, no
    [trash] parts. The file is otherwise byte-identical."""
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename.startswith("[trash]/"):
                continue
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = re.sub(rb"<Application>.*?</Application>",
                              b"<Application>Microsoft Excel</Application>", data)
                if b"<AppVersion>" in data:
                    data = re.sub(rb"<AppVersion>.*?</AppVersion>", b"<AppVersion>16.0300</AppVersion>", data)
                else:
                    data = data.replace(b"</Properties>", b"<AppVersion>16.0300</AppVersion></Properties>")
            elif item.filename == "docProps/core.xml":
                who = company.encode()
                data = re.sub(rb"(<dc:creator[^>]*>).*?(</dc:creator>)", lambda m: m.group(1) + who + m.group(2), data)
                data = re.sub(rb"(<cp:lastModifiedBy[^>]*>).*?(</cp:lastModifiedBy>)",
                              lambda m: m.group(1) + who + m.group(2), data)
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


# ────────────────────────── main ──────────────────────────
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cutoff", help="report date YYYY-MM-DD (default: the prior month-end)")
    ap.add_argument("--out-dir", help="where to write (default: the WIP History folder)")
    ap.add_argument("--rp-from-master", action="store_true",
                    help="take the RP section from Test-Master as-is (no schedule rebuild)")
    args = ap.parse_args()

    cutoff = dt.date.fromisoformat(args.cutoff) if args.cutoff else _prior_month_end(dt.date.today())
    month_start = cutoff.replace(day=1)
    prior_end = _prior_month_end(cutoff)
    month_name = calendar.month_name[cutoff.month] + f" {cutoff.year}"
    cutoff_iso = cutoff.isoformat()
    out_dir = Path(args.out_dir).expanduser() if args.out_dir else HISTORY
    out = out_dir / f"WIP {_fmt_date(cutoff)}.xlsx"
    if not MASTER.exists():
        print(f"master not mounted: {MASTER}", file=sys.stderr)
        return 2

    company, master_rows = read_master()
    rp_file = read_rp_file()
    master_rp = {r["proj"]: r for r in master_rows if r["proj"].startswith("RP")}
    fixed_rows = [r for r in master_rows if not r["proj"].startswith("RP")]
    print(f"Test-Master: {len(master_rows)} jobs ({len(master_rp)} RP) · RP file: {len(rp_file)} lines · cutoff {cutoff_iso}")

    access, cid = qbo_api.load_credentials()
    proj_map = qbo_api.build_project_customer_map(access, cid)
    cache = {}

    def totals(proj):
        if proj not in cache:
            cust = proj_map.get(proj)
            cache[proj] = qbo_totals_at(access, cid, cust["id"], cutoff_iso) if cust else None
        return cache[proj]

    # 1. MFD + CP at the cutoff (from the master)
    missing, ahead = [], []
    for row in fixed_rows:
        t = totals(row["proj"])
        if t is None:
            row["billed"], row["costs"] = 0.0, 0.0
            missing.append(row["proj"])
        else:
            row["billed"], row["costs"] = t
            for fld in ("billed", "costs"):
                tab = row.get(f"tab_{fld}")
                if tab is not None and row[fld] > tab + 1:
                    ahead.append((row["proj"], fld, row[fld], tab))
    print(f"  MFD/CP: {len(fixed_rows)} jobs pulled"
          + (f" · no QBO project: {', '.join(missing)}" if missing else "")
          + (f" · {len(ahead)} figure(s) higher than the live tab (bills dated in the month, entered after the sync)" if ahead else ""))

    # 2. RP: the roster from the schedules
    inv_by = invoices_by_project(access, cid, cutoff_iso)
    rp_rows, snap = [], []
    if args.rp_from_master:
        for row in master_rows:
            if not row["proj"].startswith("RP"):
                continue
            t = totals(row["proj"])
            row["billed"], row["costs"] = t if t else (0.0, 0.0)
            rp_rows.append(row)
        print(f"  RP: {len(rp_rows)} rows from Test-Master (--rp-from-master)")
    else:
        known_ftw = {p for p in list(rp_file) + list(master_rp) + list(proj_map) if p.endswith("-FTW")}
        files, roster = build_roster(prior_end, cutoff, known_ftw)
        cutoff_sd = max((d for d, _f in files), default=None)
        prior_sd = min((d for d, _f in files), default=None)
        print(f"  schedules: {len(files)} files {(_fmt_date(prior_sd) if prior_sd else '-')} .. "
              f"{(_fmt_date(cutoff_sd) if cutoff_sd else '-')} · {len(roster)} RP lines on them")
        month_inv = {p for p, v in inv_by.items()
                     if p.startswith("RP") and any(month_start <= d <= cutoff for d, _n, _a in v)}
        universe = sorted(set(roster) | set(rp_file) | set(master_rp) | month_inv)

        # folders only for the lines no file prices
        need = [ln for ln in universe
                if not ((master_rp.get(ln) or {}).get("contract") or (rp_file.get(ln) or {}).get("contract"))
                or not ((master_rp.get(ln) or {}).get("etc") or (rp_file.get(ln) or {}).get("etc"))]
        folders, gl_by_job = {}, {}
        if need:
            print(f"  {len(need)} line(s) need a contract or ETC - indexing the Residential folders ...")
            try:
                rp_to_folders, addr_folders = RP.index_residential(RP.RP_ROOT)
            except Exception as e:
                rp_to_folders, addr_folders = {}, []
                print(f"    folder index failed: {type(e).__name__}")
            for ln in need:
                job = ln.replace("-FTW", "")
                fl = sorted(rp_to_folders.get(job, ()), key=lambda f: (f.parent.name, f.name))
                folder = fl[0] if fl else None
                addr = (roster.get(ln) or {}).get("address") or (rp_file.get(ln) or {}).get("address") or ""
                if folder is None and addr:
                    parts = addr.split(None, 1)
                    folder = RP.match_by_address(
                        {"house": parts[0], "street": parts[1] if len(parts) > 1 else addr}, addr_folders)
                if folder is not None:
                    folders[ln] = folder
            try:
                recs, _m = RP.read_general_list(RP.ALPHA_PATH)
                gl_by_job = {r["job"]: r for r in recs}
            except Exception as e:
                print(f"    General List unreadable: {type(e).__name__}")
            jt_by_job = jt_proposals({ln.replace("-FTW", "") for ln in need})
            print(f"    JobTread: approved proposals on {len(jt_by_job)} of "
                  f"{len({ln.replace('-FTW', '') for ln in need})} jobs")
        else:
            jt_by_job = {}

        for i, ln in enumerate(universe, 1):
            r = roster.get(ln)
            m, f = master_rp.get(ln), rp_file.get(ln)
            invs = inv_by.get(ln, [])
            base = ln.replace("-FTW", "")
            base_t = totals(base) if ln.endswith("-FTW") else None
            K, ks, E, es, note = resolve_numbers(ln, master_rp, rp_file, r, folders, gl_by_job, invs,
                                                 jt=jt_by_job.get(base), base_billed=(base_t[0] if base_t else None))
            t = totals(ln)
            B, C = t if t else (None, None)
            if t is None and invs:
                B = round(sum(a for _d, _n, a in invs), 2)
                note = (note + "; " if note else "") + "no QBO project - billed from invoices"
            last_inv = invs[-1][0] if invs else None
            status, stage, detail = classify(ln, r, K, B, C, last_inv, month_start, cutoff, cutoff_sd)
            billed_out = _fully_billed(K, B)
            on_wip = status == S_ACTIVE or (status == S_DONE and not billed_out)
            in_file = ln in rp_file
            file_action = ("REMOVE" if status == S_EARLIER or (status == S_DONE and billed_out and in_file)
                           else "ADD" if (not in_file and status in (S_ACTIVE, S_DONE)) else "keep" if in_file else "")
            builder = (f or {}).get("builder") or (r or {}).get("builder") or ""
            name = ((m or {}).get("name") or (f or {}).get("address") or (r or {}).get("address")
                    or _qbo_name(proj_map, ln))
            home = ("Tract" if _is_tract(builder) else "Custom")
            if m and "Tract" in m["type"]:
                home = "Tract"
            elif m and "Custom" in m["type"]:
                home = "Custom"
            notes = "; ".join(x for x in (detail, note) if x)
            if not K:
                notes += "; NO CONTRACT - estimator to price"
            if not E:
                notes += "; NO ETC - estimator to budget"
            snap.append({"status": status, "line": ln, "name": name, "builder": builder, "stage": stage,
                         "first": (r["days"][0] if r else None), "last": (r["days"][-1] if r else None),
                         "days": (len(r["days"]) if r else 0), "contract": K, "k_src": ks, "etc": E,
                         "e_src": es, "costs": C, "billed": B, "last_inv": last_inv, "on_wip": on_wip,
                         "file_action": file_action, "notes": notes.strip("; ")})
            if on_wip:
                rp_rows.append({"type": f"Residential {home} {'Flatwork' if ln.endswith('-FTW') else 'Slab'}",
                                "proj": ln, "name": name, "bonded": "No", "contract": K, "etc": E,
                                "costs": C if C is not None else 0.0, "billed": B if B is not None else 0.0})
            if i % 25 == 0:
                print(f"    ...{i}/{len(universe)} RP lines")
        order = {s: i for i, s in enumerate(STATUS_ORDER)}
        snap.sort(key=lambda s: (order[s["status"]], s["line"]))
        rp_rows.sort(key=lambda x: (x["proj"].endswith("-FTW"), "Tract" in x["type"], x["proj"]))

    rows = fixed_rows + rp_rows

    # 3. write, scrub, verify - in that order
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = f"WIP {_fmt_date(cutoff)}"
    last, tot = write_report(ws, company, rows, cutoff)
    if snap:
        ws2 = wb.create_sheet(f"RP Snapshot {_fmt_date(cutoff)}")
        write_snapshot(ws2, company, snap, cutoff)
        ws2.sheet_view.tabSelected = False
    wb.active = 0
    ws.sheet_view.tabSelected = True
    tmp = Path(tempfile.mkdtemp()) / out.name
    wb.save(tmp)
    scrub(tmp, company)
    assert_clean(tmp)
    shutil.move(str(tmp), str(out))
    assert_clean(out)

    # 4. operator summary (stdout only)
    print(f"\nWIP {_fmt_date(cutoff)}: {len(rows)} jobs, rows 4-{last}, TOTALS row {tot}")
    div = defaultdict(lambda: [0, 0.0, 0.0, 0.0, 0.0])
    for r in rows:
        d = "MFD" if r["proj"].startswith("MFD") else "CP" if r["proj"].startswith("CP") else "RP"
        for key in (d, "ALL"):
            a = div[key]
            a[0] += 1
            a[1] += r["contract"] or 0
            a[2] += r["etc"] or 0
            a[3] += r["costs"] or 0
            a[4] += r["billed"] or 0
    print(f"  {'DIV':4} {'jobs':>4} {'contract':>14} {'ETC':>14} {'costs':>14} {'billed':>14}  GP%   %compl")
    for key in ("MFD", "CP", "RP", "ALL"):
        n, k, e, c, b = div[key]
        gp = (k - e) / k if k else 0
        pc = c / e if e else 0
        print(f"  {key:4} {n:4} {k:14,.0f} {e:14,.0f} {c:14,.0f} {b:14,.0f} {gp:6.2%} {pc:6.2%}")
    if snap:
        print(f"\nRP snapshot ({month_name}): {len(snap)} lines")
        for status in STATUS_ORDER:
            items = [s for s in snap if s["status"] == status]
            print(f"  {status}: {len(items)}  (on WIP sheet {sum(1 for s in items if s['on_wip'])})")
        adds = [s["line"] for s in snap if s["file_action"] == "ADD"]
        rems = [s["line"] for s in snap if s["file_action"] == "REMOVE"]
        nok = [s["line"] for s in snap if s["on_wip"] and not (s["contract"] and s["etc"])]
        print(f"  RP file: ADD {len(adds)}: {' '.join(adds)}")
        print(f"  RP file: REMOVE {len(rems)}: {' '.join(rems)}")
        print(f"  on the WIP sheet without a contract or ETC ({len(nok)}): {' '.join(nok)}")
    print(f"\n  verified clean -> {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
