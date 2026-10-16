#!/usr/bin/env python3
"""
credit_memo_reclass.py - move a client's backcharge credit memos from the parent customer onto the
PROJECT, renumber them to the client's backcharge PDF, and put back every payment QuickBooks unapplies
when the credit's customer changes.

Why the payment step: changing a credit memo's customer makes QuickBooks drop the credit off the payment
that used it (the credit reopens, the payment no longer adds up). Seen on BKGP23054 (10/14, put back by
hand 10/15) and BKGP23052 (10/16). So every payment a credit sits on is saved BEFORE the change and
re-saved with its exact lines right after, then checked to the cent.

    credit_memo_reclass.py before                 the BEFORE workbook: every 2026 parent-level credit of
                                                  the customer, the payments it is applied on (snapshot),
                                                  the address -> RP # candidates
    credit_memo_reclass.py match                  read the backcharge PDF attached to each credit memo in
                                                  QuickBooks, fill the AFTER columns
                                                  (new Ref # = BKGP + PDF #, PO # = the full BRIX number, project = address
                                                  RP # + -FTW when the backcharge is flatwork work)
    credit_memo_reclass.py apply                  dry run, live from QuickBooks, rows marked Approve = Y
    credit_memo_reclass.py apply --commit         write: per payment, move its credits, re-save the payment
                                                  as it was, verify; --only <payment id> limits the run
    credit_memo_reclass.py move <credit id> <RP#> one credit already on a project -> another project, payment
                                                  put back right after; --commit
    credit_memo_reclass.py restore <payment id>   dry run: re-save a payment exactly as the mirror holds it
                                                  (a credit already moved by hand that reopened); --commit

Options: --customer "Grand Homes" (default), --since 2026-01-01.
Snapshots of every record before a write: ~/Library/Logs/Proficient/credit-memo-reclass/.
Workbook + state: CompanyHealth/Analysis/<Customer> credit memos (<date>)/.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import requests  # noqa: E402

from shared import paths, payment_reapply, qbo_api, qbo_mirror as mirror  # noqa: E402

EPS = 0.005
LOG_DIR = Path.home() / "Library" / "Logs" / "Proficient" / "credit-memo-reclass"
PO_FIELD = {"DefinitionId": "1", "Name": "Customer PO #", "Type": "StringType"}

# Date = the PDF PO's date (owner 2026-10-16). Every field written is read back; a mismatch stops the run.
# Backcharge text that names flatwork work -> the job's -FTW project; anything else -> the slab.
FTW_WORDS = ("flatwork", "ftw", "driveway", "drive way", "sidewalk", "side walk", "walkway", "patio",
             "approach", "sod", "irrigation", "flat work", "turnkey", "ardex", "stoop", "porch", "curb")

_SUFFIX = {"way", "dr", "drive", "trail", "trl", "road", "rd", "lane", "ln", "court", "ct", "circle", "cir",
           "view", "st", "street", "ave", "avenue", "blvd", "pl", "place", "pkwy", "cove", "cv", "loop"}
_ADDR = re.compile(r"\b(\d{2,6})\s+([A-Za-z][A-Za-z']+)(?:\s+([A-Za-z][A-Za-z']+))?")
_PROJ = re.compile(r"\b(RP\d{3,5}(?:-FTW)?)\b", re.I)


# ── helpers ──────────────────────────────────────────────────────────────────────────────────────────────

def _fmt(d) -> str:
    try:
        return dt.date.fromisoformat(str(d)[:10]).strftime("%m/%d/%Y")
    except ValueError:
        return str(d or "")


def _po(rec: dict) -> str:
    return next((f.get("StringValue") or "" for f in rec.get("CustomField") or [] if f.get("DefinitionId") == "1"), "")


def _text(rec: dict) -> str:
    bits = [rec.get("PrivateNote") or ""]
    bits += [ln.get("Description") or "" for ln in rec.get("Line") or [] if ln.get("DetailType") == "SalesItemLineDetail"]
    return " | ".join(dict.fromkeys(b.strip() for b in bits if b.strip()))


def addr_keys(text: str) -> set:
    """('5425', 'glenscape') for every house number + street word in the text."""
    out = set()
    for num, w1, _w2 in _ADDR.findall(text or ""):
        w = w1.lower()
        if w in _SUFFIX:
            continue
        out.add((num, w))
    return out


def is_ftw(text: str) -> str:
    """The flatwork word found in the backcharge text, or '' (= slab)."""
    t = (text or "").lower()
    return next((w for w in FTW_WORDS if re.search(rf"\b{re.escape(w)}\b", t)), "")


def _links(pay: dict) -> list:
    return [(lt["TxnType"], str(lt["TxnId"]), round(float(ln["Amount"]), 2))
            for ln in pay.get("Line") or [] for lt in ln.get("LinkedTxn") or []]


def _work_dir(customer: str) -> Path:
    base = paths.analysis_dir(f"{customer} credit memos").parent
    found = sorted(base.glob(f"{customer} credit memos (*)"), key=lambda p: p.stat().st_mtime)
    return found[-1] if found else paths.analysis_dir(f"{customer} credit memos")


# ── the customer's projects, keyed by street address ─────────────────────────────────────────────────────

def project_index(con, customer: str) -> tuple:
    custs = mirror.load("Customer", con=con)
    parents = [c for c in custs if not c.get("ParentRef") and c.get("Active")
               and customer.upper() in (c.get("FullyQualifiedName") or "").upper()]
    if len(parents) != 1:
        raise SystemExit(f"Expected one active parent customer named like '{customer}', found "
                         f"{[c['FullyQualifiedName'] for c in parents]}")
    parent = parents[0]
    subs = {c["Id"]: c for c in custs if (c.get("ParentRef") or {}).get("value") == parent["Id"]}
    by_proj: dict = {}
    for c in subs.values():
        m = _PROJ.match(c.get("DisplayName") or "")
        if m:
            by_proj.setdefault(m.group(1).upper(), []).append(c)
    addr: dict = {}

    def add(text, proj):
        base = proj.upper().replace("-FTW", "")
        for k in addr_keys(text):
            addr.setdefault(k, set()).add(base)

    for c in subs.values():
        m = _PROJ.match(c.get("DisplayName") or "")
        if m:
            add((c.get("DisplayName") or "")[len(m.group(1)):], m.group(1))
    for inv in mirror.load("Invoice", con=con):
        c = subs.get(inv["CustomerRef"]["value"])
        m = c and _PROJ.match(c.get("DisplayName") or "")
        if m:
            add(" ".join([(inv.get("ShipAddr") or {}).get("Line1") or "", inv.get("PrivateNote") or ""]), m.group(1))
    return parent, subs, by_proj, addr


def resolve_project(by_proj: dict, proj: str) -> dict | None:
    """The one customer record for an exact project # (RP7186 and RP7186-FTW are different jobs)."""
    cs = by_proj.get((proj or "").upper(), [])
    act = [c for c in cs if c.get("Active")]
    pick = act or cs
    return pick[0] if len(pick) == 1 else None


def candidates(addr: dict, by_proj: dict, text: str) -> list:
    out = set()
    for k in addr_keys(text):
        out |= addr.get(k, set())
    res = []
    for b in sorted(out):
        res.append(b + (" (+FTW)" if f"{b}-FTW" in by_proj else ""))
    return res


# ── BEFORE ───────────────────────────────────────────────────────────────────────────────────────────────

def build_before(customer: str, since: str) -> dict:
    con = mirror.connect()
    try:
        parent, subs, by_proj, addr = project_index(con, customer)
        cms = [c for c in mirror.load("CreditMemo", "txn_date >= ?", (since,), con=con)
               if c["CustomerRef"]["value"] == parent["Id"]]
        moved = [c for c in mirror.load("CreditMemo", "txn_date >= ?", (since,), con=con)
                 if c["CustomerRef"]["value"] in subs]
        ids = {c["Id"] for c in cms} | {c["Id"] for c in moved}
        pays: dict = {}
        on: dict = {}
        for p in mirror.load("Payment", con=con):
            for t, i, a in _links(p):
                if t == "CreditMemo" and i in ids:
                    pays[p["Id"]] = p
                    on.setdefault(i, []).append({"payment_id": p["Id"], "applied": a})
        invs = {i: mirror.get("Invoice", i, con=con) for p in pays.values()
                for t, i, _ in _links(p) if t == "Invoice"}
    finally:
        con.close()

    def row(c):
        txt = _text(c)
        return {"id": c["Id"], "sync_token": c["SyncToken"], "date": c["TxnDate"], "ref": c.get("DocNumber") or "",
                "po": _po(c), "customer": c["CustomerRef"].get("name"), "customer_id": c["CustomerRef"]["value"],
                "amount": round(float(c["TotalAmt"]), 2), "open": round(float(c.get("RemainingCredit") or c.get("Balance") or 0), 2),
                "memo": txt, "applied_on": on.get(c["Id"], []), "candidates": candidates(addr, by_proj, txt)}

    state = {
        "customer": customer, "since": since, "built": dt.datetime.now().isoformat(timespec="seconds"),
        "parent": {"id": parent["Id"], "name": parent["FullyQualifiedName"]},
        "credits": [row(c) for c in sorted(cms, key=lambda c: (c["TxnDate"], c["Id"]))],
        "already_moved": [row(c) for c in sorted(moved, key=lambda c: (c["TxnDate"], c["Id"]))],
        "payments": pays,
        "invoices": {i: {"doc": (v or {}).get("DocNumber"), "balance": (v or {}).get("Balance"),
                         "customer": ((v or {}).get("CustomerRef") or {}).get("name")} for i, v in invs.items()},
    }
    return state


# ── PDFs ─────────────────────────────────────────────────────────────────────────────────────────────────

_BK = re.compile(r"\b(BKGP\s*-?\s*\d{3,6}[A-Z]?)\b", re.I)
_MONEY = re.compile(r"\$?\s*(\d{1,3}(?:,\d{3})*\.\d{2})\b")


_PONUM = re.compile(r"^\s*(?:[\d-]+\s+)?(\d{5})-(\d{5}[A-Z]?)\s*$", re.M)   # Hyphen BRIX PO: 00263-00683
_DOUBLED = re.compile(r"((\d)\2)+//((\d)\4)+//((\d)\6){4}")                  # bold header date: 99//1144//22002266


def _undouble(s: str) -> str:
    return s[::2]


def read_pdf(path: Path) -> dict:
    """A Hyphen BRIX backcharge PO: number, PO #, date, job address, description, amount."""
    import pdfplumber
    with pdfplumber.open(str(path)) as pdf:
        text = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    m = _BK.search(path.stem) or _BK.search(text)
    num = re.sub(r"[\s-]", "", m.group(1)).upper() if m else ""
    pm = _PONUM.search(text)
    if not num and pm:
        num = f"BKGP{pm.group(2)}"
    # PO # = the FULL BRIX number, 00240-05853 (owner 2026-10-16; credits keyed by hand before that carry 5 digits).
    po = f"{pm.group(1)}-{pm.group(2)}" if pm else (re.sub(r"\D", "", num.replace("BKGP", "", 1)) if num else "")
    d = ""
    dm = _DOUBLED.search(text)
    if dm:
        mo, dd, y = _undouble(dm.group(0).replace("////", "//")).split("/")
        try:
            d = dt.date(int(y), int(mo), int(dd)).isoformat()
        except ValueError:
            d = ""
    am = re.search(r"PROFICIENT CONCRETE LLC\s+\d+\s+(.+)", text)
    address = am.group(1).strip() if am else ""
    dsc = re.search(r"Quantity Item #.*?\n(.*?)\nTotal", text, re.S)
    desc = re.sub(r"\s+", " ", re.sub(r"[\d.,]+-", "", dsc.group(1))).strip() if dsc else ""
    desc = re.sub(r"^1\.00\s+|\s+1\.0000", "", desc)
    tm = re.search(r"^Total\s+([\d,]+\.\d{2})-?", text, re.M)
    amounts = [float(tm.group(1).replace(",", ""))] if tm else \
        sorted({float(x.replace(",", "")) for x in _MONEY.findall(text)}, reverse=True)
    return {"file": path.name, "number": num, "po": po, "amounts": amounts, "date": d,
            "address": address, "desc": desc, "text": text}


def attached_pdfs(access: str, cid: str, cm_id: str, dest: Path) -> list:
    """Download every PDF attached to the credit memo in QuickBooks into its own folder."""
    import urllib.request
    q = ("SELECT * FROM Attachable WHERE AttachableRef.EntityRef.Type = 'CreditMemo' "
         f"AND AttachableRef.EntityRef.value = '{cm_id}'")
    recs = qbo_api._api_get(f"/v3/company/{cid}/query", access, {"query": q}).get("QueryResponse", {}).get("Attachable", [])
    out = []
    for a in recs:
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", a.get("FileName") or f"{a['Id']}.pdf")
        url = a.get("TempDownloadUri") or ""
        if not name.lower().endswith(".pdf") or not url.lower().startswith("https://"):
            continue
        dest.mkdir(parents=True, exist_ok=True)
        f = dest / name
        with urllib.request.urlopen(url, timeout=60) as r:      # noqa: S310 - https only, checked above
            f.write_bytes(r.read())
        out.append(f)
    return out


def match(state: dict, work: Path) -> list:
    """Read the backcharge PDF attached to each credit memo and fill its AFTER columns."""
    con = mirror.connect()
    try:
        _parent, _subs, by_proj, addr = project_index(con, state["customer"])
    finally:
        con.close()
    access, cid = qbo_api.load_credentials()
    notes = []
    for cr in state["credits"]:
        keys = addr_keys(cr["memo"])
        files = attached_pdfs(access, cid, cr["id"], work / "PDFs" / cr["id"])
        pdfs = [x for x in (read_pdf(f) for f in files) if x["number"]] or [read_pdf(f) for f in files]
        cr["after"] = {}
        if len(pdfs) != 1:
            cr["after"]["note"] = "no PDF attached yet" if not pdfs else \
                "several PDFs attached: " + ", ".join(x["file"] for x in pdfs)
            continue
        p = pdfs[0]
        bases = set()
        for k in (addr_keys(p["address"]) or keys):
            bases |= addr.get(k, set())
        ftw = is_ftw(p["desc"] or p["text"])
        note = []
        proj = ""
        if len(bases) == 1:
            base = next(iter(bases))
            proj = f"{base}-FTW" if ftw else base
            if not resolve_project(by_proj, proj):
                note.append(f"{proj} not found as a project")
                proj = ""
        else:
            note.append("address matches " + (", ".join(sorted(bases)) or "no project"))
        if not p["number"]:
            note.append("no BKGP # on the PDF")
        if cr["amount"] not in p["amounts"]:
            note.append(f"PDF does not show ${cr['amount']:,.2f}")
        cr["after"] = {"file": p["file"], "ref": p["number"], "po": p["po"], "project": proj,
                       "why": (f"flatwork ('{ftw}')" if ftw else "slab (no flatwork word)") + f" - {p['address']}: {p['desc'][:80]}",
                       "pdf_date": p["date"], "note": "; ".join(note)}
    return notes


# ── workbook ─────────────────────────────────────────────────────────────────────────────────────────────

COLS = [  # (header, section, width)
    ("Credit Id", "B", 10), ("Date", "B", 11), ("Ref # now", "B", 12), ("PO # now", "B", 10),
    ("Customer now", "B", 18), ("Amount", "B", 11), ("Open now", "B", 10), ("Memo", "B", 48),
    ("Applied on payment", "B", 30), ("Address matches", "B", 22),
    ("PDF file", "A", 24), ("New Ref #", "A", 13), ("New PO #", "A", 10), ("New project", "A", 13),
    ("Slab / FTW", "A", 22), ("PDF date", "A", 11), ("Check", "A", 30),
    ("Approve", "I", 9),
]
FILL = {"B": "BDD7EE", "A": "C6E0B4", "I": "FFE699"}


def _payment_label(state: dict, cr: dict) -> str:
    out = []
    for a in cr["applied_on"]:
        p = state["payments"].get(a["payment_id"]) or {}
        out.append(f"#{p.get('PaymentRefNum') or '?'} {_fmt(p.get('TxnDate'))} ${a['applied']:,.2f}")
    return "; ".join(out) or "NOT FOUND on any payment"


def write_book(state: dict, path: Path, approvals: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter
    from shared import xlsx_guard, xlsx_verify

    wb = Workbook()
    ws = wb.active
    ws.title = "Credits"
    for j, (h, sec, w) in enumerate(COLS, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color=FILL[sec], end_color=FILL[sec], fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, cr in enumerate(state["credits"], 2):
        a = cr.get("after") or {}
        vals = [int(cr["id"]), dt.date.fromisoformat(cr["date"]), cr["ref"], cr["po"], cr["customer"], cr["amount"],
                cr["open"], cr["memo"], _payment_label(state, cr), ", ".join(cr["candidates"]),
                a.get("file", ""), a.get("ref", ""), a.get("po", ""), a.get("project", ""), a.get("why", ""),
                dt.date.fromisoformat(a["pdf_date"]) if a.get("pdf_date") else None, a.get("note", ""),
                approvals.get(cr["id"], "")]
        for j, v in enumerate(vals, 1):
            c = xlsx_guard.put(ws, i, j, v)
            c.alignment = Alignment(vertical="top", wrap_text=j in (8, 9, 17))
            if j in (6, 7):
                c.number_format = '#,##0.00'
            elif j in (2, 16):
                c.number_format = "mm/dd/yyyy"
    last = max(2, len(state["credits"]) + 1)
    t = Table(displayName="Credits", ref=f"A1:{get_column_letter(len(COLS))}{last}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    red = PatternFill(start_color="F4A6A6", end_color="F4A6A6", fill_type="solid")
    ws.conditional_formatting.add(f"I2:I{last}", FormulaRule(formula=['LEFT($I2,3)="NOT"'], fill=red))
    ws.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=['LEN($Q2)>0'], fill=red))
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Payments before")
    hdr = ["Payment Id", "Check #", "Date", "Payment total", "Unapplied", "Applies", "Doc #", "Customer",
           "Amount", "Open on that doc (before)"]
    for j, h in enumerate(hdr, 1):
        c = ws2.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color=FILL["B"], end_color=FILL["B"], fill_type="solid")
    r = 2
    cm_by_id = {c["id"]: c for c in state["credits"] + state["already_moved"]}
    for pid, p in sorted(state["payments"].items(), key=lambda kv: kv[1].get("TxnDate") or ""):
        for typ, tid, amt in _links(p):
            if typ == "Invoice":
                inv = state["invoices"].get(tid) or {}
                doc, cust, opn = inv.get("doc"), inv.get("customer"), inv.get("balance")
            else:
                cm = cm_by_id.get(tid) or {}
                doc, cust, opn = cm.get("ref", tid), cm.get("customer"), cm.get("open")
            vals = [int(pid), p.get("PaymentRefNum") or "", dt.date.fromisoformat(p["TxnDate"]),
                    float(p["TotalAmt"]), float(p.get("UnappliedAmt") or 0), typ, doc or "", cust or "", amt, opn]
            for j, v in enumerate(vals, 1):
                c = xlsx_guard.put(ws2, r, j, v)
                if j in (4, 5, 9, 10):
                    c.number_format = '#,##0.00'
                elif j == 3:
                    c.number_format = "mm/dd/yyyy"
            r += 1
    t2 = Table(displayName="PaymentsBefore", ref=f"A1:J{max(2, r - 1)}")
    t2.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws2.add_table(t2)
    ws2.freeze_panes = "A2"
    for j, w in enumerate([11, 10, 11, 13, 10, 11, 13, 22, 11, 14], 1):
        ws2.column_dimensions[get_column_letter(j)].width = w

    ws3 = wb.create_sheet("Already on a project")
    hdr3 = ["Credit Id", "Date", "Ref #", "PO #", "Project", "Amount", "Open now", "Applied on payment", "Status"]
    for j, h in enumerate(hdr3, 1):
        c = ws3.cell(1, j, h)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color=FILL["A"], end_color=FILL["A"], fill_type="solid")
    for i, cr in enumerate(state["already_moved"], 2):
        status = "OPEN - payment lost this credit, run restore" if cr["open"] > EPS and cr["applied_on"] else "ok"
        for j, v in enumerate([int(cr["id"]), dt.date.fromisoformat(cr["date"]), cr["ref"], cr["po"], cr["customer"],
                               cr["amount"], cr["open"], _payment_label(state, cr), status], 1):
            c = xlsx_guard.put(ws3, i, j, v)
            if j in (6, 7):
                c.number_format = '#,##0.00'
            elif j == 2:
                c.number_format = "mm/dd/yyyy"
    l3 = max(2, len(state["already_moved"]) + 1)
    t3 = Table(displayName="AlreadyMoved", ref=f"A1:I{l3}")
    t3.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws3.add_table(t3)
    ws3.conditional_formatting.add(f"I2:I{l3}", FormulaRule(formula=['LEFT($I2,4)="OPEN"'], fill=red))
    ws3.freeze_panes = "A2"
    for j, w in enumerate([10, 11, 13, 9, 13, 10, 10, 30, 40], 1):
        ws3.column_dimensions[get_column_letter(j)].width = w

    wb.save(path)
    xlsx_verify.assert_clean(path)


def read_approvals(path: Path) -> dict:
    """{credit id: row dict} as the owner left the Credits sheet (his edits to the AFTER columns win)."""
    if not path.exists():
        return {}
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Credits"]
    hdr = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, row))
        if d.get("Credit Id"):
            out[str(d["Credit Id"])] = d
    return out


# ── QuickBooks writes ────────────────────────────────────────────────────────────────────────────────────

def _get(access, cid, ent, rid) -> dict:
    return qbo_api._api_get(f"/v3/company/{cid}/{ent.lower()}/{rid}", access)[ent]


def _post(access, cid, ent, body) -> dict:
    r = requests.post(f"{qbo_api.API_BASE}/v3/company/{cid}/{ent.lower()}",
                      headers={"Authorization": f"Bearer {qbo_api._CURRENT_ACCESS or access}",
                               "Content-Type": "application/json", "Accept": "application/json"},
                      params={"minorversion": qbo_api.MINOR_VERSION}, json=body, timeout=120)
    if r.status_code != 200:
        try:
            e = r.json()["Fault"]["Error"][0]
            msg = f"{e.get('Message')}: {e.get('Detail')}"
        except Exception:                                   # noqa: BLE001
            msg = r.text[:400]
        raise RuntimeError(f"QuickBooks refused the {ent} update ({r.status_code}): {msg}")
    return r.json()[ent]


def _same_links(a: dict, b: dict) -> bool:
    return sorted(_links(a)) == sorted(_links(b))


def plan_groups(state: dict, rows: dict, only: str | None) -> tuple:
    """Approved credits grouped by the payment they sit on: [(payment_id, [credit rows])], problems."""
    con = mirror.connect()
    try:
        _parent, _subs, by_proj, _addr = project_index(con, state["customer"])
    finally:
        con.close()
    groups: dict = {}
    problems = []
    for cr in state["credits"]:
        r = rows.get(cr["id"]) or {}
        if str(r.get("Approve") or "").strip().upper() != "Y":
            continue
        proj, ref, po = (str(r.get(k) or "").strip() for k in ("New project", "New Ref #", "New PO #"))
        pdate = r.get("PDF date")
        pdate = pdate.date().isoformat() if hasattr(pdate, "date") else str(pdate or "")[:10]
        target = resolve_project(by_proj, proj)
        if not (proj and ref and po and pdate and target):
            problems.append(f"{cr['ref']} (credit {cr['id']}): needs New project / New Ref # / New PO # / PDF date"
                            + ("" if target or not proj else f" - {proj} is not one project"))
            continue
        if len(cr["applied_on"]) > 1:
            problems.append(f"{cr['ref']}: applied on {len(cr['applied_on'])} payments - do by hand")
            continue
        if not cr["applied_on"] and cr["open"] < cr["amount"] - EPS:
            problems.append(f"{cr['ref']}: used somewhere other than a payment - skipped until found")
            continue
        pid = cr["applied_on"][0]["payment_id"] if cr["applied_on"] else ""
        if only and pid != only:
            continue
        groups.setdefault(pid, []).append({**cr, "new": {"project": proj, "project_id": target["Id"],
                                                          "ref": ref, "po": po, "date": pdate}})
    return sorted(groups.items()), problems


def apply(state: dict, rows: dict, commit: bool, only: str | None) -> int:
    groups, problems = plan_groups(state, rows, only)
    for p in problems:
        print("  SKIP", p)
    if not groups:
        print("Nothing approved to change.")
        return 0
    access, cid = qbo_api.load_credentials()
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    failed = 0
    for pid, crs in groups:
        snap = state["payments"].get(pid)
        live_pay = _get(access, cid, "Payment", pid) if pid else None
        live_cms = {c["id"]: _get(access, cid, "CreditMemo", c["id"]) for c in crs}
        stale = []
        if live_pay is not None and not _same_links(live_pay, snap):
            stale.append("the payment changed since the BEFORE snapshot")
        for c in crs:
            lc = live_cms[c["id"]]
            if lc["CustomerRef"]["value"] != state["parent"]["id"] or abs(float(lc["TotalAmt"]) - c["amount"]) > EPS:
                stale.append(f"credit {c['ref']} changed since the BEFORE snapshot")
        head = (f"Payment #{(snap or {}).get('PaymentRefNum')} {_fmt((snap or {}).get('TxnDate'))} "
                f"${float((snap or {}).get('TotalAmt') or 0):,.2f}" if pid else "No payment")
        print(f"\n{head}  ({len(crs)} credit{'s' if len(crs) != 1 else ''})")
        for c in crs:
            n = c["new"]
            print(f"   {c['ref']:>10} ${c['amount']:>9,.2f}  -> {n['project']:<11} Ref {n['ref']:<12} PO {n['po']:<7} "
                  f"date {_fmt(c['date'])} -> {_fmt(n['date'])}")
        if stale:
            print("   STOP:", "; ".join(stale), "- run `before` again")
            failed += 1
            continue
        if not commit:
            continue
        expect = _open_before(access, cid, live_pay)
        (LOG_DIR / f"{stamp}-payment-{pid or 'none'}.json").write_text(json.dumps(
            {"payment": live_pay, "credits": live_cms, "open_before": {f"{k[0]}:{k[1]}": v for k, v in expect.items()},
             "plan": crs}, indent=1, default=str))
        err = None
        try:
            for c in crs:
                lc = live_cms[c["id"]]
                done = _post(access, cid, "CreditMemo", {
                    "Id": lc["Id"], "SyncToken": lc["SyncToken"], "sparse": True,
                    "CustomerRef": {"value": c["new"]["project_id"]}, "DocNumber": c["new"]["ref"],
                    "TxnDate": c["new"]["date"], "CustomField": [{**PO_FIELD, "StringValue": c["new"]["po"]}]})
                got = (done["CustomerRef"]["value"], done.get("DocNumber"), done.get("TxnDate"), _po(done))
                want = (c["new"]["project_id"], c["new"]["ref"], c["new"]["date"], c["new"]["po"])
                if got != want:
                    raise RuntimeError(f"{c['ref']}: QuickBooks saved {got}, expected {want}")
                print(f"   moved {c['ref']} -> {c['new']['project']} as {done.get('DocNumber')}")
        except Exception as e:                              # noqa: BLE001
            err = e
        if live_pay is not None:
            # Always put the payment back, even when a credit write failed half way.
            res = payment_reapply.commit(_post, pid, target=live_pay, creds=(access, cid), expect_open=expect)
            print("   payment restored and verified" if res["ok"] else "   PAYMENT NOT BACK: " + res.get("error", ""))
            if not res["ok"]:
                failed += 1
                print(f"   the BEFORE copy is in {LOG_DIR} - stopping here")
                return 1
        if err:
            print(f"   STOP: {err}")
            return 1
    if not commit:
        print("\nDry run - nothing written. Add --commit to write.")
    return 1 if failed else 0


def move(credit_id: str, project: str, customer: str, commit: bool) -> int:
    """Move one credit memo (already on a project) to another project, then put back its payment."""
    con = mirror.connect()
    try:
        _parent, _subs, by_proj, _addr = project_index(con, customer)
        pids = [p["Id"] for p in mirror.load("Payment", con=con)
                if any(t == "CreditMemo" and i == credit_id for t, i, _ in _links(p))]
    finally:
        con.close()
    target = resolve_project(by_proj, project)
    if not target:
        print(f"{project} is not one {customer} project.")
        return 1
    access, cid = qbo_api.load_credentials()
    cm = _get(access, cid, "CreditMemo", credit_id)
    pays = [p for p in (_get(access, cid, "Payment", i) for i in pids)
            if any(t == "CreditMemo" and i == credit_id for t, i, _ in _links(p))]
    print(f"{cm.get('DocNumber')} ${float(cm['TotalAmt']):,.2f}: {cm['CustomerRef'].get('name')} -> {project}")
    for p in pays:
        print(f"   on payment #{p.get('PaymentRefNum')} {_fmt(p.get('TxnDate'))} ${float(p['TotalAmt']):,.2f}")
    if len(pays) > 1:
        print("On more than one payment - do by hand.")
        return 1
    if not commit:
        print("Dry run - add --commit to write.")
        return 0
    pay = pays[0] if pays else None
    expect = _open_before(access, cid, pay)
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    (LOG_DIR / f"{dt.datetime.now():%Y%m%d-%H%M%S}-move-{credit_id}.json").write_text(
        json.dumps({"credit": cm, "payment": pay, "open_before": {f"{k[0]}:{k[1]}": v for k, v in expect.items()}},
                   indent=1, default=str))
    err = None
    try:
        done = _post(access, cid, "CreditMemo", {"Id": cm["Id"], "SyncToken": cm["SyncToken"], "sparse": True,
                                                 "CustomerRef": {"value": target["Id"]}})
        print(f"   moved to {done['CustomerRef'].get('name')}")
    except Exception as e:                                  # noqa: BLE001
        err = e
    if pay:
        res = payment_reapply.commit(_post, pay["Id"], target=pay, creds=(access, cid), expect_open=expect)
        print("   payment restored and verified" if res["ok"] else "   PAYMENT NOT BACK: " + res.get("error", ""))
        if not res["ok"]:
            return 1
    if err:
        print(f"   STOP: {err}")
        return 1
    return 0


def _open_before(access, cid, pay: dict | None) -> dict:
    """{(type, id): open} for every document on the payment, read before a change - what each must show after."""
    return {(t, i): payment_reapply.open_of(_get(access, cid, t, i)) for t, i, _ in _links(pay or {})
            if t in payment_reapply.DOCS}


def restore(payment_id: str, commit: bool) -> int:
    """Re-save a payment exactly as the mirror holds it (shared/payment_reapply - the ledger's engine)."""
    res = payment_reapply.dry_run(payment_id)
    print(f"Payment #{res['ref']} {_fmt(res['date'])} {res['customer']} ${res['total']:,.2f}")
    for b in res["back"]:
        print(f"   put back {b['type']:<10} {b['doc']:<12} ${b['amount']:>10,.2f}   open {b['open_now']:,.2f} -> {b['open_after']:,.2f}")
    for w in res["why"]:
        print("  ", w)
    if not commit or not res["ok"]:
        print("Dry run - add --commit to write." if res["ok"] else "Nothing written.")
        return 0 if res["ok"] or res["why"] == [payment_reapply.NOTHING] else 1
    done = payment_reapply.commit(_post, payment_id, res["sync_token"], res["n"], res["amount"])
    print("Restored and verified." if done["ok"] else "NOT BACK: " + done.get("error", ""))
    return 0 if done["ok"] else 1


# ── main ─────────────────────────────────────────────────────────────────────────────────────────────────

def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("step", choices=["before", "match", "apply", "restore", "move"])
    ap.add_argument("payment", nargs="?")
    ap.add_argument("project", nargs="?")
    ap.add_argument("--customer", default="Grand Homes")
    ap.add_argument("--since", default="2026-01-01")
    ap.add_argument("--commit", action="store_true")
    ap.add_argument("--only")
    a = ap.parse_args()

    if a.step == "move":
        if not (a.payment and a.project):
            ap.error("move needs a credit memo id and a project #")
        return move(a.payment, a.project, a.customer, a.commit)

    if a.step == "restore":
        if not a.payment:
            ap.error("restore needs a payment id")
        return restore(a.payment, a.commit)

    if a.step == "before":
        work = paths.analysis_dir(f"{a.customer} credit memos")
    else:
        work = _work_dir(a.customer)
    book = work / f"{a.customer} credit memos.xlsx"
    state_file = work / "state.json"

    if a.step == "before":
        state = build_before(a.customer, a.since)
        state_file.write_text(json.dumps(state, indent=1, default=str))
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        (LOG_DIR / f"{dt.datetime.now():%Y%m%d-%H%M%S}-before.json").write_text(json.dumps(state, indent=1, default=str))
        write_book(state, book, {k: v.get("Approve") or "" for k, v in read_approvals(book).items()})
        n_pay = len({x["payment_id"] for c in state["credits"] for x in c["applied_on"]})
        print(f"{len(state['credits'])} credits on {state['parent']['name']} since {_fmt(a.since)}, "
              f"${sum(c['amount'] for c in state['credits']):,.2f}, on {n_pay} payments")
        print(book)
        return 0

    if not state_file.exists():
        print("Run `before` first.")
        return 1
    state = json.loads(state_file.read_text())

    if a.step == "match":
        notes = match(state, work)
        state_file.write_text(json.dumps(state, indent=1, default=str))
        write_book(state, book, {k: v.get("Approve") or "" for k, v in read_approvals(book).items()})
        ok = sum(1 for c in state["credits"] if (c.get("after") or {}).get("project") and not c["after"].get("note"))
        print(f"{ok} of {len(state['credits'])} credits read cleanly from their attached PDF")
        for n in notes:
            print("  ", n)
        print(book)
        return 0

    return apply(state, read_approvals(book), a.commit, a.only)


if __name__ == "__main__":
    sys.exit(main())
