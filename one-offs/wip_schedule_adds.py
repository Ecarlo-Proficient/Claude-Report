"""wip_schedule_adds.py — the "Schedule Adds" tab for a month-end WIP report.

Jobs that were live on the RP crew SCHEDULE of the report date but have NO line
in that month's WIP report. Same shape the 12-31-25 report already uses
('Schedule Adds 12-31-25'): the WIP's own 18-column header, the added rows
banded yellow, and only the two QBO columns filled — COSTS TO DATE (H) and
BILLED TO DATE (M), dated on/before the cutoff. A job with no QBO project at
all keeps those two cells BLANK ON RED, exactly as the 12-31-25 tab does.

The schedule of 3-31-26 has no project-# column (col A is the superintendent,
col B the address), so each address was resolved to its job # against the
General Lista, the Residential job folders' takeoff filenames, and the
12-31-25 report — then confirmed against the QBO customer list. The resolved
map is baked in below (ROWS) with its evidence, because it took a human pass:
re-deriving it silently on every run would be guesswork, not a lookup.

Writes ONE new sheet into the existing report — never a v2 of the file — and
verifies the saved workbook with shared.xlsx_verify before returning.

Usage:
  python3 one-offs/wip_schedule_adds.py --dry-run
  python3 one-offs/wip_schedule_adds.py --commit
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))
from shared import qbo_api                                    # noqa: E402
from shared.xlsx_verify import assert_clean                   # noqa: E402

REPORT = Path(os.getenv(
    "WIP_HISTORY_DIR",
    str(Path.home() / "Library/CloudStorage/OneDrive-ProficientConcrete,LLC"
        / "Company Files - WIP Report" / "WIP History"))) / "WIP REPORT - 3-31-26 FINAL.xlsx"
SRC_SHEET = "Sheet1"
NEW_SHEET = "Schedule Adds 3-31-26"
CUTOFF = "2026-03-31"
TITLE_A2 = "REPORT DATE: MAR 31, 2026 (schedule adds)"

# Styling lifted from 'Schedule Adds 12-31-25' — do not restyle (WIP formatting
# is frozen; this tab's look is the 12-31-25 tab's look).
FONT = "Tahoma"
SIZE = 8
MONEY = '"$"#,##0_);[Red]("$"#,##0)'
BAND = PatternFill("solid", fgColor="FFFF99")     # the added-row band, cols A-G
NOQBO = PatternFill("solid", fgColor="FFC7CE")    # H/M when the job isn't in QBO
HDR = PatternFill("solid", fgColor="D9D9D9")
_thin = Side(style="thin", color="000000")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
NCOL = 18
COL_COSTS, COL_BILLED = 8, 13                     # H, M — the only QBO columns

# (project #, address as it should read, why it's here)
#   evidence for each project # is in the third field; a None project # means
#   the address has no QBO project under the number the paperwork uses.
ROWS = [
    # ── flatwork lines ──
    ("RP4362-FTW", "5021 Bryan Street",
     "WIP carries RP4362 (slab) only; schedule ran leadwalks + sidewalks"),
    ("RP5542-FTW", "2400 Country Club Road",
     "schedule SET UP FTW; base RP5542 is deleted in QBO"),
    # ── slab lines ──
    ("RP5542", "2400 Country Club Road",
     "schedule Wreck Footings; QBO project marked (deleted)"),
    ("RP6201", "1123 Moonlight Bay Drive",
     "schedule Wreck Pad; RP6201 = this address per the 2024 General Lista"),
    ("RP6378", "2107 Cartwright Drive",
     "schedule SET FORMS; RP6378 = this address per the 2024 General Lista"),
    ("RP6047", "303 Lone Tree Lane",
     "schedule POUR PIERS + SET FORMS; no RP6047 in QBO"),
    ("RP6516", "624 West Berkeley Circle",
     "schedule POUR PIERS + SET FORMS; job # from the folder takeoff; no RP6516 in QBO"),
    ("RP6586", "9330 Peninsula Drive",
     "schedule READY FOR MAKE UP; QBO names the address on the customer"),
    ("RP6764", "4933 Mangold Circle",
     "schedule TRENCH; QBO names the address on the customer"),
    ("RP6766", "1921 Mount McKinley Place",
     "schedule Wreck Slab; folder takeoff RP6766_1921 MOUNT McKINLEY PLACE"),
    ("RP7005", "7130 South Janmar Drive",
     "WIP carries RP7005-FTW only; schedule ran Wreck Slab"),
    ("RP7027", "12004 Garden Grove Drive",
     "schedule SET FORMS; QBO parent Proficient Custom Homes matches"),
    ("RP7235", "6407 Glenrose Court",
     "schedule Wreck pier and beam backfill; QBO invoice 34167 matches that scope"),
    ("RP7314", "14956-14978 Little Bluestem Lane",
     "schedule Wreck Slab + READY FOR MAKE UP; General Lista 2026"),
    ("RP7332", "7701 Black Elk Court",
     "WIP carries RP7332-FTW only; schedule ran Wreck Slab + make up"),
    ("RP7371", "5116 Clotho Road",
     "WIP carries RP7371-FTW only; schedule ran Wreck Slab"),
    ("RP7419", "4230 Valley Ridge Road",
     "schedule READY FOR MAKE UP; General Lista 2026"),
    ("RP7440", "6333 Olive Branch Avenue",
     "schedule SET FORMS; folder takeoff RP7440_6333 OLIVE BRANCH AVENUE"),
]


def qbo_totals():
    """{project # -> (costs, billed)} as of the cutoff. Missing key = not in QBO.

    Matches the QBO DisplayName EXACTLY on the project token, so RP5542 never
    picks up RP5542-FTW (a prefix match does, and did).
    """
    access, cid = qbo_api.load_credentials()
    customers = qbo_api.query_all(access, cid, "Customer")
    customers += qbo_api.query_all(access, cid, "Customer", "Active = false")
    by_proj = {}
    for c in customers:
        disp = (c.get("DisplayName") or "").strip().upper()
        m = re.match(r"^((?:RP|CP)\d{3,4}(?:-FTW)?)(?:\s|$)", disp)
        if m:
            by_proj.setdefault(m.group(1), c)
    out = {}
    for proj, _addr, _why in ROWS:
        c = by_proj.get(proj)
        if not c:
            continue
        t = qbo_api.extract_pl_totals(qbo_api.fetch_project_pl(
            access, cid, c["Id"], "2019-01-01", CUTOFF))
        billed = round(float(t.get("income", 0) or 0), 2)
        costs = round(float((t.get("cogs", 0) or 0) + (t.get("expenses", 0) or 0)), 2)
        out[proj] = (costs, billed)
    return out


def build(commit: bool):
    if not REPORT.exists():
        sys.exit(f"  ✗ report not found: {REPORT}")
    lock = REPORT.parent / f"~${REPORT.name}"
    if lock.exists():
        sys.exit(f"  ✗ {REPORT.name} is open in Excel — close it first")

    print(f"  report   : {REPORT.name}")
    print(f"  new sheet: {NEW_SHEET}  ({len(ROWS)} lines)")
    totals = qbo_totals()
    missing = [p for p, _a, _w in ROWS if p not in totals]
    print(f"  QBO      : {len(totals)}/{len(ROWS)} projects found"
          + (f" · absent: {', '.join(missing)}" if missing else ""))

    wb = load_workbook(REPORT)          # keep formulas — no data_only
    src = wb[SRC_SHEET]
    if NEW_SHEET in wb.sheetnames:
        del wb[NEW_SHEET]               # idempotent: rebuild, never v2
    ws = wb.create_sheet(NEW_SHEET)

    # header block copied from the report's own sheet so the tab reads identically
    ws.cell(1, 1, src.cell(1, 1).value).font = Font(name=FONT, size=SIZE, bold=True)
    ws.cell(2, 1, TITLE_A2).font = Font(name=FONT, size=SIZE, bold=True)
    for col in (COL_COSTS, COL_BILLED):
        c = ws.cell(2, col, "QBO")
        c.font = Font(name="Aptos Narrow", size=12)
        c.alignment = Alignment(horizontal="center")
    for col in range(1, NCOL + 1):
        c = ws.cell(3, col, src.cell(3, col).value)
        c.font = Font(name=FONT, size=SIZE, bold=True)
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    for letter, dim in src.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[letter].width = dim.width

    row = 4
    for proj, addr, _why in ROWS:
        for col in range(1, NCOL + 1):
            c = ws.cell(row, col)
            c.font = Font(name=FONT, size=SIZE)
            c.border = BORDER
            if col <= 7:
                c.fill = BAND
            if col >= 4:
                c.number_format = MONEY
        a = ws.cell(row, 1, proj)
        a.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.cell(row, 2, addr).alignment = Alignment(wrap_text=True)
        ws.cell(row, 7).number_format = "0.00%"
        ws.cell(row, 10).number_format = "0.00%"
        got = totals.get(proj)
        if got:
            ws.cell(row, COL_COSTS, got[0])
            ws.cell(row, COL_BILLED, got[1])
        else:
            ws.cell(row, COL_COSTS).fill = NOQBO
            ws.cell(row, COL_BILLED).fill = NOQBO
        row += 1

    last = row - 1
    for col in range(4, NCOL + 1):
        c = ws.cell(row, col)
        if col in (7, 10):
            num, den = (6, 4) if col == 7 else (COL_COSTS, 5)
            c.value = (f"=IF({ws.cell(row, den).coordinate}=0,\"\","
                       f"{ws.cell(row, num).coordinate}/{ws.cell(row, den).coordinate})")
            c.number_format = "0.00%"
        else:
            L = ws.cell(row, col).column_letter
            c.value = f"=SUBTOTAL(109,{L}4:{L}{last})"
            c.number_format = MONEY
        c.font = Font(name=FONT, size=SIZE, bold=True)
        c.fill = HDR
        c.border = BORDER
    for col in range(1, 4):
        ws.cell(row, col).fill = HDR
        ws.cell(row, col).border = BORDER

    if not commit:
        print("  DRY RUN — nothing written. Re-run with --commit.")
        return
    wb.save(REPORT)
    assert_clean(REPORT)                # LAST step, always
    print(f"  ✓ wrote '{NEW_SHEET}' into {REPORT.name} and verified it clean")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true")
    g.add_argument("--commit", action="store_true")
    a = ap.parse_args()
    build(a.commit)
