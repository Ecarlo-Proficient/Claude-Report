#!/usr/bin/env python3
"""
rp_review.py - the weekly RP review with the ops manager: every RP line on the WIP, WHERE
each number was grabbed from, and a SNAPSHOT of the source spreadsheet cut around the row
the number sits in (max 5 rows, that row highlighted) - the way the owner wants to see it
(2026-09-15: "instead of only showing the source file, create a snapshot of the excel of
only where it's at, max 5 rows with its row highlighted").

WHAT IT READS (all read-only)
  * the WIP master's RP tab ('Test - RP', or 'WIP-RP' if the office renamed it) - the four
    inputs per line, the «base» columns (what the last sync wrote), the QBO deep links
  * the owner's RP WIP file ('RP WIP', its CONTRACT FILE / ETC FILE paths, STATUS, ACTION;
    the 'Removed log' sheet = the finished lines and why they left)
  * the job folder on Common: the proposal PDF / takeoff bid sheet / cost sheet the number
    came from - located by VALUE (the exact amount) so the snapshot is honest: when the
    number is not in any document, the page says so
  * JobTread approved proposals (shared/jobtread, Touch ID once)
  * the daily crew schedules (shared/schedule_index for the last date; that day's sheet
    for the row snapshot)

WHAT IT WRITES
  ~/Library/Application Support/Proficient/rp-review/rp_review.json  (ACB_RP_REVIEW_DIR)
  plus a snapshot cache beside it (keyed by file path + mtime + value) so a rebuild costs
  seconds when nothing changed. The dashboard serves the JSON; the ops-manager marks live in
  the ledger DB (rp_review_mark + rp_review_mark_log) - the mark helpers are at the bottom.

  python3 ledger/rp_review.py [--no-jobtread] [--no-snapshots] [--limit N]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sqlite3
import subprocess
import sys
import time
from pathlib import Path
from typing import Dict, List, Optional

_REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(_REPO))

from openpyxl import load_workbook                     # noqa: E402
from shared import paths, schedule as SCH, schedule_index as SI, jobtread, takeoff_etc as TK  # noqa: E402

OUT_DIR = Path(os.environ.get("ACB_RP_REVIEW_DIR",
                              Path.home() / "Library" / "Application Support" / "Proficient" / "rp-review"))
LEDGER_DB = Path(os.environ.get("ACB_LEDGER_DB",
                                Path.home() / "Library" / "Application Support" / "Proficient" / "ledger.sqlite3"))
RP_TABS = ("Test - RP", "WIP-RP")
MAX_ROWS = 5
MAX_COLS = 8
STALE_DAYS = 14
SNAP_VER = 2                 # bump when the snapshot shape changes - old cache entries are then ignored

_log_lines: List[str] = []
_WB: Dict[str, object] = {}          # in-run memo: data_only workbooks by path (the master is opened once, not 94 times)


def _wb(path: Path):
    k = str(path)
    if k not in _WB:
        _WB[k] = load_workbook(path, data_only=True)
    return _WB[k]


def log(msg: str) -> None:
    print(msg, flush=True)
    _log_lines.append(msg)


# ───────────────────────────── helpers ─────────────────────────────
def _num(v) -> Optional[float]:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _fmt(v) -> str:
    """A cell for the snapshot: numbers with commas, dates mm/dd/yyyy, text as is."""
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%m/%d/%Y")
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v - round(v)) > 0.004 else f"{v:,.0f}"
    s = str(v)
    return s if len(s) <= 60 else s[:57] + "..."


def _master_path() -> Path:
    p = paths.get_path("WIP_EXCEL_PATH", paths.onedrive_base() / "Company Files - WIP Report" / "WIP - MASTER.xlsx")
    if not p.exists():
        alt = p.with_name("WIP - MASTER new.xlsx")
        if alt.exists():
            return alt
    return p


def _rp_file() -> Path:
    return paths.get_path("RP_WIP_FILE", paths.onedrive_base() / "RP WIP TO FIX_Final.xlsx")


# ───────────────────────────── snapshots ─────────────────────────────
class SnapCache:
    def __init__(self, path: Path):
        self.path = path
        try:
            self.d = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            self.d = {}
        self.hits = 0

    def key(self, f: Path, extra: str) -> Optional[str]:
        try:
            return f"v{SNAP_VER}|{f}|{int(f.stat().st_mtime)}|{extra}"
        except OSError:
            return None

    def get(self, k):
        v = self.d.get(k) if k else None
        if v is not None:
            self.hits += 1
        return v

    def put(self, k, v):
        if k:
            self.d[k] = v

    def save(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self.path.write_text(json.dumps(self.d), encoding="utf-8")


def _col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def _snap_rows(ws, r0: int, cols: List[int], hi_row: int, title: str, path: Path, anchor: str) -> dict:
    """MAX_ROWS rows centred on hi_row (clipped to the sheet), the given columns."""
    lo = max(1, hi_row - 2)
    hi = min(ws.max_row, lo + MAX_ROWS - 1)
    lo = max(1, hi - MAX_ROWS + 1)
    rows = []
    for r in range(lo, hi + 1):
        rows.append({"n": r, "hi": r == hi_row,
                     "cells": [_fmt(ws.cell(r, c).value) for c in cols]})
    return {"kind": "xlsx", "file": path.name, "path": str(path), "sheet": ws.title, "anchor": anchor,
            "cols": [_col_letter(c) for c in cols], "rows": rows}


def _snap_rows_multi(ws, hi_rows: List[int], cols: List[int], path: Path, anchor: str) -> dict:
    """Up to MAX_ROWS rows: every highlighted row plus the row above each, oldest first."""
    want = sorted({r for h in hi_rows for r in (h - 1, h) if r >= 1})
    while len(want) > MAX_ROWS:
        extra = [r for r in want if r not in hi_rows]
        if not extra:
            break
        want.remove(extra[0])
    rows = [{"n": r, "hi": r in hi_rows, "cells": [_fmt(ws.cell(r, c).value) for c in cols]} for r in want]
    return {"kind": "xlsx", "file": path.name, "path": str(path), "sheet": ws.title, "anchor": anchor,
            "cols": [_col_letter(c) for c in cols], "rows": rows}


def _snap_cost_bands(path: Path, value: float, cache: SnapCache) -> Optional[dict]:
    """The ETC is the SUM of the cost sheet's band subtotals (SL + PR for a slab, FW for
    flatwork - shared/takeoff_etc). Find the band set that adds up to `value` and cut those
    subtotal rows (each highlighted)."""
    k = cache.key(path, f"bands={value:.2f}")
    hit = cache.get(k)
    if hit is not None:
        return hit or None
    out = None
    try:
        wb = _wb(path)
        for n in wb.sheetnames:
            if "COST" not in n.upper():
                continue
            bands = TK._cost_sheet_totals(wb[n])
            subs = {b: v for b, v in bands.items() if v.get("sub") is not None and v.get("cell")}
            combos = [("SL", "PR"), ("SL",), ("FW",), ("SL", "PR", "FW"), ("PR",)]
            for combo in combos:
                if all(b in subs for b in combo) and abs(sum(subs[b]["sub"] for b in combo) - value) < 1.0:
                    ws = wb[n]
                    rows = [ws[subs[b]["cell"]].row for b in combo]
                    col = ws[subs[combo[0]]["cell"]].column
                    out = _snap_rows_multi(ws, rows, _window_cols(col, ws.max_column), path, " + ".join(subs[b]["cell"] for b in combo))
                    out["note"] = " + ".join(f"{b} {subs[b]['sub']:,.2f}" for b in combo) + f" = {value:,.2f}"
                    break
            if out:
                break
    except Exception as e:                                        # noqa: BLE001
        log(f"    ! cost bands {path.name}: {type(e).__name__}")
    cache.put(k, out or {})
    return out


def _window_cols(anchor_col: int, max_col: int) -> List[int]:
    lo = max(1, anchor_col - (MAX_COLS - 2))
    hi = min(max_col, lo + MAX_COLS - 1)
    lo = max(1, hi - MAX_COLS + 1)
    return list(range(lo, hi + 1))


def _snap_xlsx_value(path: Path, value: float, cache: SnapCache, sheets_like: Optional[str] = None) -> Optional[dict]:
    """Find `value` in the workbook (bid/cost sheets first) and cut the rows around it."""
    k = cache.key(path, f"v={value:.2f}|{sheets_like or ''}")
    hit = cache.get(k)
    if hit is not None:
        return hit or None
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:                                        # noqa: BLE001
        cache.put(k, {}); log(f"    ! cannot open {path.name}: {type(e).__name__}"); return None
    names = wb.sheetnames
    order = sorted(names, key=lambda n: (0 if (sheets_like and sheets_like in n.upper()) else
                                         1 if ("BID" in n.upper() or "COST" in n.upper()) else 2))
    out = None
    for n in order:
        ws = wb[n]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
            for c in row:
                v = c.value
                if isinstance(v, (int, float)) and not isinstance(v, bool) and abs(v - value) < 0.6:
                    ws2 = _wb(path)[n]                            # random access for the window
                    cols = _window_cols(c.column, ws2.max_column)
                    out = _snap_rows(ws2, c.row, cols, c.row, n, path, c.coordinate)
                    break
            if out:
                break
        if out:
            break
    wb.close()
    cache.put(k, out or {})
    return out


def _snap_xlsx_row(path: Path, sheet: str, row: int, cols: List[int], cache: SnapCache, anchor_col: int) -> Optional[dict]:
    k = cache.key(path, f"r={sheet}|{row}|{','.join(map(str, cols))}")
    hit = cache.get(k)
    if hit is not None:
        return hit or None
    try:
        ws = _wb(path)[sheet]
    except Exception as e:                                        # noqa: BLE001
        cache.put(k, {}); log(f"    ! cannot open {path.name}: {type(e).__name__}"); return None
    out = _snap_rows(ws, row, cols, row, sheet, path, f"{_col_letter(anchor_col)}{row}")
    cache.put(k, out)
    return out


def _snap_pdf(path: Path, value: Optional[float], cache: SnapCache) -> Optional[dict]:
    """pdftotext -layout; the 5 lines around the line carrying the amount (or SUB TOTAL)."""
    k = cache.key(path, f"pdf={value if value is None else round(value, 2)}")
    hit = cache.get(k)
    if hit is not None:
        return hit or None
    try:
        txt = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True,   # noqa: S603,S607
                             text=True, timeout=60).stdout
    except (OSError, subprocess.SubprocessError):
        cache.put(k, {}); return None
    lines = [ln.rstrip() for ln in txt.splitlines()]
    needle = f"{value:,.2f}" if value is not None else None
    idx = None
    for i, ln in enumerate(lines):
        if needle and needle in ln:
            idx = i
    if idx is None:
        for i, ln in enumerate(lines):
            if "SUB TOTAL" in ln.upper():
                idx = i
    if idx is None:
        cache.put(k, {}); return None
    nonempty = [i for i, ln in enumerate(lines) if ln.strip()]
    pos = nonempty.index(idx) if idx in nonempty else 0
    pick = nonempty[max(0, pos - 2): pos + 3][:MAX_ROWS]
    found = bool(needle and needle in lines[idx])
    m = re.search(r"\$\s*([\d,]+\.\d{2})", lines[idx])
    out = {"kind": "pdf", "file": path.name, "path": str(path), "anchor": f"line {idx + 1}",
           "rows": [{"n": i + 1, "hi": i == idx, "cells": [re.sub(r"\s{3,}", "   ", lines[i]).strip()]} for i in pick],
           "cols": ["text"], "found": found,
           "note": "" if found else (f"this proposal says {m.group(1)}, not {needle}" if (m and needle) else "the amount is not in this proposal")}
    cache.put(k, out)
    return out


def _find_in_folder(folder: Path, value: float, cache: SnapCache, want: str) -> Optional[dict]:
    """The document in the job folder that carries `value`: proposal PDFs and takeoff
    workbooks at the top level only (the ARCHIVE subfolder is skipped on purpose)."""
    if not folder.exists() or not folder.is_dir():
        return None
    try:
        items = sorted(folder.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True)
    except OSError:
        return None
    pdfs = [p for p in items if p.suffix.lower() == ".pdf" and re.search(r"proposal|bid|flatwork|ftw", p.name, re.I)]
    xls = [p for p in items if p.suffix.lower() in (".xlsm", ".xlsx") and not p.name.startswith("~$")]
    seq = (pdfs + xls) if want == "contract" else (xls + pdfs)
    for p in seq[:12]:
        if p.suffix.lower() == ".pdf":
            s = _snap_pdf(p, value, cache)
            if s and s.get("found"):
                return s
        else:
            s = _snap_xlsx_value(p, value, cache, "COST" if want == "etc" else "BID")
            if not s and want == "etc":
                s = _snap_cost_bands(p, value, cache)
            if s:
                return s
    return None


# ───────────────────────────── readers ─────────────────────────────
def read_master(path: Path) -> dict:
    wb = load_workbook(path)                       # formulas kept: we read INPUT cells + base columns
    tab = next((t for t in RP_TABS if t in wb.sheetnames), None)
    if not tab:
        raise SystemExit(f"no RP tab ({' / '.join(RP_TABS)}) in {path.name}")
    ws = wb[tab]
    hr = next(r for r in range(1, 16) if ws.cell(r, 2).value == "PROJECT #")
    hdr = {str(ws.cell(hr, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    rows = {}
    for r in range(hr + 1, ws.max_row + 1):
        p = str(ws.cell(r, 2).value or "").strip().upper()
        if not p.startswith("RP"):
            continue
        d = {k: ws.cell(r, c).value for k, c in hdr.items() if k}
        d["_row"] = r
        for k in ("COSTS TO DATE", "BILLED TO DATE"):
            c = ws.cell(r, hdr[k]) if k in hdr else None
            d[k + " LINK"] = c.hyperlink.target if (c is not None and c.hyperlink) else None
        rows[p] = d
    m = re.search(r"REPORT DATE:\s*(.+)", str(ws.cell(2, 2).value or ""))
    return {"tab": tab, "hdr": hdr, "rows": rows, "report_date": m.group(1).strip() if m else "",
            "path": path, "renamed": tab != RP_TABS[0]}


def read_rp_file(path: Path) -> dict:
    wb = load_workbook(path)
    ws = wb["RP WIP"]
    rows, band = {}, ""
    for r in range(3, ws.max_row + 1):
        v = str(ws.cell(r, 1).value or "").strip()
        if v.startswith("⚠") or v.upper().startswith("FTW BACKLOG"):
            band = v; continue
        j = v.upper()
        if not j.startswith("RP"):
            continue
        rows.setdefault(j, {"row": r, "address": ws.cell(r, 2).value, "builder": ws.cell(r, 3).value,
                            "contract": _num(ws.cell(r, 4).value), "etc": _num(ws.cell(r, 5).value),
                            "status": ws.cell(r, 12).value, "action": ws.cell(r, 13).value,
                            "contract_file": ws.cell(r, 15).value, "etc_file": ws.cell(r, 16).value,
                            "band": band})
    removed = []
    if "Removed log" in wb.sheetnames:
        lg = wb["Removed log"]
        hdr = [str(lg.cell(1, c).value or "").strip() for c in range(1, lg.max_column + 1)]
        for r in range(2, lg.max_row + 1):
            if not lg.cell(r, 2).value:
                continue
            d = {hdr[c - 1]: lg.cell(r, c).value for c in range(1, len(hdr) + 1) if hdr[c - 1]}
            d["_row"] = r
            removed.append(d)
    return {"rows": rows, "removed": removed, "path": path}


def _schedule_snap(job: str, day: Optional[dt.date], files: Dict[dt.date, Path], cache: SnapCache) -> Optional[dict]:
    f = files.get(day) if day else None
    if not f:
        return None
    k = cache.key(f, f"sched={job}")
    hit = cache.get(k)
    if hit is not None:
        return hit or None
    out = None
    try:
        wb = load_workbook(f, data_only=True, read_only=True)
        name = next((n for n in wb.sheetnames if "MAIN" in n.upper()), wb.sheetnames[0])
        ws = wb[name]
        base = job.replace("-FTW", "")
        want_ftw = job.endswith("-FTW")
        hits = []
        section = ""
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 600)):
            first = str(row[0].value or "").strip() if row else ""
            texts = [str(c.value or "").strip().upper() for c in row[:12]]
            if first and all(not str(c.value or "").strip() for c in row[1:6]):
                section = first.upper()
            if any(t == base or t.startswith(base + " ") for t in texts):
                hits.append((row[0].row, "FLATWORK" in section))
        pick = next((r for r, ftw in hits if ftw == want_ftw), hits[0][0] if hits else None)
        wb.close()
        if pick:
            ws2 = _wb(f)[name]
            # columns B..G = project #, address, city, builder, description, date (the super and
            # crew columns A / H are people - left out on purpose)
            out = _snap_rows(ws2, pick, list(range(2, min(ws2.max_column, 7) + 1)), pick, name, f, f"B{pick}")
    except Exception as e:                                        # noqa: BLE001
        log(f"    ! schedule {f.name}: {type(e).__name__}")
    cache.put(k, out or {})
    return out


# ───────────────────────────── the build ─────────────────────────────
def classify_contract(line: str, d: dict, r: dict, typed: bool):
    v = _num(d.get("ORIGINAL CONTRACT"))
    cf = str(r.get("contract_file") or "")
    if typed and v:
        hint = "equals the invoice amount" if abs(v - (_num(d.get("BILLED TO DATE")) or -1)) < 1 else "no paper named"
        return "typed", "Typed on the master", f"blank at the last sync, typed in by hand since ({hint}); OneDrive version history names who"
    if not v:
        return "blank", "BLANK", "no contract anywhere - the estimator enters it in JobTread"
    if cf.lower().endswith(".pdf"):
        return "pdf", "Proposal PDF", Path(cf).name
    if cf.lower().endswith((".xlsm", ".xlsx")):
        return "takeoff", "Takeoff bid sheet", Path(cf).name
    if cf:
        return "folder", "Job folder", "read from the proposal / takeoff in the folder"
    if "FTW BACKLOG" in str(r.get("band") or "").upper():
        return "gl", "General List", "flatwork price typed in the General List"
    return "master", "WIP master", "carried from an earlier master; no document on file"


def classify_etc(line: str, d: dict, r: dict, typed: bool):
    v = _num(d.get("ORIGINAL ESTIMATED COST"))
    ef = str(r.get("etc_file") or "")
    if typed and v:
        return "typed", "Typed on the master", "blank at the last sync, typed in by hand since, no cost sheet named"
    if not v:
        return "blank", "BLANK", "no budget anywhere - the estimator enters it in JobTread"
    if ef.lower().endswith((".xlsm", ".xlsx")):
        return "takeoff", "Takeoff cost sheet", Path(ef).name
    if ef:
        return "folder", "Job folder", "read from the takeoff cost sheet in the folder"
    if "manual" in str(r.get("action") or "").lower():
        return "manual", "Typed by hand", str(r.get("action"))[:90]
    return "master", "WIP master", "carried from an earlier master; no cost sheet on file"


def build(no_jobtread: bool = False, no_snapshots: bool = False, limit: int = 0) -> dict:
    t0 = time.time()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cache = SnapCache(OUT_DIR / "snapshot_cache.json")
    mp, fp = _master_path(), _rp_file()
    log(f"master: {mp.name}   RP file: {fp.name}")
    M = read_master(mp)
    F = read_rp_file(fp)
    log(f"RP tab '{M['tab']}' ({'RENAMED - the sync writes Test - RP' if M['renamed'] else 'as expected'}): "
        f"{len(M['rows'])} lines · RP file: {len(F['rows'])} active, {len(F['removed'])} removed")

    # last schedule date per job (cached index; refresh when the share is mounted)
    try:
        seen = SI.last_seen(refresh=Path(str(SCH.SCHEDULE_DIR)).exists())
    except Exception as e:                                        # noqa: BLE001
        log(f"  schedule index unavailable: {type(e).__name__}"); seen = {}
    sched_files = {d: f for d, f in SCH.all_schedule_files()} if Path(str(SCH.SCHEDULE_DIR)).exists() else {}

    jt: Dict[str, list] = {}
    if not no_jobtread:
        nums = {ln.replace("-FTW", "") for ln in M["rows"]} | {str(x.get("JOB #") or "").upper().replace("-FTW", "") for x in F["removed"]}
        jt = jobtread.approved_proposals(nums, log=log)
        log(f"JobTread: {len(jt)} of {len(nums)} jobs have an approved proposal")

    as_of = dt.date.today()
    current = []
    lines = list(M["rows"].items())[:limit] if limit else list(M["rows"].items())
    for i, (line, d) in enumerate(lines, 1):
        r = F["rows"].get(line, {})
        K, E = _num(d.get("ORIGINAL CONTRACT")), _num(d.get("ORIGINAL ESTIMATED COST"))
        bK, bE = _num(d.get("«base» base_contract")), _num(d.get("«base» base_etc"))
        typed_k = (K or 0) != (bK or 0) and "«base» base_contract" in d
        typed_e = (E or 0) != (bE or 0) and "«base» base_etc" in d
        kc, kl, kd = classify_contract(line, d, r, typed_k)
        ec, el, ed = classify_etc(line, d, r, typed_e)
        base = line.replace("-FTW", "")
        docs = jt.get(base) or []
        jt_price, jt_cost, jt_date = (docs[-1] if docs else (None, None, ""))
        last = seen.get(line) or seen.get(base)
        stale = bool(last and (as_of - last).days > STALE_DAYS)
        costs, billed = _num(d.get("COSTS TO DATE")), _num(d.get("BILLED TO DATE"))
        flags = []
        if typed_k or typed_e:
            flags.append("contract / ETC typed on the master after the last sync, no source named - which paper?")
        fk, fe = _num(r.get("contract")), _num(r.get("etc"))
        if K and fk and abs(K - fk) > 1:
            flags.append(f"the RP file says contract {fk:,.0f}")
        if E and fe and abs(E - fe) > 1:
            flags.append(f"the RP file says ETC {fe:,.0f}")
        if jt_price and K and not line.endswith("-FTW") and abs(jt_price - K) > 1:
            flags.append(f"JobTread price {jt_price:,.0f} differs from the WIP")
        if not docs:
            flags.append("not in JobTread")
        if stale:
            flags.append(f"not on the schedule since {last.strftime('%m/%d/%Y')}")
        if K and billed and billed >= K * 0.99 and not stale:
            flags.append("billed out, still on the schedule")
        if K and costs and E and costs > E:
            flags.append(f"costs {costs:,.0f} already over the ETC {E:,.0f}")

        snaps = {}
        if not no_snapshots:
            # where the CONTRACT sits
            src = r.get("contract_file")
            if kc == "pdf" and src and Path(src).exists():
                snaps["contract"] = _snap_pdf(Path(src), K, cache)
            elif kc == "takeoff" and src and Path(src).exists():
                snaps["contract"] = _snap_xlsx_value(Path(src), K, cache, "BID")
            elif kc == "folder" and src and K:
                snaps["contract"] = _find_in_folder(Path(src), K, cache, "contract")
            elif kc in ("typed", "master") and K:
                snaps["contract"] = None
            if kc in ("pdf", "takeoff", "folder") and K and not snaps.get("contract"):
                flags.append("contract amount not found in the folder's proposal or takeoff")
            elif snaps.get("contract") and snaps["contract"].get("kind") == "pdf" and not snaps["contract"].get("found"):
                flags.append(snaps["contract"].get("note") or "the proposal PDF says a different number")
            # where the ETC sits
            src = r.get("etc_file")
            if ec == "takeoff" and src and Path(src).exists() and E:
                snaps["etc"] = _snap_xlsx_value(Path(src), E, cache, "COST") or _snap_cost_bands(Path(src), E, cache)
            elif ec == "folder" and src and E:
                snaps["etc"] = _find_in_folder(Path(src), E, cache, "etc")
            if ec in ("takeoff", "folder") and E and not snaps.get("etc"):
                flags.append("ETC amount not found in the folder's cost sheet")
            # the master row, the RP file row, the schedule row
            hdr = M["hdr"]
            cols = [hdr[k] for k in ("PROJECT #", "PROJECT NAME", "ORIGINAL CONTRACT", "APPROVED COs",
                                     "ORIGINAL ESTIMATED COST", "CO COSTS", "COSTS TO DATE", "BILLED TO DATE") if k in hdr]
            snaps["master"] = _snap_xlsx_row(mp, M["tab"], d["_row"], cols, cache, hdr.get("ORIGINAL CONTRACT", 2))
            if r:
                snaps["rpfile"] = _snap_xlsx_row(fp, "RP WIP", r["row"], list(range(1, 9)), cache, 4)
            snaps["schedule"] = _schedule_snap(line, last, sched_files, cache)

        current.append({
            "line": line, "ftw": line.endswith("-FTW"), "name": d.get("PROJECT NAME"), "builder": d.get("BUILDER"),
            "type": d.get("TYPE"), "status": d.get("STATUS"), "category": d.get("CATEGORY"),
            "contract": K, "etc": E, "costs": costs, "billed": billed,
            "costs_link": d.get("COSTS TO DATE LINK"), "billed_link": d.get("BILLED TO DATE LINK"),
            "src": {"contract": {"kind": kc, "label": kl, "detail": kd, "file": r.get("contract_file")},
                    "etc": {"kind": ec, "label": el, "detail": ed, "file": r.get("etc_file")},
                    "costs": {"kind": "qbo", "label": "QuickBooks", "detail": f"project P&L, synced {M['report_date']}"},
                    "billed": {"kind": "qbo", "label": "QuickBooks", "detail": f"invoices on the project, synced {M['report_date']}"}},
            "jt": {"price": jt_price, "cost": jt_cost, "date": jt_date, "count": len(docs),
                   "scope_note": "the base job's proposal - may be the slab, not the flatwork" if (docs and line.endswith("-FTW")) else ""},
            "rp_status": r.get("status"), "rp_action": r.get("action"), "band": r.get("band"),
            "last_seen": last.isoformat() if last else None, "stale": stale,
            "flags": flags, "problem": bool(flags) or not K or not E,
            "snaps": {k: v for k, v in snaps.items() if v},
        })
        if i % 10 == 0:
            log(f"  {i}/{len(lines)} lines · {time.time() - t0:.0f}s · cache hits {cache.hits}")

    finished = []
    for x in F["removed"]:
        job = str(x.get("JOB #") or "").upper()
        base = job.replace("-FTW", "")
        docs = jt.get(base) or []
        last_day = x.get("LAST DAY ON SCHEDULE")
        ld = None
        if isinstance(last_day, (dt.date, dt.datetime)):
            ld = last_day if isinstance(last_day, dt.date) else last_day.date()
        else:
            m = re.match(r"(\d{2})/(\d{2})/(\d{4})", str(last_day or ""))
            if m:
                ld = dt.date(int(m.group(3)), int(m.group(1)), int(m.group(2)))
        snaps = {}
        if not no_snapshots:
            snaps["removed"] = _snap_xlsx_row(fp, "Removed log", x["_row"], [2, 3, 5, 6, 7, 8, 9, 10], cache, 7)
            snaps["schedule"] = _schedule_snap(job, ld, sched_files, cache)
        finished.append({
            "line": job, "ftw": job.endswith("-FTW"), "name": x.get("ADDRESS"), "builder": x.get("BUILDER"),
            "contract": _num(x.get("CONTRACT $")), "etc": _num(x.get("ETC $")),
            "billed": _num(x.get("BILLED TO DATE $ (QBO)")), "costs": _num(x.get("COSTS TO DATE $ (QBO)")),
            "last_invoice": _fmt(x.get("LAST INVOICE")), "last_day": _fmt(last_day), "last_stage": x.get("LAST STAGE"),
            "why": x.get("WHY REMOVED"), "source": x.get("SOURCE"), "removed_on": _fmt(x.get("REMOVED ON")),
            "approved_by": x.get("APPROVED BY"),
            "jt": {"price": docs[-1][0] if docs else None, "cost": docs[-1][1] if docs else None, "date": docs[-1][2] if docs else ""},
            "snaps": {k: v for k, v in snaps.items() if v},
        })

    cache.save()
    kinds = {}
    for c in current:
        kinds[c["src"]["contract"]["kind"]] = kinds.get(c["src"]["contract"]["kind"], 0) + 1
    out = {
        "built_at": dt.datetime.now().isoformat(timespec="seconds"),
        "master": {"file": mp.name, "tab": M["tab"], "renamed": M["renamed"], "report_date": M["report_date"],
                   "mtime": dt.datetime.fromtimestamp(mp.stat().st_mtime).isoformat(timespec="seconds")},
        "rp_file": {"file": fp.name, "mtime": dt.datetime.fromtimestamp(fp.stat().st_mtime).isoformat(timespec="seconds")},
        "jobtread": {"queried": not no_jobtread, "approved": len(jt)},
        "schedule_mounted": bool(sched_files),
        "counts": {"current": len(current), "finished": len(finished), "contract_kinds": kinds,
                   "problems": sum(1 for c in current if c["problem"]),
                   "in_jobtread": sum(1 for c in current if c["jt"]["count"]),
                   "typed": sum(1 for c in current if c["src"]["contract"]["kind"] == "typed" or c["src"]["etc"]["kind"] == "typed"),
                   "stale": sum(1 for c in current if c["stale"])},
        "current": current, "finished": finished,
        "log": _log_lines[-40:],
    }
    (OUT_DIR / "rp_review.json").write_text(json.dumps(out, default=str), encoding="utf-8")
    log(f"wrote {OUT_DIR / 'rp_review.json'} in {time.time() - t0:.0f}s (cache hits {cache.hits})")
    return out


def load() -> Optional[dict]:
    p = OUT_DIR / "rp_review.json"
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


# ───────────────────────────── marks (the ledger's write surface) ─────────────────────────────
# The owner and the ops manager never edit the prepared numbers (2026-09-15: "we need our own
# space for our own numbers with the numbers you have prepopulated in case it's right, we check
# or mark x with notes/changes and you read and make sense of it"). Each line keeps ONE
# standing answer (rp_review_mark) plus every answer ever given (rp_review_mark_log):
#   decision   current: confirmed | fix        finished: agree | keep       '' clears
#   mode       ops (the ops manager is here, deciding together) | me (the owner alone)
#   our_contract / our_etc   THEIR numbers (prepopulated with the prepared value)
#   contract_ok / etc_ok     1 = check, 0 = X, NULL = not answered
#   note       free text - what could not be settled, what changed
MODES = ("ops", "me")
DECISIONS = {"current": ("confirmed", "fix"), "finished": ("agree", "keep")}
_COLS = "project_no, kind, decision, mode, note, at, our_contract, our_etc, contract_ok, etc_ok"


def _ensure(con: sqlite3.Connection) -> None:
    con.execute("CREATE TABLE IF NOT EXISTS rp_review_mark ("
                "project_no TEXT NOT NULL, kind TEXT NOT NULL, decision TEXT NOT NULL, mode TEXT NOT NULL, "
                "note TEXT, at TEXT NOT NULL, our_contract REAL, our_etc REAL, contract_ok INTEGER, etc_ok INTEGER, "
                "PRIMARY KEY (project_no, kind))")
    con.execute("CREATE TABLE IF NOT EXISTS rp_review_mark_log ("
                "id INTEGER PRIMARY KEY AUTOINCREMENT, project_no TEXT NOT NULL, kind TEXT NOT NULL, "
                "decision TEXT NOT NULL, mode TEXT NOT NULL, note TEXT, at TEXT NOT NULL, "
                "our_contract REAL, our_etc REAL, contract_ok INTEGER, etc_ok INTEGER)")
    for t in ("rp_review_mark", "rp_review_mark_log"):        # older tables: add the answer columns
        have = {r[1] for r in con.execute(f"PRAGMA table_info({t})")}
        for c, typ in (("our_contract", "REAL"), ("our_etc", "REAL"), ("contract_ok", "INTEGER"), ("etc_ok", "INTEGER")):
            if c not in have:
                con.execute(f"ALTER TABLE {t} ADD COLUMN {c} {typ}")


def _rowdict(r) -> dict:
    return {"project_no": r[0], "kind": r[1], "decision": r[2], "mode": r[3], "note": r[4], "at": r[5],
            "our_contract": r[6], "our_etc": r[7], "contract_ok": r[8], "etc_ok": r[9]}


def read_marks() -> dict:
    """{'<project_no>|<kind>': {...}} - absent-safe (no DB / no table -> {})."""
    if not LEDGER_DB.exists():
        return {}
    try:
        con = sqlite3.connect(f"file:{LEDGER_DB}?mode=ro", uri=True)
        try:
            rows = con.execute(f"SELECT {_COLS} FROM rp_review_mark").fetchall()
        finally:
            con.close()
    except sqlite3.OperationalError:
        return {}
    return {f"{r[0]}|{r[1]}": _rowdict(r) for r in rows}


def read_mark_log(limit: int = 400) -> list:
    if not LEDGER_DB.exists():
        return []
    try:
        con = sqlite3.connect(f"file:{LEDGER_DB}?mode=ro", uri=True)
        try:
            rows = con.execute(f"SELECT {_COLS} FROM rp_review_mark_log ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
        finally:
            con.close()
    except sqlite3.OperationalError:
        return []
    return [_rowdict(r) for r in rows]


def set_mark(project_no: str, kind: str, decision: str, mode: str, note: str, now: str,
             our_contract=None, our_etc=None, contract_ok=None, etc_ok=None) -> None:
    """Write one answer (and append it to the log). decision '' clears the standing mark."""
    LEDGER_DB.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(str(LEDGER_DB))
    try:
        _ensure(con)
        vals = (project_no, kind, decision, mode, note, now, our_contract, our_etc, contract_ok, etc_ok)
        if decision:
            con.execute(f"INSERT INTO rp_review_mark ({_COLS}) VALUES (?,?,?,?,?,?,?,?,?,?) "
                        "ON CONFLICT(project_no, kind) DO UPDATE SET decision=excluded.decision, mode=excluded.mode, "
                        "note=excluded.note, at=excluded.at, our_contract=excluded.our_contract, our_etc=excluded.our_etc, "
                        "contract_ok=excluded.contract_ok, etc_ok=excluded.etc_ok", vals)
        else:
            con.execute("DELETE FROM rp_review_mark WHERE project_no=? AND kind=?", (project_no, kind))
        con.execute(f"INSERT INTO rp_review_mark_log ({_COLS}) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (project_no, kind, decision or "cleared", mode, note, now, our_contract, our_etc, contract_ok, etc_ok))
        con.commit()
    finally:
        con.close()


def answers_report() -> str:
    """The answers as text - what a later session reads to 'make sense of it': every line the
    owner / ops manager marked, their numbers where they differ from the prepared ones, the X's,
    the notes. Prepared values come from the last build."""
    data = load() or {}
    prepared = {c["line"]: c for c in data.get("current", [])}
    prepared.update({f["line"] + "|finished": f for f in data.get("finished", [])})
    marks = read_marks()
    if not marks:
        return "no answers yet"
    out = [f"RP review answers - {len(marks)} line(s) marked (build {data.get('built_at', '?')})", ""]
    for key in sorted(marks):
        m = marks[key]; pn, kind = key.split("|")
        p = prepared.get(pn if kind == "current" else pn + "|finished", {})
        who = "OPS MANAGER + owner" if m["mode"] == "ops" else "owner alone"
        bits = [f"{pn} [{kind}] {m['decision'].upper()} by {who} on {m['at'][:16].replace('T', ' ')}"]
        for fld, ok, ours in (("contract", m["contract_ok"], m["our_contract"]), ("etc", m["etc_ok"], m["our_etc"])):
            mine = p.get(fld)
            mark = {1: "check", 0: "X"}.get(ok, "-")
            diff = f" THEIR {fld.upper()} {ours:,.2f} vs prepared {mine:,.2f}" if (ours is not None and mine is not None and abs(ours - mine) > 0.5) else ""
            if ok is not None or diff:
                bits.append(f"    {fld}: {mark}{diff}")
        if m["note"]:
            bits.append(f"    note: {m['note']}")
        out.extend(bits)
    return "\n".join(out)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--no-jobtread", action="store_true", help="skip the JobTread pull")
    ap.add_argument("--no-snapshots", action="store_true", help="skip the source snapshots (fast)")
    ap.add_argument("--limit", type=int, default=0, help="first N lines only (testing)")
    ap.add_argument("--answers", action="store_true", help="print the owner's / ops manager's answers and stop")
    a = ap.parse_args()
    if a.answers:
        print(answers_report())
        return 0
    build(a.no_jobtread, a.no_snapshots, a.limit)
    return 0


if __name__ == "__main__":
    sys.exit(main())
