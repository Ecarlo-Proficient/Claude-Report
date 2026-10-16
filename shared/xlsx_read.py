"""shared/xlsx_read.py - THE way a tool opens a workbook it takes numbers or people's typing from.

`open_source(path)` separates two cases that used to look alike:
  - the file is not there        -> None. The caller decides what "no file" means (a first run, no budget yet).
  - the file is there but cannot be opened (a share hiccup, a file mid-copy, a corrupt save)
                                 -> SourceUnreadable, naming the file.

Returning an empty result for the second case is how a locked WIP master became "no contract" on a P&L, and
how an unreadable Bill Tracker became "no Notes" and was saved over. A file that exists is either read or the
run stops on it - never quietly treated as empty.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


class SourceUnreadable(Exception):
    """A workbook exists but could not be opened. Its message names the file."""

    def __init__(self, path: Path, cause: BaseException):
        self.path = Path(path)
        super().__init__(f"{self.path.name} exists but could not be opened "
                         f"({type(cause).__name__}: {cause}) - {self.path.parent}")


def open_source(path, **kwargs):
    """openpyxl.load_workbook(path, **kwargs), or None when the file does not exist.
    Raises SourceUnreadable when it exists and cannot be opened."""
    path = Path(path)
    if not path.exists():
        return None
    try:
        return load_workbook(str(path), **kwargs)
    except Exception as e:  # openpyxl surfaces zip, XML, format and OS errors alike for one unreadable file
        raise SourceUnreadable(path, e) from e
