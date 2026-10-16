"""The Revised Bills sheet in every vendor's statement workbook (statement_reconciler._revised_rows +
shared/revised_sheet, 10/12/2026): which bills the vendor must re-issue (tax added, amount differs, a PM's
REVISED memo), the clerk's status read back from the old sheet, the history kept, Stage forward only to
Done when the statement agrees, the sheet written and read back, the marked sheet's captions beside the
legend with the page from row 1. No network, no share."""
import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "statement-reconciler"))

import statement_reconciler as r  # noqa: E402
from shared import revised_sheet as rvs, xlsx_verify  # noqa: E402

TODAY = dt.date(2026, 10, 12)


def _row(cat, stmt_ref, qbo_ref, stmt_amt, qbo_amt, bill_id, date="2026-09-10"):
    return r.ReconRow(category=cat, date=date, ref=stmt_ref or qbo_ref, stmt_amount=stmt_amt, qbo_amount=qbo_amt, po="", address="",
                      stmt_ref=stmt_ref, qbo_ref=qbo_ref, stmt_date=date if stmt_ref else "", qbo_date=date if qbo_ref else "",
                      qbo_memo="", qbo_bill_id=bill_id)


def _bill(bid, doc, total, memo, created="2026-09-30", date="2026-09-10"):
    return r.QboBill(bill_id=bid, doc_number=doc, txn_date=date, open_balance=total, total_amount=total, memo=memo, created=created)


def test_rows_are_the_bills_the_vendor_must_reissue():
    rows = [_row("VENDOR_TAX_VIOLATION", "400621", "400621", 8768.25, 8100.0, "11"),
            _row("CLERK_AMOUNT_MISMATCH", "401386", "401386", 6200.0, 6083.65, "12"),
            _row("MATCHED", "399126", "399126", 4752.18, 4752.18, "13")]
    bills = [_bill("11", "400621", 8100.0, "RP7622 - x\n\nREVISED - TAX EXEMPT JOB"),
             _bill("12", "401386", 6083.65, "RP7110 - y\n\nREVISED - CORRECTED ADDRESS"),
             _bill("13", "399126", 4752.18, "RP1 - matched, no revision"),
             _bill("14", "402000", 500.0, "RP2 - z\n\nREVISED - NO DOWELS ORDERED"),          # a PM revision not on the statement
             _bill("15", "R1", 50.0, "REVISED - CP861 - 4231 STATE HWY 161", created="2026-03-01")]   # paper-era line-1: not a follow-up
    out = r._revised_rows(rows, bills, {}, TODAY, originals={"14": {"original": 600.0, "revised_on": "2026-10-01"}})
    by = {x["bill_id"]: x for x in out}
    assert set(by) == {"11", "12", "14"}
    assert by["11"]["reason"] == "Tax-exempt job - please remove the sales tax" and by["11"]["amount_billed"] == 8768.25 and by["11"]["corrected_amount"] == 8100.0
    assert by["12"]["reason"] == "Corrected address" and by["12"]["invoice"] == "401386"
    assert by["14"]["amount_billed"] == 600.0 and by["14"]["corrected_amount"] == 500.0 and by["14"]["revised_on"] == "2026-10-01" and by["14"]["reason"] == "No dowels ordered"
    assert all(x["stage"] == "To send" for x in out) and by["14"]["standing"]["text"] == "the clerk has not told the vendor yet - 11 days"
    assert by["11"]["revised_on"] == TODAY.isoformat() and by["11"]["standing"]["text"] == "the clerk has not told the vendor yet - 0 days"   # first seen today


def test_the_clerks_status_comes_back_history_is_kept_and_done_moves_forward_only():
    rows = [_row("VENDOR_TAX_VIOLATION", "400621", "400621", 8768.25, 8100.0, "11"),
            _row("MATCHED", "401386", "401386", 6083.65, 6083.65, "12")]            # 12 was mismatched last run, re-issued since
    bills = [_bill("11", "400621", 8100.0, "x"), _bill("12", "401386", 6083.65, "y")]
    prior = {"11": {"key": "11", "stage": "sent to vendor", "sent_on": "2026-10-03", "replied_on": "", "received_on": "", "entered_on": "", "notes": "emailed"},
             "12": {"key": "12", "invoice": "401386", "invoice_date": "2026-09-18", "amount_billed": 6200.0, "corrected_amount": 6083.65, "reason": "Corrected address",
                    "stage": "Entered in QBO", "sent_on": "2026-09-25", "replied_on": "2026-09-28", "received_on": "2026-10-01", "entered_on": "2026-10-02", "notes": ""},
             "99": {"key": "99", "invoice": "OLD1", "invoice_date": "2026-06-01", "amount_billed": 10.0, "corrected_amount": 9.0, "reason": "old", "stage": "Dropped"}}
    out = r._revised_rows(rows, bills, prior, TODAY)
    by = {x["bill_id"]: x for x in out}
    assert by["11"]["stage"] == "Sent to vendor" and by["11"]["sent_on"] == "2026-10-03" and by["11"]["notes"] == "emailed"
    assert by["11"]["standing"]["text"] == "waiting on the vendor's reply - 9 days"
    assert by["12"]["stage"] == "Done - statement agrees" and by["12"]["entered_on"] == "2026-10-02" and by["12"]["amount_billed"] == 6200.0   # history row, now matched
    assert by["99"]["stage"] == "Dropped" and by["99"]["invoice"] == "OLD1"              # the history stays
    assert [x["bill_id"] for x in out][0] == "11" and out[-1]["standing"]["final"]        # open first, finished last


def test_sheet_round_trip_keeps_what_the_clerk_typed(tmp_path):
    from openpyxl import Workbook, load_workbook
    path = tmp_path / "Reconciliation - V - as of 10-02-2026.xlsx"
    wb = Workbook(); wb.active.title = "Summary"; wb.active["A1"] = "x"
    rows = [{"bill_id": "11", "invoice": "400621", "invoice_date": "2026-09-10", "amount_billed": 8768.25, "corrected_amount": 8100.0, "reason": "Tax-exempt job - please remove the sales tax", "stage": "To send", "revised_on": "2026-10-01"},
            {"bill_id": "14", "invoice": "402000", "invoice_date": "2026-09-12", "amount_billed": None, "corrected_amount": 500.0, "reason": "No dowels ordered", "stage": "To send"}]
    assert rvs.write_sheet(wb, rows) == 2
    wb.save(path); xlsx_verify.assert_clean(path)
    wb2 = load_workbook(path)
    assert wb2.sheetnames == ["Summary", rvs.SHEET] and wb2[rvs.SHEET].column_dimensions["L"].hidden and wb2[rvs.SHEET].column_dimensions["M"].hidden   # _Key + _Since hidden
    ws = wb2[rvs.SHEET]; hdr = [c.value for c in ws[1]]
    assert hdr[:5] == ["Invoice #", "Invoice date", "Amount billed", "Corrected amount", "Reason"] and hdr[5] == "Stage" and hdr[-2:] == ["_Key", "_Since"]
    ws.cell(row=2, column=hdr.index("Stage") + 1, value="Vendor replied")
    ws.cell(row=2, column=hdr.index("Sent on") + 1, value=dt.date(2026, 10, 3))
    ws.cell(row=2, column=hdr.index("Replied on") + 1, value="10/09/2026")
    ws.cell(row=2, column=hdr.index("Notes") + 1, value="they will re-issue")
    wb2.save(path)
    prior = rvs.harvest(path)
    assert prior["11"]["stage"] == "Vendor replied" and prior["11"]["sent_on"] == "2026-10-03" and prior["11"]["replied_on"] == "2026-10-09" and prior["11"]["notes"] == "they will re-issue"
    assert prior["14"]["stage"] == "To send" and prior["14"]["amount_billed"] is None and prior["14"]["corrected_amount"] == 500.0
    # the next run: the rows rebuilt, her columns back
    out = r._revised_rows([_row("VENDOR_TAX_VIOLATION", "400621", "400621", 8768.25, 8100.0, "11")], [_bill("11", "400621", 8100.0, "x")], prior, TODAY)
    wb3 = load_workbook(path); rvs.write_sheet(wb3, out); wb3.save(path); xlsx_verify.assert_clean(path)
    back = rvs.harvest(path)
    assert back["11"]["stage"] == "Vendor replied" and back["11"]["notes"] == "they will re-issue" and back["14"]["invoice"] == "402000"   # 14 kept as history
    assert back["11"]["since"] == "2026-10-01" and next(x for x in out if x["bill_id"] == "11")["revised_on"] == "2026-10-01"   # the first-seen day survives the rebuild
    assert wb3[rvs.SHEET].data_validations.dataValidation[0].formula1 == f'"{rvs.STAGE_LIST}"'


def test_marked_sheet_starts_at_the_top_with_the_caption_beside_the_legend(tmp_path):
    from PIL import Image
    from openpyxl import Workbook
    import statement_markup as sm
    res = sm.Result(pages=[Image.new("RGB", (400, 300), "white")], located=3, total=3, lines=[1, 2, 3])
    wb = Workbook(); wb.active.title = "Summary"
    tmp = r._write_marked_sheet(wb, [("Statement V 10-02-2026.pdf", res)])
    ws = wb["Statement (marked)"]
    assert ws["A1"].value is None and ws["A3"].value is None
    assert ws._images[0].anchor == "A1"                                                            # the page from A1
    cap = ws.cell(row=1, column=r.LEGEND_COL).value
    assert cap.startswith("Statement V 10-02-2026.pdf - 3 of 3 bill rows found and coloured") and cap.endswith("all accounted for")
    assert ws.cell(row=3, column=r.LEGEND_COL).value == "Colour key"
    import shutil; shutil.rmtree(tmp, ignore_errors=True)


def test_stage_words_and_standing():
    assert rvs.stage_of(" sent to vendor ") == "Sent to vendor" and rvs.stage_of("DONE") == "Done - statement agrees" and rvs.stage_of("received") == "Revised bill received"
    assert rvs.stage_of("") == "" and rvs.stage_of("whatever") == "" and len(rvs.STAGE_LIST) < 255
    assert rvs.standing({"stage": "Done - statement agrees"}, TODAY)["final"]
    assert rvs.standing({"stage": "", "revised_on": "2026-10-10"}, TODAY)["text"] == "the clerk has not told the vendor yet - 2 days"
    assert rvs.iso("10/3/2026") == "2026-10-03" and rvs.iso(dt.datetime(2026, 10, 3, 9)) == "2026-10-03" and rvs.iso("x") == ""


def test_paid_section_lists_the_check_and_keeps_its_notes(tmp_path):
    from openpyxl import Workbook, load_workbook
    rows = [_row("LIKELY_VENDOR_LAG", "398000", "398000", 1500.0, 1500.0, "21", date="2026-08-05"),
            _row("LIKELY_VENDOR_LAG", "398001", "398001", 250.0, 250.0, "22", date="2026-08-06"),
            _row("MATCHED", "399126", "399126", 4752.18, 4752.18, "13")]
    payments = {"21": [{"check_no": "2601", "check_date": "2026-09-20", "amount": 1000.0, "pay_type": "Check"},
                       {"check_no": "2610", "check_date": "2026-09-27", "amount": 500.0, "pay_type": "Check"}],
                "22": [{"check_no": "", "check_date": "2026-09-22", "amount": 250.0, "pay_type": "CreditCard"}]}
    paid = r._paid_rows(rows, {"22": "they said 2 weeks"}, "2026-10-02", payments=payments)
    assert [p["invoice"] for p in paid] == ["398000", "398001"]
    assert paid[0]["check_no"] == "2601 + 2610" and paid[0]["check_date"] == "2026-09-27" and paid[0]["paid_amount"] == 1500.0 and paid[0]["amount"] == 1500.0
    assert paid[1]["check_no"] == "(creditcard)" and paid[1]["notes"] == "they said 2 weeks"
    path = tmp_path / "Reconciliation - V - as of 10-02-2026.xlsx"
    wb = Workbook(); wb.active.title = "Summary"
    rvs.write_sheet(wb, [{"bill_id": "11", "invoice": "400621", "invoice_date": "2026-09-10", "amount_billed": 8768.25, "corrected_amount": 8100.0, "reason": "tax", "stage": "To send"}], paid)
    wb.save(path); xlsx_verify.assert_clean(path)
    ws = load_workbook(path)[rvs.SHEET]
    assert [t for t in ws.tables] == [rvs.TABLE, rvs.PAID_TABLE] and ws.tables[rvs.PAID_TABLE].ref.startswith("A6:H")
    assert ws["A5"].value == rvs.PAID_TITLE and [c.value for c in ws[6]][:6] == ["Invoice #", "Invoice date", "Amount", "Check #", "Check date", "Paid amount"]
    assert rvs.harvest(path) and "11" in rvs.harvest(path) and "21" not in rvs.harvest(path)        # the first reader stops at the second section
    assert rvs.harvest_paid(path) == {"21": "", "22": "they said 2 weeks"}


def test_payments_for_reads_the_copy(monkeypatch):
    import json, sqlite3
    from shared import qbo_mirror
    con = sqlite3.connect(":memory:")
    con.execute("CREATE TABLE qbo_billpayment (id TEXT, txn_date TEXT, doc_number TEXT, total REAL, name TEXT, deleted INT, json TEXT)")
    monkeypatch.setattr(qbo_mirror, "_unpack", lambda b: json.loads(b))
    pay = lambda pid, date, no, lines, ptype="Check": con.execute("INSERT INTO qbo_billpayment VALUES (?,?,?,?,?,0,?)", (pid, date, no, sum(a for _b, a in lines), None, json.dumps({"PayType": ptype, "Line": [{"Amount": a, "LinkedTxn": [{"TxnId": b, "TxnType": "Bill"}]} for b, a in lines]})))
    pay("p1", "2026-09-20", "2601", [("21", 1000.0), ("99", 5.0)])
    pay("p2", "2026-09-27", "2610", [("21", 500.0)])
    pay("p3", "2025-01-01", "1", [("21", 7.0)])                                          # before `since`: left out
    out = rvs.payments_for(["21", "22"], con=con, since="2026-01-01")
    assert [p["check_no"] for p in out["21"]] == ["2601", "2610"] and out["21"][0]["amount"] == 1000.0 and "22" not in out and "99" not in out
