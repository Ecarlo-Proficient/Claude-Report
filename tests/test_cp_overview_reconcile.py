"""one-offs/cp_overview_reconcile.py - the built-in formula evaluator gives the
numbers Excel would for the shapes the CP workbooks use."""
import importlib.util
from pathlib import Path

from openpyxl import Workbook

_p = Path(__file__).resolve().parents[1] / "one-offs" / "cp_overview_reconcile.py"
_spec = importlib.util.spec_from_file_location("cp_overview_reconcile", _p)
rc = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(rc)


def _book():
    wb = Workbook()
    ws = wb.active
    ws.title = "Overview"
    job = wb.create_sheet("CP790")
    job["F17"], job["G17"] = 791905.60, 778868.99
    ws["J7"] = "='CP790'!F17"
    ws["K7"] = "='CP790'!G17"
    ws["L7"] = "=J7-K7"
    ws["N7"] = "=J7*0.1"
    ws["T7"] = '=IF(K7=0,"",J7/K7)'
    ws["C7"], ws["D7"] = 0, 0
    ws["F7"] = '=IF(OR(C7=0,D7=0),"",C7-D7)'          # lazy IF: no #DIV/0 on the other branch
    ws["G7"] = '=IF(F7="","",F7/C7)'
    ws["C8"], ws["D8"] = 100, 0
    ws["C9"], ws["D9"] = 50, 20
    ws["J8"], ws["J9"] = 10, 30
    ws["X1"] = "=SUMPRODUCT((D7:D9<>0)*(C7:C9<>0)*J7:J9)"
    ws["X2"] = "=SUM(C7:C9)"
    return rc.Book(wb)


def test_cross_sheet_refs_and_arithmetic():
    bk = _book()
    assert bk.value("Overview", "L7") == 791905.60 - 778868.99
    assert abs(bk.value("Overview", "N7") - 79190.56) < 1e-6


def test_lazy_if_and_blank():
    bk = _book()
    assert bk.value("Overview", "F7") == ""
    assert bk.value("Overview", "G7") == ""
    assert abs(bk.value("Overview", "T7") - 791905.60 / 778868.99) < 1e-12


def test_sumproduct_mask_and_sum():
    bk = _book()
    assert bk.value("Overview", "X1") == 30            # only row 9 has both a contract and an ETC
    assert bk.value("Overview", "X2") == 150
