"""statement_set: one vendor's statements as ONE running record (owner 10/07/2026).
The merge (full list / past-due letter / activity statement), duplicates,
unreadable statements, changes between updates, and the clerk's notes surviving
a re-run of the Excel. No network, no share."""
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "statement-reconciler"))

import statement_set as ss  # noqa: E402


@dataclass
class L:                      # a StmtLine stand-in
    date: str
    ref: str
    amount: float
    address: str = ""


def _doc(name, as_of, lines, due=None, template="", origin="inbox", sha=None):
    d = ss.Doc(path=Path(f"/x/{name}"), as_of=as_of, lines=lines, template=template,
               amount_due=round(sum(x.amount for x in lines), 2) if due is None else due,
               origin=origin, sha=sha or name)
    return ss.classify(d)


def test_a_newer_full_list_replaces_the_older_one():
    old = _doc("aug.pdf", "2026-08-31", [L("2026-07-01", "A", 10), L("2026-08-01", "B", 20)])
    new = _doc("sep.pdf", "2026-09-30", [L("2026-08-01", "B", 18.5), L("2026-09-02", "C", 5)])
    m, docs = ss.plan([new, old])
    assert {x.ref for x in m.lines} == {"B", "C"} and m.amount_due == 23.5 and m.gap == 0
    assert (old.status, new.status) == ("replaced", "current")
    ch = ss.changes_for(docs, m, {})
    assert [r["ref"] for r in ch["gone"]] == ["A"] and [r["ref"] for r in ch["new"]] == ["C"]
    assert ch["changed"] == [{"ref": "B", "date": "2026-08-01", "amount": 18.5, "was": 20,
                              "where": ""}]


def test_past_due_letter_plus_activity_statement_is_one_open_list():
    """Cowtown 10/01 + 10/02: the letter itemizes everything through August net of
    September's payments; the activity statement adds September. Its balance
    forward and payments are already inside the letter - they must not count."""
    letter = _doc("letter.pdf", "2026-10-01",
                  [L("2026-07-15", "396001", 100), L("2026-08-20", "399001", 50)],
                  template="vendor_cowtown")
    act = _doc("act.pdf", "2026-10-02",
               [L("2026-08-31", "", 400), L("2026-09-21", "PMT", -250),
                L("2026-09-04", "400300", 30), L("2026-09-21", "401417", 7),
                L("2026-09-21", "PMT", -7)])
    assert letter.kind == "pastdue" and letter.hi == "2026-08-31"
    assert act.kind == "activity" and act.lo == "2026-08-31" and act.ties_out
    m, _ = ss.plan([letter, act])
    assert {x.ref for x in m.lines} == {"396001", "399001", "400300", "401417"}
    assert m.amount_due == act.amount_due and m.as_of == "2026-10-02"
    assert m.gap == round(187 - 180, 2)          # the vendor's two documents differ by 7


def test_a_lone_activity_statement_keeps_its_balance_forward():
    act = _doc("act.pdf", "2026-09-01", [L("2026-07-31", "", 400), L("2026-08-04", "1", 30),
                                         L("2026-08-31", "PMT", -100)])
    m, _ = ss.plan([act])
    assert m.gap == 0 and len(m.lines) == 3


def test_an_old_activity_statement_is_replaced_by_the_letter_that_itemizes_its_month():
    sep1 = _doc("sep1.pdf", "2026-09-01", [L("2026-07-31", "", 400), L("2026-08-04", "1", 30)],
                origin="legacy")
    letter = _doc("letter.pdf", "2026-10-01", [L("2026-08-04", "1", 30)], template="vendor_cowtown")
    _m, _ = ss.plan([sep1, letter])
    assert (sep1.status, letter.status) == ("replaced", "current")


def test_same_file_twice_is_a_duplicate_and_the_filed_copy_wins():
    a = _doc("RCI 10-07.pdf", "2026-09-30", [L("2026-09-01", "X", 5)], sha="same")
    b = _doc("RCI  Ready Cable 10-07.pdf", "2026-09-30", [L("2026-09-01", "X", 5)], sha="same",
             origin="current")
    ss.plan([a, b])
    assert (a.status, b.status) == ("duplicate", "current")


def test_a_statement_that_does_not_tie_out_is_kept_out_of_the_merge():
    bad = _doc("bad.pdf", "2026-10-31", [L("2026-10-01", "Z", 5)], due=999.0)
    good = _doc("good.pdf", "2026-09-30", [L("2026-09-01", "Y", 5)])
    m, _ = ss.plan([bad, good])
    assert bad.status == "unreadable" and {x.ref for x in m.lines} == {"Y"}


def test_changes_compare_the_new_batch_with_the_update_before_it():
    sep = _doc("sep.pdf", "2026-09-01", [L("2026-08-01", "A", 10)], origin="current")
    letter = _doc("letter.pdf", "2026-10-01", [L("2026-08-01", "A", 4)], template="vendor_cowtown")
    act = _doc("act.pdf", "2026-10-02", [L("2026-08-31", "", 4), L("2026-09-02", "B", 6)])
    m, docs = ss.plan([sep, letter, act])
    ch = ss.changes_for(docs, m, {})
    assert ch["frm"] == "2026-09-01" and [r["ref"] for r in ch["new"]] == ["B"]
    assert ch["changed"][0]["was"] == 10 and ch["changed"][0]["amount"] == 4
    # next run, the old statement is in History: the stored changes are reused
    assert ss.changes_for([letter, act], m, {"changes": ch}) == ch


def test_names_are_month_first():
    assert ss.history_name("2026-09-30", "Duplicate", "a.pdf") == "09-30-2026 Duplicate - a.pdf"
    assert ss.report_name("RCI Ready Cable", "2026-09-30") == \
        "Reconciliation - RCI Ready Cable - as of 09-30-2026.xlsx"


def test_the_clerks_note_survives_a_rerun_and_follows_the_bill(tmp_path):
    import statement_reconciler as sr
    bills = [sr.QboBill("77", "C1", "2026-09-01", 100.0, 100.0, memo="Approved")]
    rows = [sr.ReconRow("MATCHED", "2026-09-01", "C1", 100.0, 100.0, "", "",
                        stmt_ref="C1", qbo_ref="C1", stmt_date="2026-09-01",
                        qbo_date="2026-09-01", qbo_memo="Approved", qbo_bill_id="77"),
            sr.ReconRow("MISSING_IN_QBO", "2026-09-02", "N9", 5.0, 0.0, "", "",
                        "ENTER BILL IN QBO", stmt_ref="N9", stmt_date="2026-09-02")]
    out = tmp_path / "r.xlsx"
    lines = [L("2026-09-01", "C1", 100.0), L("2026-09-02", "N9", 5.0)]
    sr.write_excel(out, "V", "2026-09-30", 105.0, bills, rows, line_sum=105.0, stmt_lines=lines)
    from openpyxl import load_workbook
    wb = load_workbook(out)
    ws = wb["Summary"]
    for r in range(1, ws.max_row + 1):            # she types on both rows
        if ws.cell(r, 2).value == "C1":
            ws.cell(r, 13).value = "PM approved by phone"
        if ws.cell(r, 2).value == "N9":
            ws.cell(r, 13).value = "asked vendor for copy"
    wb.save(out)
    notes = ss.fold_notes({}, ss.harvest_notes(out), "10/07/2026")
    assert notes["bill:77"]["note"] == "PM approved by phone"
    assert notes["ref:N9"]["note"] == "asked vendor for copy"
    sr.write_excel(out, "V", "2026-09-30", 105.0, bills, list(reversed(rows)), line_sum=105.0,
                   stmt_lines=lines, notes=notes)    # re-run: rows in another order
    ws = load_workbook(out)["Summary"]
    got = {ws.cell(r, 2).value: ws.cell(r, 13).value for r in range(1, ws.max_row + 1)
           if ws.cell(r, 2).value in ("C1", "N9")}
    assert got == {"C1": "PM approved by phone", "N9": "asked vendor for copy"}
    # she clears one: the next harvest drops it
    wb = load_workbook(out)
    ws = wb["Summary"]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value == "N9":
            ws.cell(r, 13).value = None
    wb.save(out)
    assert "ref:N9" not in ss.fold_notes(notes, ss.harvest_notes(out), "10/08/2026")


def test_a_history_twin_never_knocks_out_the_live_statement():
    live = _doc("RCI.pdf", "2026-09-30", [L("2026-09-01", "X", 5)], sha="same", origin="current")
    twin = _doc("09-30-2026 Duplicate - RCI copy.pdf", "2026-09-30", [], sha="same", origin="history")
    again = _doc("RCI again.pdf", "2026-09-30", [L("2026-09-01", "X", 5)], sha="same", origin="inbox")
    m, _ = ss.plan([twin, again, live])
    assert live.status == "current" and again.status == "duplicate" and m.lines


def test_excel_layout_owner_10_08(tmp_path):
    """No blank row 3 under the frozen title, her column is 'Notes', the tool's is
    'Finding', no Statements / Changes sheets."""
    import statement_reconciler as sr
    from openpyxl import load_workbook
    rows = [sr.ReconRow("MISSING_IN_QBO", "2026-09-02", "N9", 5.0, 0.0, "", "", "ENTER BILL IN QBO",
                        stmt_ref="N9", stmt_date="2026-09-02")]
    out = tmp_path / "r.xlsx"
    sr.write_excel(out, "V", "2026-09-30", 5.0, [], rows, line_sum=5.0, stmt_lines=[L("2026-09-02", "N9", 5.0)])
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary"]
    ws = wb["Summary"]
    assert ws["A3"].value == "TIE-OUT" and ws.freeze_panes == "A3"
    head = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Stmt Ref #")
    assert head[11] == "Finding" and head[12] == "Notes"
