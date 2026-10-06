"""project-pnl layout review (2026-10-02, rounds 1-23): the rules the owner
corrected, pinned so a later layout change cannot quietly undo them.

- Billed is GROSS (retainage included): a draw's gross = net billed + retained.
- Every TO-DATE overhead is rate x BILLED, never the contract: the draw rows,
  the coverage TOTAL, and the P&L-by-account check against them. (The
  projection alone uses the contract.)
- The draw coverage sits on the P&L, newest draw first, each name linking to
  its own sheet; the P&L checks live on Reconciliations under the QuickBooks
  tie-out, every formula sheet-qualified.
- Dollar figures carry a $.
"""
import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "project-pnl"))
_spec = importlib.util.spec_from_file_location(
    "project_pnl_export", ROOT / "project-pnl" / "project_pnl_export.py")
pnl = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pnl)

# (name, period, net billed, costs, retainage held, retainage billed)
DRAWS = [("Draw 1", "11/01/25–11/30/25", 29190.38, 1828.81, 3243.38, 0.0),
         ("Draw 2", "01/01/26–01/30/26", 126674.10, 69547.35, 14074.90, 0.0)]


def _coverage(ws, top=20, contract=None):
    return pnl.write_draw_coverage(
        ws, top, 1, DRAWS, {}, {"__outside": {"total": 500.0}}, None,
        overhead_pct=10.0, contract_ref=contract,
        sheet_links={"__next": "Next Draw", "11/01/25–11/30/25": "Draw 1",
                     "01/01/26–01/30/26": "Draw 2"})


def _row_of(ws, text):
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == text)


def test_coverage_gross_includes_retainage_and_overhead_is_on_billed():
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    cov = _coverage(ws)
    r2 = _row_of(ws, "Draw 2")
    assert ws.cell(r2, 5).value == round(126674.10 + 14074.90, 2)    # gross billed
    assert ws.cell(r2, 9).value == f"=E{r2}*0.1"                     # OH = 10% x gross billed
    assert ws.cell(r2, 10).value == f"=E{r2}-F{r2}-I{r2}"            # net = billed - costs - OH
    tot = _row_of(ws, "TOTAL  (draws)")
    assert cov["cov_oh_tot"] == f"'P&L'!I{tot}"
    assert cov["cov_gross_tot"] == f"'P&L'!E{tot}"


def test_coverage_is_newest_first_and_links_to_each_draw_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    _coverage(ws, contract="'P&L'!B4")
    nxt, d2, d1 = (_row_of(ws, t) for t in ("Next draw (forming)", "Draw 2", "Draw 1"))
    tot = _row_of(ws, "TOTAL  (draws)")
    assert d2 < d1 < tot < nxt          # the next draw sits UNDER the total (10/06)
    assert ws.cell(d2, 1).hyperlink.target == "#'Draw 2'!B2"
    assert ws.cell(nxt, 1).hyperlink.target == "#'Next Draw'!B2"
    # % Billed: this draw and every older one, over the contract
    assert ws.cell(d2, 12).value == f"=IF('P&L'!B4=0,\"\",SUM(E{d2}:E{d1})/'P&L'!B4)"


def test_checks_are_sheet_qualified_and_compare_overhead_on_billed():
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    cov = _coverage(ws)
    refs = {"billed": "C5", "costs": "C6", "opex": "N4", "ret": "I8",
            "oh_billed": "R6", **{k: v for k, v in cov.items() if k.startswith("cov_")}}
    checks = pnl._wire_pl_support(wb, None, DRAWS, {}, refs)
    by = {c[0]: c for c in checks}
    assert by["Draws gross billed = Billed to Date"][1] == "'P&L'!C5"
    assert by["② overhead (on billed) = draws overhead"][1:3] == ("-'P&L'!R6", cov["cov_oh_tot"])
    assert by["Draws retained = retainage held net of releases"][1] == "'P&L'!I8"
    for _label, plf, supf, _link in checks:
        assert "'P&L'!" in plf and "!" in supf


def test_checks_join_the_reconciliations_table():
    wb = Workbook()
    checks = [("Draws gross billed = Billed to Date", "'P&L'!C5", "'P&L'!E30", None)]
    pnl.build_sheet_reconciliations(
        wb, "CP000", {"name": "Test"}, {}, 100.0, 50.0, 0.0,
        {"billed": "Transactions!D5", "cogs": "Transactions!D6", "exp": "Transactions!D7"},
        "10/02/2026", pl_checks=checks)
    ws = wb["Reconciliations"]
    hdr = _row_of(ws, "P&L ties to its support")
    assert ws.cell(hdr + 1, 1).value == checks[0][0]
    assert ws.cell(hdr + 1, 2).value == "='P&L'!C5"
    status = next(ws.cell(r, 2).value for r in range(1, ws.max_row + 1)
                  if ws.cell(r, 1).value == "RECONCILIATION STATUS")
    assert f"ABS(D{hdr + 1})<0.005" in status
    assert ws.cell(hdr + 1, 2).number_format.startswith('"$"')


def test_dollar_columns_widen_for_the_sign():
    wb = Workbook()
    ws = wb.active
    ws.column_dimensions["B"].width = 12
    c = ws.cell(1, 2, value=1234.5)
    c.number_format = pnl.CURR_FMT
    pnl._widen_for_dollar(wb)
    assert ws.column_dimensions["B"].width == 13.5
    assert pnl.CURR_FMT.startswith('"$"') and pnl.RET_FMT.startswith('"$"')


def test_billing_outside_the_draws_joins_the_checks():
    # CP585's shape: three untagged pre-period invoices, a release outside every
    # draw window, and not-billed retainage that was never placed on a draw
    groups = {"__untagged": {"net_billed": 96777.0, "retainage_held": 0.0,
                             "retainage_billed": 0.0},
              "__retainage_billed": {"net_billed": 12663.0, "retainage_held": 0.0,
                                     "retainage_billed": 12663.0},
              "__retainage": {"total": 10753.0}}
    out = pnl._outside_draws(groups)
    assert out == {"gross": 107530.0, "ret": -1910.0}
    wb = Workbook()
    ws = wb.active
    ws.title = "P&L"
    cov = _coverage(ws)
    refs = {"billed": "C5", "costs": "C6", "ret": "I8", "oh_billed": "R6",
            "outside": out, "oh_rate": 0.10,
            **{k: v for k, v in cov.items() if k.startswith("cov_")}}
    by = {c[0].split(" = ")[0]: c for c in pnl._wire_pl_support(wb, None, DRAWS, {}, refs)}
    assert by["Draws gross billed + 107,530.00 outside the draws"][2] == cov["cov_gross_tot"] + "+107530.0"
    assert by["Draws retained - 1,910.00 outside the draws"][2] == cov["cov_ret_tot"] + "-1910.0"
    assert by["② overhead (on billed)"][2] == cov["cov_oh_tot"] + "+10753.0"


def test_only_untagged_invoices_after_the_first_draw_are_forming():
    groups = {"__untagged": {"invoices": [{"date": "2025-01-20", "doc_num": "old"},
                                          {"date": "2026-09-25", "doc_num": "new"}]}}
    costs = {"__disregarded": {"anchor": "2026-04-20"}}
    assert [i["doc_num"] for i in pnl._forming_invoices(groups, costs)] == ["new"]


def test_legacy_quickbooks_income_is_gross_through_the_pl_end():
    # MFD295's shape: invoices at their total with the retainage held added
    # back, a release at its total, and nothing after the P&L's end date
    groups = {"d1": {"invoices": [{"date": "2025-03-20", "amount": 90.0, "retainage": 10.0}]},
              "__retainage_billed": {"invoices": [{"date": "2024-12-31", "amount": 6.0}]},
              "d9": {"invoices": [{"date": "2026-09-25", "amount": 50.0, "retainage": 5.0}]}}
    t = pnl._synth_pl_totals([], [], groups, "1", {}, {}, "2026-09-20")
    assert t["income"] == 106.0


def test_credits_net_into_cost_as_their_own_transactions():
    from shared import qbo_costs as qc
    card = {"Id": "1", "Credit": True, "TotalAmt": 71.57,
            "Line": [{"Amount": 71.57, "AccountBasedExpenseLineDetail": {}}]}
    charge = {"Id": "2", "TotalAmt": 150.0,
              "Line": [{"Amount": 150.0, "AccountBasedExpenseLineDetail": {}}]}
    vc = {"Id": "3", "TotalAmt": 940.0, "VendorRef": {"value": "9", "name": "White Cap"},
          "Line": [{"Amount": 940.0, "ItemBasedExpenseLineDetail": {}}]}
    out = qc.with_credits([card, charge], [vc])
    assert [round(sum(l["Amount"] for l in t["Line"]), 2) for t in out] == [-71.57, 150.0, -940.0]
    assert out[2]["_tx_type"] == "VendorCredit" and out[2]["EntityRef"]["name"] == "White Cap"
    assert card["Line"][0]["Amount"] == 71.57          # the pulled record is not changed
    assert pnl._pay_state(0.0, -940.0)[0] == "CREDIT"
    assert "vendorcredit%3FtxnId%3D3" in pnl._qbo_txn_url("VendorCredit", "3", "1")


def test_a_combined_draw_is_cut_by_each_invoice_date():
    tx = {"income": [{"date": "2026-09-01", "billed": 70.0,
                      "docs": [{"date": "2026-09-01", "billed": 50.0},
                               {"date": "2026-09-07", "billed": 20.0}]}]}
    assert pnl._after_cut(tx, "2026-09-01")["income"] == 20.0


def test_retainage_lines_are_told_apart_by_the_item_account():
    pnl.set_retainage_items([{"Id": "1812", "IncomeAccountRef": {"value": "a"}},
                             {"Id": "PC00", "IncomeAccountRef": {"value": "i"}}],
                            [{"Id": "a", "AccountType": "Other Current Asset"},
                             {"Id": "i", "AccountType": "Income"}])
    try:
        line = lambda item, amt, d="": {"Amount": amt, "Description": d,   # noqa: E731
                                        "DetailType": "SalesItemLineDetail",
                                        "SalesItemLineDetail": {"ItemRef": {"value": item}}}
        invs = [
            # a draw: work + 10% held on the receivable item
            {"Id": "a", "DocNumber": "1", "TxnDate": "2026-03-10", "TotalAmt": 90.0,
             "PrivateNote": "Period: 03/01/2026 - 03/31/2026",
             "Line": [line("PC00", 100.0), line("1812", -10.0, "Retainage")]},
            # a PC00 "Retainage" invoice = real income in QuickBooks
            {"Id": "b", "DocNumber": "2", "TxnDate": "2026-04-10", "TotalAmt": 30.0,
             "PrivateNote": "Retainage", "Line": [line("PC00", 30.0, "Retainage")]},
            # a release on the receivable item = collecting, never billing
            {"Id": "c", "DocNumber": "3", "TxnDate": "2026-05-10", "TotalAmt": 10.0,
             "PrivateNote": "Retainage release", "Line": [line("1812", 10.0, "Retainage")]},
        ]
        g = pnl.group_invoices_by_draw(invs)
        billed = sum(i["gross"] for k, x in g.items() for i in x.get("invoices", [])
                     if k != "__retainage")
        assert billed == 130.0
        assert g["__retainage_billed"]["retainage_billed"] == 10.0
    finally:
        pnl._RET_ITEMS.clear()


def test_a_pc00_retainage_invoice_is_booked_so_its_release_is_not_billed_twice():
    sys.path.insert(0, str(ROOT / "project-pnl"))
    import completed_pnl as cp
    # CP610's shape after the re-bill: #32656 billed retainage AS INCOME on
    # PC00 (gross), later released on the Retainage Receivable item
    src = {"invoices": [{"doc": "100", "gross": 200.0, "withheld": 0.0, "ret_billed": 0.0},
                        {"doc": "32656", "gross": 23256.0, "withheld": 0.0, "ret_billed": 0.0},
                        {"doc": "40000", "gross": 0.0, "withheld": 0.0, "ret_billed": 23256.0}],
           "not_billed": 0.0, "sections": [], "ret_income_docs": ["32656"]}
    assert cp._totals(src)["billed"] == 23456.0
    src["ret_income_docs"] = []                        # without the mark it double counts
    assert cp._totals(src)["billed"] == 46712.0


def test_the_pc00_retainage_invoice_is_marked_on_the_workbook():
    pnl.set_retainage_items([{"Id": "1812", "IncomeAccountRef": {"value": "a"}},
                             {"Id": "PC00", "IncomeAccountRef": {"value": "i"}}],
                            [{"Id": "a", "AccountType": "Other Current Asset"},
                             {"Id": "i", "AccountType": "Income"}])
    try:
        inv = {"Id": "b", "DocNumber": "32656", "TxnDate": "2026-04-10", "TotalAmt": 30.0,
               "PrivateNote": "Retainage", "Line": [{"Amount": 30.0, "Description": "Retainage",
                                                     "DetailType": "SalesItemLineDetail",
                                                     "SalesItemLineDetail": {"ItemRef": {"value": "PC00"}}}]}
        g = pnl.group_invoices_by_draw([inv])
        wb = Workbook()
        pnl.mark_retainage_income(wb, g)
        assert wb.defined_names[pnl.RET_INCOME_NAME].attr_text == '"32656"'
    finally:
        pnl._RET_ITEMS.clear()


def test_retainage_still_owed_and_not_billed_yet():
    # CP610: draws entered NET, retainage invoiced on PC00 (#32656) + a
    # not-billed record (#32960), then re-billed on 99 - Retainage (#34729, open)
    groups = {"d4": {"invoices": [{"gross": 100.0, "retainage": 0.0, "retainage_billed": 0.0,
                                   "balance": 0.0}]},
              "__untagged": {"invoices": [{"gross": 23256.0, "retainage": 0.0, "retainage_billed": 0.0,
                                           "balance": 0.0, "ret_income": True}]},
              "__retainage_billed": {"invoices": [{"gross": 0.0, "retainage": 0.0,
                                                   "retainage_billed": 24572.0, "balance": 24572.0}]},
              "__retainage": {"pool": 1316.0}}
    rows = [{"withheld": 0.0, "billed_ret": 24572.0}, {"not_billed_ret": 1316.0}]
    st = pnl._retainage_state(groups, rows, 24572.0)
    assert st["still_owed"] == 24572.0 and st["not_billed_yet"] == 0.0 and st["net_draws"]
    # CP595 before the owner's invoices: pay app 21,185.17, 17,730.00 invoiced
    groups = {"d3": {"invoices": [{"gross": 100.0, "retainage": 0.0, "retainage_billed": 0.0,
                                   "balance": 0.0}]},
              "__untagged": {"invoices": [{"gross": 17730.0, "retainage": 0.0, "retainage_billed": 0.0,
                                           "balance": 0.0, "ret_income": True}]}}
    st = pnl._retainage_state(groups, [{"withheld": 0.0, "billed_ret": 0.0}], 21185.17)
    assert st["not_billed_yet"] == 3455.17
    # a job that withholds on its draws: held 79,190.56, nothing released
    groups = {"d1": {"invoices": [{"gross": 791905.6, "retainage": 79190.56,
                                   "retainage_billed": 0.0, "balance": 0.0}]}}
    st = pnl._retainage_state(groups, [{"withheld": 79190.56, "billed_ret": 0.0}], None)
    assert st["still_owed"] == 79190.56 and st["not_billed_yet"] == 79190.56 and not st["net_draws"]


def test_a_misspelled_retainage_line_is_still_marked_as_booked():
    # CP610 #32960: PC00, its line reads "REITANAGE", the memo says "not billed"
    pnl.set_retainage_items([{"Id": "PC00", "IncomeAccountRef": {"value": "i"}}],
                            [{"Id": "i", "AccountType": "Income"}])
    try:
        inv = {"Id": "x", "DocNumber": "32960", "TxnDate": "2025-06-01", "TotalAmt": 1316.0,
               "PrivateNote": "Draw #4 - Retainage not billed",
               "Line": [{"Amount": 1316.0, "Description": "REITANAGE",
                         "DetailType": "SalesItemLineDetail",
                         "SalesItemLineDetail": {"ItemRef": {"value": "PC00"}}}]}
        draw = {"Id": "y", "DocNumber": "100", "TxnDate": "2025-06-01", "TotalAmt": 500.0,
                "PrivateNote": "Draw #4 - retainage held per contract",
                "Line": [{"Amount": 500.0, "Description": "Foundation",
                          "DetailType": "SalesItemLineDetail",
                          "SalesItemLineDetail": {"ItemRef": {"value": "PC00"}}}]}
        g = pnl.group_invoices_by_draw([inv, draw])
        marked = {i["doc_num"] for x in g.values() for i in x.get("invoices", []) if i.get("ret_income")}
        assert marked == {"32960"}         # the draw whose MEMO mentions retainage is not
    finally:
        pnl._RET_ITEMS.clear()
