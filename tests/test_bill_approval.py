"""shared/bill_approval - the ONE approval rule (Bill Tracker + statement reconciler),
and the reconciler's workbook showing its states (owner 10/08/2026: bills are approved
in QBO now, but older ones still carry NOT APPROVED; 10/13/2026: QuickBooks' own
approval-queue report answers for the new ones, as of the time it synced)."""
import datetime as dt
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "statement-reconciler"))

from shared import bill_approval as ba  # noqa: E402

STAMP = dt.datetime(2026, 10, 13, 16, 19)


def test_three_states_without_a_report():
    assert ba.approval_state("NOT APPROVED - CP861", "2026-08-01", 10) == ba.APPROVAL_NO
    assert ba.approval_state("  not approved", dt.date(2026, 10, 1), 10) == ba.APPROVAL_NO   # tag wins
    assert ba.approval_state("CP861", "2026-10-01T09:00:00", 10) == ba.APPROVAL_CHECK         # new, unpaid
    assert ba.approval_state("CP861", "2026-10-01", 0) == ba.APPROVAL_YES                     # new, paid
    assert ba.approval_state("CP861", "2026-09-15", 10) == ba.APPROVAL_YES                    # old, no tag
    assert ba.approval_state("Approved 5/1 (was Not Approved)", "2026-08-01", 10) == ba.APPROVAL_YES
    assert ba.approval_state(None, "", None) == ba.APPROVAL_YES


def _report_xlsx(path: Path, stamp: str = "Tuesday, October 13, 2026 | 04:19 PM GMT -05:00") -> Path:
    """The shape QuickBooks exports: title rows, a header, Vendor -> Class groups in the
    'Bill number' column, one row per bill LINE, '--' for no class, totals, the run stamp."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    d = dt.datetime
    rows = [
        ("Proficient Concrete, LLC",), ("Bill Approval Status ALL",), ("All dates",),
        ("Vendor", "Class", "Bill number", "Date", "Amount", "Paid status", "Approval status"),
        (None, None, "RCI READY CABLE"),
        (None, None, "Commercial"),
        (None, None, "MCK777522", d(2026, 7, 30), 185.11, "Unpaid", "Requires Approval"),
        (None, None, "WAX793470", d(2026, 9, 17), 1803.22, "Unpaid", "Pending Approval"),
        (None, None, "WAX793470", d(2026, 9, 17), 1803.22, "Unpaid", "Pending Approval"),
        (None, None, "Total for Commercial", None, 3791.55),
        (None, None, "Residential"),
        (None, None, None, d(2026, 9, 23), 402.67, "Unpaid", "Pending Approval"),   # entered with no bill #
        (None, None, "Total for Residential", None, 402.67),
        (None, None, "24", d(2026, 9, 16), 2400, "Paid", "Requires Approval"),
        (None, None, "Total for --", None, 2400),
        (None, None, "Total for RCI READY CABLE", None, 6594.22),
        (None, None, "Croell, Inc"),
        (None, None, "Residential"),
        (None, None, "Total for Residential", None, 0),
        (None, None, "Total for Croell, Inc", None, 0),
        (stamp,),
    ]
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def test_report_reads_quickbooks_shape(tmp_path):
    rep = ba.ApprovalReport(_report_xlsx(tmp_path / "Bill Approval Status ALL.xlsx"), now=STAMP + dt.timedelta(hours=2))
    assert rep.as_of == STAMP and rep.stamp_from_report and rep.fresh
    assert len(rep.bills) == 3                       # the 2-line bill counted once
    assert rep.lookup("RCI Ready Cable", "wax 793470")["state"] == ba.APPROVAL_PENDING   # case + spacing folded
    assert rep.lookup("RCI READY CABLE", "MCK777522")["state"] == ba.APPROVAL_REQUIRES
    assert rep.lookup("RCI READY CABLE", "", "2026-09-23", 402.67)["state"] == ba.APPROVAL_PENDING  # no bill #
    assert rep.lookup("RCI READY CABLE", "ZZZ") is None
    assert rep.lookup("Croell, Inc", "24") is None    # a bill number belongs to its vendor
    assert "10/13/2026 04:19 PM" in rep.describe() and "3 bills" in rep.describe()


def test_report_decides_the_new_bills(tmp_path):
    rep = ba.ApprovalReport(_report_xlsx(tmp_path / "r.xlsx"), now=STAMP + dt.timedelta(hours=2))
    kw = dict(report=rep, vendor="RCI READY CABLE")
    # on the report -> its state (even an old-process bill)
    assert ba.approval_state("CP800", "2026-09-17", 1803.22, docnum="WAX793470", **kw) == ba.APPROVAL_PENDING
    assert ba.approval_state("CP800", "2026-07-30", 185.11, docnum="MCK777522", **kw) == ba.APPROVAL_REQUIRES
    assert ba.approval_state("RP7533", "2026-09-24", 402.67, docnum="", txn_date="2026-09-23", amount=402.67, **kw) \
        == ba.APPROVAL_PENDING
    # new, unpaid, NOT on a fresh report that was run after it was entered -> approved
    assert ba.approval_state("CP800", "2026-09-26", 30607.04, docnum="RWC792172", **kw) == ba.APPROVAL_YES
    # entered after the report ran -> the report cannot know it -> check QBO
    assert ba.approval_state("CP800", "2026-10-14", 5, docnum="NEW1", **kw) == ba.APPROVAL_CHECK
    # the memo tag still wins; a paid bill is approved whatever the report says
    assert ba.approval_state("NOT APPROVED x", "2026-09-17", 1803.22, docnum="WAX793470", **kw) == ba.APPROVAL_NO
    assert ba.approval_state("CP800", "2026-09-17", 0, docnum="WAX793470", **kw) == ba.APPROVAL_YES
    # old-process bill, no tag, not on the report -> approved as before
    assert ba.approval_state("CP800", "2026-09-01", 10, docnum="OLD", **kw) == ba.APPROVAL_YES


def test_stale_report_falls_back_to_check_qbo(tmp_path):
    rep = ba.ApprovalReport(_report_xlsx(tmp_path / "r.xlsx"), now=STAMP + dt.timedelta(days=3))
    assert not rep.fresh and "too old" in rep.describe()
    kw = dict(report=rep, vendor="RCI READY CABLE")
    assert ba.approval_state("CP800", "2026-09-17", 1803.22, docnum="WAX793470", **kw) == ba.APPROVAL_CHECK
    assert ba.approval_state("CP800", "2026-09-26", 30607.04, docnum="RWC792172", **kw) == ba.APPROVAL_CHECK
    assert ba.approval_state("CP800", "2026-09-01", 10, docnum="OLD", **kw) == ba.APPROVAL_YES


def test_report_without_a_stamp_uses_the_file_time(tmp_path):
    rep = ba.ApprovalReport(_report_xlsx(tmp_path / "r.xlsx", stamp="no stamp here"))
    assert rep.as_of is not None and not rep.stamp_from_report


def test_missing_report_is_none_not_a_crash(tmp_path):
    assert ba.load(tmp_path / "nope.xlsx") is None
    assert "check QBO" in ba.report_line(None)


def test_reconciler_workbook_shows_the_states(tmp_path, monkeypatch):
    import statement_reconciler as sr
    from openpyxl import load_workbook
    rep = ba.ApprovalReport(_report_xlsx(tmp_path / "r.xlsx"), now=STAMP + dt.timedelta(hours=2))
    monkeypatch.setattr(ba, "current", lambda reload=False: rep)
    bills = [sr.QboBill("1", "A1", "2026-08-01", 10.0, 10.0, memo="NOT APPROVED", created="2026-08-01", vendor="RCI READY CABLE"),
             sr.QboBill("2", "B2", "2026-10-14", 20.0, 20.0, memo="CP861", created="2026-10-14", vendor="RCI READY CABLE"),
             sr.QboBill("3", "C3", "2026-08-01", 30.0, 30.0, memo="CP861", created="2026-08-01", vendor="RCI READY CABLE"),
             sr.QboBill("4", "WAX793470", "2026-09-17", 1803.22, 1803.22, memo="CP800", created="2026-09-17", vendor="RCI READY CABLE"),
             sr.QboBill("5", "MCK777522", "2026-07-30", 185.11, 185.11, memo="CP800", created="2026-07-30", vendor="RCI READY CABLE"),
             sr.QboBill("6", "D6", "2026-09-26", 40.0, 40.0, memo="CP800", created="2026-09-26", vendor="RCI READY CABLE")]
    rows = [sr.ReconRow("MATCHED", b.txn_date, b.doc_number, b.open_balance, b.open_balance, "", "",
                        stmt_ref=b.doc_number, qbo_ref=b.doc_number, stmt_date=b.txn_date,
                        qbo_date=b.txn_date, qbo_memo=b.memo, qbo_bill_id=b.bill_id) for b in bills]
    out = tmp_path / "w.xlsx"
    total = round(sum(b.open_balance for b in bills), 2)
    sr.write_excel(out, "RCI READY CABLE", "2026-10-07", total, bills, rows, line_sum=total)
    ws = load_workbook(out)["Summary"]
    approved = {r[1]: r[10] for r in ws.iter_rows(values_only=True) if r[1] in ("A1", "B2", "C3", "WAX793470", "MCK777522", "D6")}
    assert approved == {"A1": "Not Approved", "B2": "Check QBO", "C3": "Yes", "WAX793470": "Pending Approval",
                        "MCK777522": "Requires Approval", "D6": "Yes"}
    text = " ".join(str(c) for r in ws.iter_rows(values_only=True) for c in r if c)
    assert "approval queue" in text and "NOT APPROVED by memo" in text and "synced from QuickBooks 10/13/2026" in text
    import statement_markup as sm
    assert [sm.bucket_for("MATCHED", s) for s in (ba.APPROVAL_NO, ba.APPROVAL_PENDING, ba.APPROVAL_REQUIRES,
                                                  ba.APPROVAL_CHECK, ba.APPROVAL_YES)] == \
        ["approve", "approve", "approve", "checkqbo", "matched"]
