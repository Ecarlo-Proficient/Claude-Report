"""shared/xlsx_verify frozen-pane rule: one-axis freezes keep one selection; a
both-axes freeze passes only in Excel's own three-selection form."""
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "project-pnl"))

from shared.xlsx_verify import assert_clean  # noqa: E402


def _book(tmp_path, freeze, fix=False):
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "x"
    ws.freeze_panes = freeze
    if fix:
        import cp_overview
        cp_overview._excel_selections(ws)
    p = tmp_path / "b.xlsx"
    wb.save(p)
    return p


def test_rows_only_freeze_passes(tmp_path):
    assert_clean(_book(tmp_path, "A6"))


def test_openpyxl_default_two_way_freeze_rejected(tmp_path):
    with pytest.raises(ValueError, match="frozen-pane"):
        assert_clean(_book(tmp_path, "C6"))


def test_excel_form_two_way_freeze_passes(tmp_path):
    assert_clean(_book(tmp_path, "C6", fix=True))


def _dv_book(tmp_path, formula):
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "x"
    dv = DataValidation(type="list", formula1=formula)
    dv.add("B2:B9")
    ws.add_data_validation(dv)
    p = tmp_path / "dv.xlsx"
    wb.save(p)
    return p


@pytest.mark.parametrize("formula", [
    "IF(ISREF(G4 col_Flat),list_Flat,list_Fnd)",   # intersection: Excel stripped it 2026-10-06
    "Phases[PHASE]",                                  # structured ref straight in the rule
    "$A$1:$A$3,$C$1:$C$3",                            # union
    "IF(" + "A1=1," * 70 + "1)",                      # over 255 characters
])
def test_dropdown_rule_excel_rejects_fails(tmp_path, formula):
    with pytest.raises(ValueError, match="dropdown"):
        assert_clean(_dv_book(tmp_path, formula))


@pytest.mark.parametrize("formula", ["list_Flat", "$A$1:$A$9", '"YES,NO"', "INDIRECT(\"x\")"])
def test_dropdown_rule_excel_accepts_passes(tmp_path, formula):
    assert_clean(_dv_book(tmp_path, formula))
