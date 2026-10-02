"""Open_Invoices.xlsx aging tabs + Lease Invoices tab (the user 2026-10-01).

Pure logic - hand-built invoice records and a fake QBO mirror, no Notion, no QBO.
  * columns: Open Balance + Total Amount right after Due Date, ONE Aging column,
    the lien columns last
  * client -> invoices with no project rows, except JPI on the MFD tab
  * the lease tab = every open QBO invoice with no project #, nothing else
"""
import datetime as dt
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))                    # repo root -> shared
sys.path.insert(0, str(ROOT / "invoice-sync"))   # the modules under test

import aging_sheet as ag  # noqa: E402
import export_invoices_xlsx as ex  # noqa: E402
from shared import lien_clock, xlsx_verify  # noqa: E402

TODAY = dt.date(2026, 10, 1)


def _rec(div, parent, proj, inv, days_past_due, bal=100.0):
    issued = TODAY - dt.timedelta(days=days_past_due + 30)
    return dict(
        parent=parent, division=div, project_num=proj, invoice_num=inv,
        invoice_date=issued, due_date=issued + dt.timedelta(days=30),
        days_past_due=days_past_due, open_balance=bal, total_amount=bal, memo=f"{proj} draw",
        lien=lien_clock.lien_state(div, issued, TODAY, memo="", note=""), lien_status="",
        notes="", last_action=None, qbo_link="", prev_draw="",
        vendor_status=ag.PREV_FIRST, vendor_bills=None, vendor_amount=None, this_draw_amount=None,
    )


def _sheet(records, division, tmp_path, **kw):
    wb = Workbook()
    ws = wb.active
    ag.build_aging_sheet(
        ws, records, today=TODAY, litigation_excluded=0,
        drop_columns=ag.RP_DROP_COLUMNS if division == "RP" else (),
        split_clients=ag.PROJECT_SPLIT_CLIENTS.get(division), **kw,
    )
    out = tmp_path / "t.xlsx"
    wb.save(out)
    xlsx_verify.assert_clean(out)
    return ws


def _header(ws):
    return [c.value for c in ws[ag.HEADER_ROW] if c.value]


def _labels(ws):
    """Column A from the ALL CLIENTS row down to TOTAL."""
    out = []
    for r in range(ag.HEADER_ROW + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        out.append(v)
        if v == "TOTAL":
            break
    return out


def test_column_order(tmp_path):
    ws = _sheet([_rec("RP", "BUILDER", "RP6001", "1", 10)], "RP", tmp_path)
    h = _header(ws)
    assert h[1:8] == ["Invoice #", "Date", "Open Balance", "Total Amount", "Due Date", "Aging", "Project #"]
    assert h[-3:] == ["Lien", "Lien status", "Litigation"]
    assert "Division" not in h and "Prev Draw" not in h


def test_aging_column_holds_the_bucket_name(tmp_path):
    recs = [_rec("RP", "BUILDER", "RP6001", "1", -3), _rec("RP", "BUILDER", "RP6002", "2", 45),
            _rec("RP", "BUILDER", "RP6003", "3", 120)]
    ws = _sheet(recs, "RP", tmp_path)
    col = _header(ws).index("Aging") + 1
    ages = [ws.cell(row=r, column=col).value for r in range(5, 8)]   # 3 all, 4 client name
    assert ages == ["90+", "31-60", "Current"]          # oldest invoice first


def test_rp_is_client_then_invoices(tmp_path):
    recs = [_rec("RP", "BUILDER", "RP6001", "1", 5), _rec("RP", "BUILDER", "RP6002", "2", 5)]
    ws = _sheet(recs, "RP", tmp_path)
    assert _labels(ws) == ["ALL CLIENTS (1)", "BUILDER", "RP6001 draw", "RP6002 draw", "Total BUILDER", "TOTAL"]


def test_jpi_on_mfd_keeps_project_rows_others_do_not(tmp_path):
    recs = [_rec("MFD", "JPI Construction, LLC", "MFD101", "1", 5),
            _rec("MFD", "JPI Construction, LLC", "MFD202", "2", 5),
            _rec("MFD", "OTHER GC", "MFD303", "3", 5),
            _rec("MFD", "OTHER GC", "MFD304", "4", 5)]
    ws = _sheet(recs, "MFD", tmp_path)
    assert _labels(ws) == [
        "ALL CLIENTS (2)",
        "JPI Construction, LLC", "MFD101", "MFD101 draw", "Total MFD101", "MFD202", "MFD202 draw",
        "Total MFD202", "Total JPI Construction, LLC",
        "OTHER GC", "MFD303 draw", "MFD304 draw", "Total OTHER GC",
        "TOTAL",
    ]


def test_lease_tab_takes_only_invoices_with_no_project(monkeypatch):
    def line(item, desc=""):
        return {"DetailType": "SalesItemLineDetail", "Description": desc,
                "SalesItemLineDetail": {"ItemRef": {"name": item}}}
    invoices = [
        {"Id": "1", "DocNumber": "100", "CustomerRef": {"value": "9", "name": "PUMP CO"},
         "TxnDate": "2026-09-01", "DueDate": "2026-09-15", "Balance": 50, "TotalAmt": 50,
         "Line": [line("Other Charges:Equipment Lease", "Monthly lease")]},
        {"Id": "2", "DocNumber": "101", "CustomerRef": {"value": "8", "name": "GC:RP6001 - JOB"},
         "TxnDate": "2026-09-01", "DueDate": "2026-09-15", "Balance": 70, "TotalAmt": 70,
         "Line": [line("PC00s:PC00")]},
        {"Id": "3", "DocNumber": "102", "CustomerRef": {"value": "9", "name": "PUMP CO"},
         "PrivateNote": "CP800 retainage", "Balance": 10, "TotalAmt": 10, "Line": []},
    ]
    monkeypatch.setattr(ex.qbo_mirror, "serves", lambda entity: True)
    monkeypatch.setattr(ex.qbo_mirror, "query", lambda entity, where="": invoices)
    monkeypatch.setattr(ex.qbo_mirror, "load", lambda entity: [
        {"Id": "9", "DisplayName": "PUMP CO"}, {"Id": "8", "DisplayName": "GC"}])
    rows = ex._lease_records(TODAY, "")
    assert [r["invoice_num"] for r in rows] == ["100"]
    assert rows[0]["project_num"] == "Equipment Lease"
    assert rows[0]["parent"] == "PUMP CO"
    assert rows[0]["days_past_due"] == 16


def test_lease_tab_skipped_without_a_mirror(monkeypatch):
    monkeypatch.setattr(ex.qbo_mirror, "serves", lambda entity: False)
    assert ex._lease_records(TODAY, "") is None


def test_lease_sheet_has_no_lien_or_draw_columns(tmp_path):
    wb = Workbook()
    ws = wb.active
    rec = dict(parent="PUMP CO", division="LEASE", project_num="Equipment Lease", invoice_num="100",
               invoice_date=TODAY, due_date=TODAY, days_past_due=0, open_balance=50.0,
               total_amount=50.0, memo="lease", qbo_link="")
    ag.build_aging_sheet(ws, [rec], today=TODAY, litigation_excluded=None,
                         drop_columns=ex.LEASE_DROP_COLUMNS, header_overrides={ag.C_PROJ: "Item"})
    out = tmp_path / "lease.xlsx"
    wb.save(out)
    xlsx_verify.assert_clean(out)
    assert _header(ws) == ["Client / Invoice", "Invoice #", "Date", "Open Balance",
                           "Total Amount", "Due Date", "Aging", "Item"]


def test_first_sheet_is_frozen():
    """The `Open Invoices` tab is copied into Outlook for clients (the user
    2026-10-01: "make sure not to change the first sheet"). Its columns and
    order are pinned; changing them needs the owner's say-so."""
    assert [c[0] for c in ex.COLUMNS] == [
        "Division", "Project #", "Client", "Date", "Invoice #", "Net Terms", "Due Date",
        "Past Due", "Memo", "Total Amount", "Open Balance", "Status", "Aging Bucket",
    ]


def test_expanded_even_rows_and_totals_follow_the_filter(tmp_path):
    lit = dict(_rec("RP", "BUILDER", "RP6002", "2", 5, bal=40.0), litigation=True)
    ws = _sheet([_rec("RP", "BUILDER", "RP6001", "1", 5, bal=60.0), lit], "RP", tmp_path)
    for r in range(3, 7):
        assert not ws.row_dimensions[r].hidden            # expanded by default
        assert ws.row_dimensions[r].height in (ag.ROW_PT, ag.CLIENT_ROW_PT)
    h = _header(ws)
    ob = h.index("Open Balance") + 1
    assert ws.cell(row=3, column=ob).value == "=SUBTOTAL(9,D4:D7)"     # ALL CLIENTS (top)
    assert ws.cell(row=4, column=ob).value in (None, "")                # client NAME row: no total
    assert ws.cell(row=7, column=1).value == "Total BUILDER"
    assert ws.cell(row=7, column=ob).value == "=SUBTOTAL(9,D5:D6)"     # client total at the BOTTOM
    litig = h.index("Litigation") + 1
    assert [ws.cell(row=r, column=litig).value for r in (5, 6)] == ["No", "Yes"]
    notes = ws.cell(row=5, column=h.index("Notion Notes") + 1)
    assert not notes.alignment.wrap_text                   # a long note never grows its row


def test_all_open_groups_by_client_within_division(tmp_path):
    lease = dict(parent="PUMP CO", division="Lease", project_num="Equipment Lease", invoice_num="9",
                 invoice_date=TODAY, due_date=TODAY, days_past_due=0, open_balance=5.0,
                 total_amount=5.0, memo="lease", qbo_link="")
    recs = [dict(_rec("CP", "GC", "CP800", "1", 5), litigation=True), _rec("MFD", "GC", "MFD1", "2", 5), lease]
    wb = Workbook()
    ws = wb.active
    ag.build_aging_sheet(ws, recs, today=TODAY, litigation_excluded=None,
                         drop_columns=ex.ALL_OPEN_DROP_COLUMNS, group_by_division=True)
    out = tmp_path / "all.xlsx"
    wb.save(out)
    xlsx_verify.assert_clean(out)
    h = _header(ws)
    assert h[-2:] == ["Division", "Litigation"]
    div = h.index("Division") + 1
    assert _labels(ws) == ["ALL CLIENTS (2)", "GC", "CP800 draw", "Total GC", "GC", "MFD1 draw", "Total GC",
                           "PUMP CO", "lease", "Total PUMP CO", "TOTAL"]
    # every row of a group carries its division, so a Division filter keeps the group whole
    assert [ws.cell(row=r, column=div).value for r in range(4, 13)] == \
        ["CP", "CP", "CP", "MFD", "MFD", "MFD", "Lease", "Lease", "Lease"]


def test_invoices_oldest_first_within_a_client(tmp_path):
    recs = [_rec("RP", "BUILDER", "RP6001", "1", 5), _rec("RP", "BUILDER", "RP6002", "2", 90),
            _rec("RP", "BUILDER", "RP6001", "3", 40)]
    ws = _sheet(recs, "RP", tmp_path)
    assert _labels(ws)[2:5] == ["RP6002 draw", "RP6001 draw", "RP6001 draw"]
    inv = _header(ws).index("Invoice #") + 1
    assert [ws.cell(row=r, column=inv).value for r in (5, 6, 7)] == ["2", "3", "1"]


def test_headers_leave_room_for_the_filter_arrow(tmp_path):
    ws = _sheet([_rec("RP", "BUILDER", "RP6001", "1", 5)], "RP", tmp_path)
    for col, name in enumerate(_header(ws), start=1):
        width = ws.column_dimensions[ws.cell(row=ag.HEADER_ROW, column=col).column_letter].width
        assert width >= len(name) * 1.3 + 3, name
        assert ws.cell(row=ag.HEADER_ROW + 2, column=col).alignment.vertical == "center"


def test_tight_top_and_nothing_under_total(tmp_path):
    ws = _sheet([_rec("RP", "BUILDER", "RP6001", "1", 5)], "RP", tmp_path, title="RP AGING")
    assert ws["A1"].value.startswith("RP AGING - as of ")
    assert "open invoices" in ws["B1"].value
    assert ws["A2"].value == "Client / Invoice"
    last = max(r for r in range(1, ws.max_row + 1) if any(c.value not in (None, "") for c in ws[r]))
    assert ws.cell(row=last, column=1).value == "TOTAL"          # no KEY / BY AGE / notes below


def test_tab_order_opens_on_all_open():
    wb = Workbook()
    wb.active.title = "Open Invoices"
    for name in ("CP Aging", "MFD Aging", "RP Aging", ex.LEASE_SHEET, "Cash Flow"):
        wb.create_sheet(name)
    wb.create_sheet(ex.ALL_OPEN_SHEET, 1)
    ex._arrange_tabs(wb)
    assert wb.sheetnames == ["All Open", "CP Aging", "MFD Aging", "RP Aging", "Lease Invoices",
                             "Open Invoices", "Cash Flow"]
    assert wb.active.title == "All Open"
    assert [ws.sheet_view.tabSelected for ws in wb.worksheets] == [True] + [False] * 6
