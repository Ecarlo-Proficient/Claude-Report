"""bill-tracker: the bill's memo is a column on the Bills + Inventory sheets (owner 10/14/2026),
carried through the bill-grain collapse, written as ONE line, widths keyed by header name."""
import datetime as dt
import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "bill-tracker"))
_spec = importlib.util.spec_from_file_location("ebs", ROOT / "bill-tracker" / "excel_bill_sync.py")
ebs = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(ebs)
import bill_rows  # noqa: E402


def _line(bill_id, line_id, project, amount, memo="RP7001 - slab\nNOT APPROVED"):
    return {"key": f"{bill_id}-{line_id}", "bill_id": bill_id, "line_id": line_id,
            "bill_date": dt.date(2026, 10, 1), "vendor": "A Vendor", "vendor_id": "9",
            "bill_doc": "INV-1", "po_num": "", "bill_total": 300.0, "bill_balance": 300.0,
            "division": "RP", "project_num": project, "bill_type": "COGS", "account": "Job Materials",
            "cost_code": "SL1", "line_amount": amount, "line_desc": f"line {line_id}",
            "bill_memo": memo, "inv_doc": "", "inv_id": "", "inv_memo": "", "inv_date": None,
            "inv_total": None, "inv_balance": None, "payment_date": None, "approved": True,
            "approval": "approved", "bill_created": None, "gc_name": "GC", "customer_name": "GC:RP7001",
            "class_name": "Residential", "auto_status": "Unpaid", "pay_status": "Unpaid",
            "invoice_status": "Awaiting Invoice", "match_basis": "", "match_note": ""}


def test_column_sits_after_line_description():
    i = ebs.HEADERS.index("Line Description")
    assert ebs.HEADERS[i + 1] == "Bill Memo"
    assert ebs.KINDS[i + 1] == "text"
    # every width maps onto a real column, and the divider / key indexes still line up
    assert set(ebs.COL_WIDTHS) <= set(range(1, len(ebs.HEADERS) + 1))
    assert ebs.KEY_COL_INDEX == len(ebs.HEADERS)
    assert sorted(ebs.DIVIDER_COL_INDEXES) == [ebs.HEADERS.index("│") + 1, ebs.HEADERS.index("┃") + 1]


def test_collapse_carries_the_memo():
    rows = [_line("1", "1", "RP7001", 100.0), _line("1", "2", "RP7002", 200.0)]
    bill = bill_rows.collapse_rows(rows, grain="bill")
    assert len(bill) == 1
    assert bill[0]["bill_memo"] == "RP7001 - slab\nNOT APPROVED"
    assert bill[0]["project_num"] == bill_rows.MULTI_MARKER
    assert bill_rows.multi_project_bill_ids(rows) == {"1"}


def test_row_writer_puts_the_memo_in_its_column_as_one_line():
    wb = Workbook()
    ws = wb.active
    ebs._apply_header(ws)
    ebs._write_bill_row(ws, 3, _line("1", "1", "RP7001", 300.0), {}, dt.date(2026, 10, 14))
    col = ebs.HEADERS.index("Bill Memo") + 1
    assert ws.cell(row=3, column=col).value == "RP7001 - slab / NOT APPROVED"
    assert ws.cell(row=2, column=col).value == "Bill Memo"
    assert ws.cell(row=3, column=col + 1).value == 300.0            # Line Amount still follows
    assert ws.cell(row=3, column=ebs.KEY_COL_INDEX).value == "1-1"  # _Key still last
    assert ws.column_dimensions[ebs._col_letter(col)].width == 32


def test_approved_is_one_icon_and_the_ledger_reads_the_word_back():
    from shared import bill_approval as ba
    col = ebs.HEADERS.index("Approved") + 1
    assert ebs.KINDS[col - 1] == "icon"
    assert ebs.COL_WIDTHS[col] <= 6
    wb = Workbook()
    ws = wb.active
    ebs._apply_header(ws)
    for i, (state, icon) in enumerate((("approved", ba.ICON_YES), ("not approved", ba.ICON_NO),
                                       ("pending approval", ba.ICON_NO), ("requires approval", ba.ICON_NO),
                                       ("check QBO", ba.ICON_CHECK)), start=3):
        r = _line("1", str(i), "RP7001", 1.0)
        r["approval"] = state
        r["approved"] = state not in ba.UNAPPROVED_STATES
        ebs._write_bill_row(ws, i, r, {}, dt.date(2026, 10, 14))
        assert ws.cell(row=i, column=col).value == icon
        # the loader turns the icon back into the word the ledger filters on
        assert ba.icon_state(icon) == ({"pending approval": "not approved", "requires approval": "not approved"}.get(state, state))
    assert ba.icon_state("approved") == "approved"      # older files / the Bill List pass through
    assert ba.icon_state(None) is None
