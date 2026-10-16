"""shared/xlsx_save - a workbook people work in is only ever replaced by one that passed every check,
and replacing it keeps the file itself (its owner and permissions on the office server)."""
import os
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from shared.xlsx_save import NotSaved, save_verified, _temp_for  # noqa: E402


def _book(value):
    wb = Workbook()
    wb.active["A1"] = value
    return wb


def test_first_save_creates_the_file(tmp_path):
    path = tmp_path / "Tracker.xlsx"
    save_verified(_book("new"), path)
    assert load_workbook(path).active["A1"].value == "new"
    assert not _temp_for(path).exists()


def test_replacing_keeps_the_same_file_and_its_permissions(tmp_path):
    path = tmp_path / "Tracker.xlsx"
    save_verified(_book("old"), path)
    os.chmod(path, 0o640)
    before = os.stat(path)
    save_verified(_book("new"), path)
    after = os.stat(path)
    assert load_workbook(path).active["A1"].value == "new"
    assert after.st_ino == before.st_ino                  # written into, not swapped for a new file
    assert after.st_mode & 0o777 == 0o640


def test_a_failed_check_leaves_the_file_untouched(tmp_path):
    path = tmp_path / "Tracker.xlsx"
    save_verified(_book("good"), path)
    good = path.read_bytes()
    with pytest.raises(NotSaved, match="bad table"):
        save_verified(_book("broken"), path, checks=lambda p: ["bad table"])
    assert path.read_bytes() == good
    assert not _temp_for(path).exists()


def test_the_repo_wide_excel_check_always_runs(tmp_path, monkeypatch):
    from shared import xlsx_save

    def boom(p):
        raise ValueError("would trigger Excel's repair prompt")
    monkeypatch.setattr(xlsx_save.xlsx_verify, "assert_clean", boom)
    path = tmp_path / "Tracker.xlsx"
    with pytest.raises(NotSaved, match="repair prompt"):
        save_verified(_book("x"), path)
    assert not path.exists()


def test_post_process_runs_on_the_temp_before_the_checks(tmp_path):
    path = tmp_path / "Tracker.xlsx"
    seen = []
    save_verified(_book("x"), path, post_process=lambda p: seen.append(("post", p.name)),
                  checks=lambda p: seen.append(("check", p.name)) or [])
    assert [s[0] for s in seen] == ["post", "check"]
    assert all(name != path.name for _, name in seen)    # never on the real file
